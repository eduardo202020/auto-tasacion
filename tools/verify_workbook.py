"""Verifica el contrato estructural de un Resultado_Final.xlsx generado.

No evalúa el contenido de negocio de un PDF. Comprueba que un archivo que
Power Automate pretende consumir conserva las hojas, tablas y restricciones
mínimas de la interfaz versionada.
"""
from __future__ import annotations

import argparse
from pathlib import Path
import re
import sys

from openpyxl import load_workbook


SHEETS = ("PARA_PROCESAR", "REVISION_IA", "CONTROL")
TABLES = {
    "PARA_PROCESAR": "tblParaProcesar",
    "REVISION_IA": "tblRevisionIa",
    "CONTROL": "tblControl",
}
PARA_PROCESAR_HEADERS = (
    "PRESTAMO", "TIPO DE INMUEBLE", "VALOR DEL BIEN", "MONEDA", "IMPORTE", "COL_F",
    "DIRECCION", "DIRECCION1", "EXTERIOR", "INTERIOR", "REFERENCIA", "COL_L",
    "UBICACION", "UBICACION1", "MUNICIPIO", "DIST_COD", "COL_Q", "PROV_COD",
    "COL_S", "DEPT_COD", "CLASE", "PISOS", "SOTANOS", "AÑO", "ID_CASO",
)
REQUIRED_READY_FIELDS = (
    "TIPO DE INMUEBLE", "MONEDA", "IMPORTE", "DIRECCION",
    "DIRECCION1", "MUNICIPIO", "DIST_COD", "PROV_COD", "DEPT_COD", "CLASE",
    "PISOS", "SOTANOS", "AÑO", "ID_CASO",
)
IMPORTE_FORMAT = re.compile(r"^\d{1,3}(?:,\d{3})*\.\d{2}$")


def is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def validate(path: Path) -> list[str]:
    """Devuelve incidencias de contrato; una lista vacía representa éxito."""
    issues: list[str] = []
    if not path.is_file():
        return [f"No existe el archivo: {path}"]
    try:
        workbook = load_workbook(path, data_only=True)
    except Exception as error:  # openpyxl normaliza distintos errores OOXML.
        return [f"No se pudo abrir el XLSX: {error}"]

    if tuple(workbook.sheetnames) != SHEETS:
        issues.append(f"Hojas inválidas: se esperaba {SHEETS}, se obtuvo {tuple(workbook.sheetnames)}")

    for sheet, table in TABLES.items():
        if sheet not in workbook.sheetnames:
            continue
        if table not in workbook[sheet].tables:
            issues.append(f"Falta la tabla {table} en la hoja {sheet}")

    if "PARA_PROCESAR" not in workbook.sheetnames:
        return issues

    worksheet = workbook["PARA_PROCESAR"]
    headers = tuple(cell.value for cell in worksheet[1])
    if headers != PARA_PROCESAR_HEADERS:
        issues.append("Los encabezados u orden de PARA_PROCESAR no coinciden con el contrato")
        return issues

    index = {header: position for position, header in enumerate(headers)}
    for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if all(is_blank(value) for value in values):
            continue
        missing = [field for field in REQUIRED_READY_FIELDS if is_blank(values[index[field]])]
        if missing:
            issues.append(f"Fila {row_number} de PARA_PROCESAR tiene campos requeridos vacíos: {', '.join(missing)}")
        importe = values[index["IMPORTE"]]
        if not is_blank(importe) and not IMPORTE_FORMAT.fullmatch(str(importe)):
            issues.append(f"Fila {row_number} de PARA_PROCESAR tiene un IMPORTE fuera del formato #,##0.00")
        case_id = values[index["ID_CASO"]]
        if not is_blank(case_id) and not str(case_id).startswith("TAS-"):
            issues.append(f"Fila {row_number} de PARA_PROCESAR tiene un ID_CASO inválido")

    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description="Valida la estructura del XLSX de tasaciones")
    parser.add_argument("workbook", type=Path, help="Ruta a Resultado_Final.xlsx")
    args = parser.parse_args()

    issues = validate(args.workbook)
    if issues:
        for issue in issues:
            print(f"ERROR: {issue}", file=sys.stderr)
        return 1
    print(f"OK: {args.workbook} cumple el contrato estructural de tasaciones")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
