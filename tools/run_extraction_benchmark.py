"""Ejecuta el extractor sobre un ZIP externo sin copiar PDFs al repositorio.

El informe JSON contiene solamente metadatos técnicos, decisiones, importes y
códigos de validación. No incluye texto de los documentos ni metadata
operativa sensible como el préstamo o el expediente.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
import re
from pathlib import Path
import sys
import zipfile
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "cloud-run"))

from service import PARA_PROCESAR_COLUMNS, build_workbook, process_pdf_entries  # noqa: E402
from verify_workbook import validate as validate_workbook  # noqa: E402


GOLDEN_FILENAMES = tuple(f"D{number:02d}.pdf" for number in range(1, 11))
# The operator workbook's calculated headers can be represented differently by
# Excel readers. The golden's documented A:Y contract is stable, so only this
# benchmark maps its semantic fields by those positions. The corpus-150 path
# still uses its explicit control headers.
GOLDEN_MASIVO_INDEX = {
    "tipo": 1,
    "moneda": 3,
    "importe": 4,
    "direccion": 5,
    "direccion1": 7,
    "exterior": 8,
    "interior": 9,
    "referencia": 10,
    "ubicacion": 11,
    "ubicacion1": 13,
    "municipio": 14,
    "clase": 20,
    "pisos": 21,
    "sotanos": 22,
    "anio": 23,
}
SAFE_CONTROL_FIELDS = (
    "Tipo documento", "Motivo tipo documento", "Tasadora id", "Tasadora detectada",
    "Origen tasadora", "Perfil plantilla", "Version perfil", "Template id", "Familia layout",
    "Estrategia extraccion", "Valor elegido tipo", "Valor elegido US$", "Valor elegido S/",
    "Valor comercial US$", "Valor comercial S/", "Candidatos valor comercial",
    "Valor reconstruccion US$", "Valor reconstruccion S/", "Candidatos valor reconstruccion",
    "Tipo inmueble", "Nro pisos edificio", "Nro sotanos edificio", "Año construccion",
    "Pagina direccion", "Pagina tipo inmueble", "Pagina valor comercial", "Pagina valor reconstruccion",
    "Pagina pisos/sotanos", "Ruta final", "Campos faltantes", "Incidencias de validacion",
)


def _nonempty_rows(worksheet) -> Iterable[dict[str, Any]]:
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows, ())]
    for values in rows:
        if not any(value not in (None, "") for value in values):
            continue
        yield {headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))}


def control_from_workbook(path: Path) -> dict[str, dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    candidates = ("tblControlLote", "CONTROL_LOTE")
    sheet = next((workbook[name] for name in candidates if name in workbook.sheetnames), None)
    if sheet is None:
        raise ValueError("La hoja de control debe llamarse tblControlLote o CONTROL_LOTE")
    controls: dict[str, dict[str, str]] = {}
    for row in _nonempty_rows(sheet):
        filename = Path(str(row.get("ARCHIVO_PDF") or "")).name
        if not filename:
            continue
        controls[filename] = {
            "id_ingreso": str(row.get("ID_INGRESO") or ""),
            "prestamo": str(row.get("PRESTAMO") or ""),
            "moneda_operador": str(row.get("MONEDA_OPERADOR") or ""),
            "seguro_origen": str(row.get("SEGURO_ORIGEN") or ""),
            "nro_expediente": str(row.get("NRO_EXPEDIENTE") or ""),
        }
    return controls


def baseline_routes(path: Path | None) -> dict[str, str]:
    if path is None:
        return {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "CONTROL" not in workbook.sheetnames:
        raise ValueError("El workbook baseline no contiene CONTROL")
    routes: dict[str, str] = {}
    for row in _nonempty_rows(workbook["CONTROL"]):
        filename = Path(str(row.get("PDF_Archivo") or "")).name
        if filename:
            routes[filename] = str(row.get("Ruta final") or "")
    return routes


def golden_control(path: Path) -> dict[str, dict[str, str]]:
    """Lee el golden en el orden D01..D10, sin usarlo para otro dataset."""
    workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=True)
    if "MASIVO" not in workbook.sheetnames:
        raise ValueError("El golden del operador debe contener la hoja MASIVO")
    rows = list(_nonempty_rows(workbook["MASIVO"]))[: len(GOLDEN_FILENAMES)]
    if len(rows) != len(GOLDEN_FILENAMES):
        raise ValueError("El golden no contiene diez filas operativas")
    controls: dict[str, dict[str, str]] = {}
    for filename, row in zip(GOLDEN_FILENAMES, rows):
        controls[filename] = {
            "id_ingreso": filename.rsplit(".", 1)[0],
            "prestamo": str(row.get("PRESTAMO") or ""),
            "moneda_operador": str(row.get("MONEDA") or ""),
            "seguro_origen": str(row.get("SEGURO INMUEBLE") or ""),
            "nro_expediente": "",
        }
    return controls


def golden_expected(path: Path | None) -> dict[str, dict[str, Any]]:
    if path is None:
        return {}
    workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=True)
    rows = list(_nonempty_rows(workbook["MASIVO"]))[: len(GOLDEN_FILENAMES)]

    def value(row: dict[str, Any], key: str) -> Any:
        values = list(row.values())
        index = GOLDEN_MASIVO_INDEX[key]
        return values[index] if index < len(values) else ""

    return {
        filename: {
            **{key: value(row, key) for key in ("tipo", "moneda", "importe", "pisos", "sotanos", "anio")},
            "direccion_estructurada": {
                key: value(row, key)
                for key in ("direccion", "direccion1", "exterior", "interior", "referencia", "ubicacion", "ubicacion1")
            },
            "geografia": {"municipio": value(row, "municipio")},
            "clase": value(row, "clase"),
            "decision_final": "LISTO_DETERMINISTA",
        }
        for filename, row in zip(GOLDEN_FILENAMES, rows)
    }

def _normal_amount(value: object) -> str:
    try:
        return f"{float(str(value).replace(',', '')):.2f}"
    except (TypeError, ValueError):
        return ""


def _safe_candidates(value: object) -> list[dict[str, Any]]:
    try:
        parsed = json.loads(str(value or "[]"))
    except json.JSONDecodeError:
        return []
    if not isinstance(parsed, list):
        return []
    allowed = {"field", "currency", "amount", "page", "anchor", "section", "source_priority"}
    return [{key: item.get(key) for key in allowed if key in item} for item in parsed if isinstance(item, dict)]


def _entry_stream(zip_paths: list[Path], controls: dict[str, dict[str, str]], descriptors: list[dict[str, str]]):
    for zip_path in zip_paths:
        with zipfile.ZipFile(zip_path) as archive:
            for member in archive.infolist():
                if member.is_dir() or not member.filename.lower().endswith(".pdf"):
                    continue
                filename = Path(member.filename).name
                content = archive.read(member)
                descriptors.append({
                    "filename": filename,
                    "source_member": member.filename,
                    "sha256": hashlib.sha256(content).hexdigest(),
                })
                yield member.filename, content, controls.get(filename, {})


def _result_entry(descriptor: dict[str, str], control: dict[str, Any], baseline: dict[str, str]) -> dict[str, Any]:
    fields = {field: control.get(field, "") for field in SAFE_CONTROL_FIELDS}
    return {
        "sha256": descriptor["sha256"],
        "filename": descriptor["filename"],
        "document_type": fields.pop("Tipo documento", ""),
        "document_reason": fields.pop("Motivo tipo documento", ""),
        "provider_id": fields.pop("Tasadora id", ""),
        "provider_name": fields.pop("Tasadora detectada", ""),
        "provider_source": fields.pop("Origen tasadora", ""),
        "profile_id": fields.pop("Perfil plantilla", ""),
        "template_id": fields.pop("Template id", ""),
        "layout_family": fields.pop("Familia layout", ""),
        "strategy": fields.pop("Estrategia extraccion", ""),
        "commercial_candidates": _safe_candidates(fields.pop("Candidatos valor comercial", "")),
        "reconstruction_candidates": _safe_candidates(fields.pop("Candidatos valor reconstruccion", "")),
        "route": fields.pop("Ruta final", ""),
        "baseline_route": baseline.get(descriptor["filename"], ""),
        "fields": fields,
    }


def _macro_rows_by_case(ready_rows: list[list[Any]]) -> dict[str, dict[str, Any]]:
    """Build final-row semantics only in memory for golden comparisons."""
    case_index = PARA_PROCESAR_COLUMNS.index("ID_CASO")
    return {
        str(row[case_index]): dict(zip(PARA_PROCESAR_COLUMNS, row))
        for row in ready_rows
        if len(row) > case_index and row[case_index]
    }


def _normalized_text(value: object) -> str:
    return " ".join(str(value or "").upper().split())


def _field_with_suffix(fields: dict[str, Any], suffix: str) -> Any:
    """Read a stable semantic field when a legacy header has accent encoding variance."""
    return next((value for name, value in fields.items() if name.endswith(suffix)), "")


def _golden_comparison(
    entries: list[dict[str, Any]],
    expected: dict[str, dict[str, Any]],
    macro_rows_by_filename: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    compared: list[dict[str, Any]] = []
    for entry in entries:
        reference = expected.get(entry["filename"])
        if reference is None:
            continue
        fields = entry["fields"]
        macro = macro_rows_by_filename.get(entry["filename"], {})
        selected_amount = fields.get("Valor elegido S/") or fields.get("Valor elegido US$")
        actual = {
            "tipo": fields.get("Tipo inmueble", ""),
            "moneda": "PEN" if fields.get("Valor elegido S/") not in ("", None) else "USD" if fields.get("Valor elegido US$") not in ("", None) else "",
            "importe": _normal_amount(selected_amount),
            "pisos": fields.get("Nro pisos edificio", ""),
            "sotanos": fields.get("Nro sotanos edificio", ""),
            "anio": _field_with_suffix(fields, "o construccion"),
            "direccion_estructurada": {
                "direccion": macro.get("DIRECCION", ""),
                "direccion1": macro.get("DIRECCION1", ""),
                "exterior": macro.get("EXTERIOR", ""),
                "interior": macro.get("INTERIOR", ""),
                "referencia": macro.get("REFERENCIA", ""),
                "ubicacion": macro.get("UBICACION", ""),
                "ubicacion1": macro.get("UBICACION1", ""),
            },
            "geografia": {"municipio": macro.get("MUNICIPIO", "")},
            "clase": macro.get("CLASE", ""),
            "decision_final": entry["route"],
        }
        matches = {
            key: (
                _normal_amount(actual[key]) == _normal_amount(reference[key])
                if key == "importe"
                else str(actual[key]).strip() == str(reference[key]).strip()
            )
            for key in ("tipo", "moneda", "importe", "pisos", "sotanos", "anio", "clase", "decision_final")
        }
        matches["direccion_estructurada"] = all(
            _normalized_text(actual["direccion_estructurada"].get(key)) == _normalized_text(value)
            for key, value in reference["direccion_estructurada"].items()
        )
        matches["geografia"] = all(
            _normalized_text(actual["geografia"].get(key)) == _normalized_text(value)
            for key, value in reference["geografia"].items()
        )
        compared.append({
            "filename": entry["filename"],
            "route": entry["route"],
            "matches": matches,
        })
    return {
        "total": len(compared),
        "critical_matches": {
            field: sum(1 for item in compared if item["matches"].get(field))
            for field in (
                "tipo", "moneda", "importe", "pisos", "sotanos", "anio", "clase",
                "direccion_estructurada", "geografia", "decision_final",
            )
        },
        "cases": compared,
    }

def _counter_with_empty(counter: Counter[str], keys: tuple[str, ...]) -> dict[str, int]:
    return {key: int(counter.get(key, 0)) for key in keys}


def _benchmark_quality(entries: list[dict[str, Any]], controls: list[dict[str, Any]]) -> dict[str, int]:
    ready_routes = {"LISTO_DETERMINISTA", "LISTO_IA_VERIFICADO"}
    ready_nonpositive = 0
    for entry in entries:
        if entry["route"] not in ready_routes:
            continue
        fields = entry["fields"]
        amount = fields.get("Valor elegido S/") or fields.get("Valor elegido US$")
        try:
            ready_nonpositive += int(float(str(amount)) <= 0)
        except (TypeError, ValueError):
            ready_nonpositive += 1

    invalid_loans = sum(
        not re.fullmatch(r"\d{20}", str(row.get("PRESTAMO") or ""))
        for row in controls
    )
    currency_conflicts = sum(
        "MONEDA" in str(entry["fields"].get("Incidencias de validacion") or "").upper()
        for entry in entries
    )
    no_geography = sum(
        any(token in str(entry["fields"].get("Campos faltantes") or "").upper()
            for token in ("MUNICIPIO", "DISTRITO", "PROVINCIA", "DEPARTAMENTO"))
        for entry in entries
    )
    amount_ambiguities = sum(
        "VALOR_" in str(entry["fields"].get("Incidencias de validacion") or "")
        and "AMBIGUO" in str(entry["fields"].get("Incidencias de validacion") or "")
        for entry in entries
    )
    return {
        "ready_with_nonpositive_amount": ready_nonpositive,
        "invalid_control_loans": invalid_loans,
        "currency_conflicts": currency_conflicts,
        "without_geography": no_geography,
        "unknown_documents": sum(entry["document_type"] == "DESCONOCIDO" for entry in entries),
        "amount_ambiguities": amount_ambiguities,
        "provider_ocr_matches": sum(entry["provider_source"] == "OCR_LOCAL" for entry in entries),
    }


def _coverage(entries: list[dict[str, Any]]) -> dict[str, dict[str, int]]:
    return {
        "provider_id": dict(sorted(Counter(entry["provider_id"] or "NO_IDENTIFICADO" for entry in entries).items())),
        "template_id": dict(sorted(Counter(entry["template_id"] or "SIN_TEMPLATE" for entry in entries).items())),
        "layout_family": dict(sorted(Counter(entry["layout_family"] or "SIN_LAYOUT" for entry in entries).items())),
        "profile_id": dict(sorted(Counter(entry["profile_id"] or "SIN_PROFILE" for entry in entries).items())),
    }


def _write_temp_workbook(ready, review, control) -> list[str]:
    """Ejecuta el verificador sin persistir un resultado en el repositorio."""
    import tempfile

    with tempfile.TemporaryDirectory(prefix="auto-tasacion-benchmark-") as directory:
        target = Path(directory) / "Resultado_Final.xlsx"
        target.write_bytes(build_workbook(ready, review, control).getvalue())
        return validate_workbook(target)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dataset", required=True, choices=("golden-10", "corpus-150"))
    parser.add_argument("--zip", dest="zip_paths", type=Path, action="append", required=True)
    parser.add_argument("--control-workbook", type=Path)
    parser.add_argument("--operator-workbook", type=Path)
    parser.add_argument("--baseline-workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True, help="JSON externo al repositorio")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output = args.output.resolve()
    try:
        output.relative_to(ROOT)
    except ValueError:
        pass
    else:
        raise SystemExit("--output debe estar fuera del repositorio para no versionar evidencia")
    if any(not path.is_file() for path in args.zip_paths):
        raise SystemExit("Cada --zip debe existir")

    if args.dataset == "golden-10":
        if args.operator_workbook is None:
            raise SystemExit("golden-10 requiere --operator-workbook")
        controls = golden_control(args.operator_workbook)
    else:
        if args.control_workbook is None:
            raise SystemExit("corpus-150 requiere --control-workbook")
        controls = control_from_workbook(args.control_workbook)

    descriptors: list[dict[str, str]] = []
    ready, review, control = process_pdf_entries(
        _entry_stream(args.zip_paths, controls, descriptors), use_environment_reviewer=False,
    )
    rows_by_filename = {Path(str(row.get("PDF_Archivo") or "")).name: row for row in control}
    baseline = baseline_routes(args.baseline_workbook)
    entries = [_result_entry(item, rows_by_filename.get(item["filename"], {}), baseline) for item in descriptors]
    route_counter = Counter(entry["route"] for entry in entries)
    summary = {
        "total": len(entries),
        "routes": _counter_with_empty(route_counter, (
            "LISTO_DETERMINISTA", "LISTO_IA_VERIFICADO", "PENDIENTE_IA",
            "REGLA_NEGOCIO_PENDIENTE", "DOCUMENTO_NO_TASACION", "REVISION_HUMANA",
        )),
        "document_types": dict(sorted(Counter(entry["document_type"] for entry in entries).items())),
        "coverage": _coverage(entries),
        "quality": _benchmark_quality(entries, control),
        "baseline_comparable": {
            "total": sum(1 for entry in entries if entry["baseline_route"]),
            "before_routes": dict(sorted(Counter(entry["baseline_route"] for entry in entries if entry["baseline_route"]).items())),
            "after_routes": dict(sorted(Counter(entry["route"] for entry in entries if entry["baseline_route"]).items())),
        },
        "workbook_contract_issues": _write_temp_workbook(ready, review, control),
    }
    report: dict[str, Any] = {
        "dataset": args.dataset,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "summary": summary,
        "entries": entries,
    }
    if args.dataset == "golden-10":
        macro_by_case = _macro_rows_by_case(ready)
        macro_by_filename = {
            filename: macro_by_case.get(str(row.get("ID_CASO") or ""), {})
            for filename, row in rows_by_filename.items()
        }
        report["golden_comparison"] = _golden_comparison(
            entries, golden_expected(args.operator_workbook), macro_by_filename,
        )
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"dataset": args.dataset, **summary}, ensure_ascii=False, separators=(",", ":")))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
