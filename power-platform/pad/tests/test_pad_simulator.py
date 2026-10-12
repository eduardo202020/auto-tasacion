from __future__ import annotations

import importlib.util
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


MODULE_PATH = Path(__file__).resolve().parents[1] / "pad_simulator.py"
SPEC = importlib.util.spec_from_file_location("pad_simulator", MODULE_PATH)
assert SPEC is not None and SPEC.loader is not None
pad = importlib.util.module_from_spec(SPEC)
import sys
sys.modules[SPEC.name] = pad
SPEC.loader.exec_module(pad)


class PadSimulationTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.input_directory = self.root / "tasaciones"
        self.runtime_directory = self.input_directory / "_pad_3270_simulado"
        self.input_directory.mkdir()
        self.config_path = self.root / "config.json"
        self.write_config()
        self.contract = pad.load_contract()

    def tearDown(self) -> None:
        self.temp.cleanup()

    def write_config(self, **overrides: object) -> None:
        value = {
            "schema_version": 1,
            "mode": "SIMULACION",
            "input_directory": str(self.input_directory),
            "runtime_directory": str(self.runtime_directory),
            "poll_seconds": 0,
            "stability_checks": 1,
            "stability_interval_seconds": 0,
            "processing_lease_seconds": 60,
            "suspicious_floors_threshold": 100,
            "suspicious_basements_threshold": 20,
            "synthetic_failure_case_ids": [],
        }
        value.update(overrides)
        self.config_path.write_text(json.dumps(value), encoding="utf-8")

    @property
    def settings(self):
        return pad.Settings.load(self.config_path)

    def valid_values(self, case_id: str, **overrides: object) -> list[object]:
        values = {
            "PRESTAMO": None,
            "TIPO DE INMUEBLE": "DEPARTAMENTO",
            "VALOR DEL BIEN": None,
            "MONEDA": "PEN",
            "IMPORTE": "100,000.00",
            "COL_F": None,
            "DIRECCION": "AV. PRUEBA",
            "DIRECCION1": "AV. PRUEBA 123",
            "EXTERIOR": None,
            "INTERIOR": None,
            "REFERENCIA": None,
            "COL_L": None,
            "UBICACION": None,
            "UBICACION1": None,
            "MUNICIPIO": "LIMA",
            "DIST_COD": "001",
            "COL_Q": None,
            "PROV_COD": "01",
            "COL_S": None,
            "DEPT_COD": "15",
            "CLASE": "1",
            "PISOS": 5,
            "SOTANOS": 0,
            self.contract.PARA_PROCESAR_HEADERS[-2]: 2020,
            "ID_CASO": case_id,
        }
        values.update(overrides)
        return [values.get(header) for header in self.contract.PARA_PROCESAR_HEADERS]

    def create_workbook(
        self,
        lot_id: str,
        rows: list[list[object]] | None = None,
        revision_rows: int = 0,
        sheet_names: tuple[str, ...] | None = None,
        headers: tuple[str, ...] | None = None,
    ) -> Path:
        output = self.input_directory / f"Resultado_Final_{lot_id}.xlsx"
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet_names = sheet_names or tuple(self.contract.SHEETS)
        for sheet_name in sheet_names:
            worksheet = workbook.create_sheet(sheet_name)
            if sheet_name == "PARA_PROCESAR":
                current_headers = headers or tuple(self.contract.PARA_PROCESAR_HEADERS)
                worksheet.append(list(current_headers))
                data = rows if rows is not None else [self.valid_values("TAS-CASE-1")]
                for row in data:
                    worksheet.append(row)
                final_column = get_column_letter(len(current_headers))
                table = Table(displayName="tblParaProcesar", ref=f"A1:{final_column}{max(2, len(data) + 1)}")
            elif sheet_name == "REVISION_IA":
                worksheet.append(["ID_CASO"])
                for index in range(revision_rows):
                    worksheet.append([f"TAS-REV-{index + 1}"])
                table = Table(displayName="tblRevisionIa", ref=f"A1:A{max(2, revision_rows + 1)}")
            else:
                worksheet.append(["ID_CASO"])
                worksheet.append([None])
                table = Table(displayName="tblControl", ref="A1:A2")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            worksheet.add_table(table)
        workbook.save(output)
        workbook.close()
        return output

    def scan(self) -> list[dict[str, str]]:
        return pad.LocalMonitor(self.settings).scan_once()

    def read_state(self, lot_id: str) -> dict:
        records = pad.StateStore(self.settings).lot_records(lot_id)
        self.assertEqual(len(records), 1)
        return records[0]

    def run_single_job(self, lot_id: str) -> tuple[str, str]:
        state = self.read_state(lot_id)
        job_path = pad.StateStore(self.settings).job_path(lot_id, state["sha256"])
        return pad.Mock3270Simulator(self.settings).run_job(job_path)

    def test_one_pdf_is_validated_simulated_and_preserves_original(self) -> None:
        source = self.create_workbook("TAS-ONE")
        self.assertEqual(self.scan()[0]["outcome"], "VALIDADO")
        status, outcome = self.run_single_job("TAS-ONE")
        self.assertEqual((status, outcome), ("SIMULADO_FINALIZADO", "claimed"))
        self.assertTrue(source.exists())
        self.assertTrue(list(self.settings.results_directory.glob("*_operaciones.csv")))
        self.assertIn(
            "ALTA_SEGURO_INMUEBLE_SIMULADA",
            {event["event"] for event in self.read_state("TAS-ONE")["events"]},
        )

    def test_multiple_pdfs_have_independent_jobs(self) -> None:
        self.create_workbook("TAS-A")
        self.create_workbook("TAS-B", rows=[self.valid_values("TAS-CASE-B")])
        self.assertEqual({item["outcome"] for item in self.scan()}, {"VALIDADO"})
        self.assertEqual(self.run_single_job("TAS-A")[0], "SIMULADO_FINALIZADO")
        self.assertEqual(self.run_single_job("TAS-B")[0], "SIMULADO_FINALIZADO")

    def test_duplicate_event_does_not_create_a_second_state_or_job(self) -> None:
        self.create_workbook("TAS-DUP")
        self.scan()
        result = self.scan()
        self.assertEqual(result[0]["outcome"], "duplicate")
        self.assertEqual(len(pad.StateStore(self.settings).lot_records("TAS-DUP")), 1)
        self.assertEqual(len(list(self.settings.pending_directory.glob("*.json"))), 1)

    def test_changed_source_for_same_lot_is_quarantined(self) -> None:
        source = self.create_workbook("TAS-CHANGE")
        self.scan()
        self.create_workbook("TAS-CHANGE", rows=[self.valid_values("TAS-CASE-CHANGED", IMPORTE="200,000.00")])
        self.assertEqual(self.scan()[0]["outcome"], "origin_changed")
        self.assertTrue(source.exists())
        states = pad.StateStore(self.settings).lot_records("TAS-CHANGE")
        self.assertIn("ERROR_ORIGEN_CAMBIO", {state["status"] for state in states})

    def test_wrong_name_is_ignored_and_temp_excel_is_ignored(self) -> None:
        self.create_workbook("TAS-NAME").rename(self.input_directory / "resultado.xlsx")
        (self.input_directory / "~$Resultado_Final_TAS-TEMP.xlsx").write_bytes(b"temporary")
        self.assertEqual(self.scan(), [])

    def test_corrupt_xlsx_is_structural_error(self) -> None:
        path = self.input_directory / "Resultado_Final_TAS-CORRUPT.xlsx"
        path.write_bytes(b"not-an-xlsx")
        self.assertEqual(self.scan()[0]["outcome"], "ERROR_ESTRUCTURAL")
        self.assertEqual(self.read_state("TAS-CORRUPT")["status"], "ERROR_ESTRUCTURAL")

    def test_missing_sheet_and_missing_columns_are_structural_errors(self) -> None:
        self.create_workbook("TAS-SHEET", sheet_names=("PARA_PROCESAR", "CONTROL"))
        self.create_workbook("TAS-COLS", headers=tuple(self.contract.PARA_PROCESAR_HEADERS[:-1]))
        outcomes = {entry["file"]: entry["outcome"] for entry in self.scan()}
        self.assertEqual(outcomes["Resultado_Final_TAS-SHEET.xlsx"], "ERROR_ESTRUCTURAL")
        self.assertEqual(outcomes["Resultado_Final_TAS-COLS.xlsx"], "ERROR_ESTRUCTURAL")

    def test_no_valid_rows_requires_manual_review(self) -> None:
        row = self.valid_values("TAS-BAD", IMPORTE="not-a-number")
        self.create_workbook("TAS-NOVALID", rows=[row])
        self.assertEqual(self.scan()[0]["outcome"], "REVISION_MANUAL")
        self.assertFalse(list(self.settings.pending_directory.glob("*.json")))

    def test_revision_rows_and_suspicious_data_do_not_go_to_simulation(self) -> None:
        suspicious = self.valid_values("TAS-SUSPICIOUS", PISOS=101)
        valid = self.valid_values("TAS-VALID")
        self.create_workbook("TAS-REVIEW", rows=[suspicious, valid], revision_rows=2)
        self.scan()
        state = self.read_state("TAS-REVIEW")
        self.assertEqual(state["validation"]["valid_case_ids"], ["TAS-VALID"])
        self.assertEqual(state["validation"]["manual_count"], 1)
        self.assertEqual(self.run_single_job("TAS-REVIEW")[0], "SIMULADO_CON_INCIDENCIAS")

    def test_double_job_start_is_idempotent(self) -> None:
        self.create_workbook("TAS-DOUBLE")
        self.scan()
        self.assertEqual(self.run_single_job("TAS-DOUBLE")[0], "SIMULADO_FINALIZADO")
        state = self.read_state("TAS-DOUBLE")
        filename = pad.StateStore(self.settings).job_path("TAS-DOUBLE", state["sha256"]).name
        completed = self.settings.completed_jobs_directory / filename
        duplicate = self.settings.pending_directory / filename
        shutil.copy2(completed, duplicate)
        self.assertEqual(pad.Mock3270Simulator(self.settings).run_job(duplicate)[1], "terminal")

    def test_configured_simulation_failure_is_reported(self) -> None:
        self.write_config(synthetic_failure_case_ids=["TAS-FAIL"])
        self.create_workbook("TAS-FAILURE", rows=[self.valid_values("TAS-FAIL")])
        self.scan()
        self.assertEqual(self.run_single_job("TAS-FAILURE")[0], "SIMULADO_CON_INCIDENCIAS")

    def test_recovery_claim_and_already_delivered_state(self) -> None:
        self.create_workbook("TAS-RECOVER")
        self.scan()
        state = self.read_state("TAS-RECOVER")
        state["status"] = "SIMULANDO_3270"
        state["updated_at"] = "2000-01-01T00:00:00+00:00"
        pad.atomic_json(pad.StateStore(self.settings).state_path("TAS-RECOVER", state["sha256"]), state)
        self.assertEqual(self.run_single_job("TAS-RECOVER"), ("SIMULADO_FINALIZADO", "recovered"))
        state = self.read_state("TAS-RECOVER")
        state["status"] = "ENTREGADO"
        pad.atomic_json(pad.StateStore(self.settings).state_path("TAS-RECOVER", state["sha256"]), state)
        completed = next(self.settings.completed_jobs_directory.glob("*.json"))
        duplicate = self.settings.pending_directory / completed.name
        shutil.copy2(completed, duplicate)
        self.assertEqual(pad.Mock3270Simulator(self.settings).run_job(duplicate)[1], "terminal")

    def test_explicit_reprocess_is_audited_and_creates_one_new_job(self) -> None:
        self.create_workbook("TAS-REPROCESS")
        self.scan()
        self.run_single_job("TAS-REPROCESS")
        state = self.read_state("TAS-REPROCESS")
        first_reports = set(self.settings.results_directory.glob("*_operaciones.csv"))
        rebuilt = pad.StateStore(self.settings).reprocess("TAS-REPROCESS", state["sha256"], "prueba_auditada")
        self.assertEqual(rebuilt["status"], "VALIDADO")
        self.assertTrue(list(self.settings.history_directory.glob("*.json")))
        self.assertTrue(pad.StateStore(self.settings).job_path("TAS-REPROCESS", state["sha256"]).exists())
        self.assertEqual(self.run_single_job("TAS-REPROCESS")[0], "SIMULADO_FINALIZADO")
        self.assertTrue(first_reports.issubset(set(self.settings.results_directory.glob("*_operaciones.csv"))))
        self.assertEqual(len(list(self.settings.results_directory.glob("*_operaciones.csv"))), 2)

    def test_atomic_json_retries_a_transient_permission_error(self) -> None:
        target = self.root / "locked-state.json"
        original_replace = pad.os.replace
        attempts = 0

        def transient_replace(source, destination):
            nonlocal attempts
            attempts += 1
            if attempts < 3:
                raise PermissionError("synthetic transient lock")
            return original_replace(source, destination)

        with patch.object(pad.os, "replace", side_effect=transient_replace):
            pad.atomic_json(target, {"ok": True}, attempts=3, base_delay_seconds=0)
        self.assertEqual(attempts, 3)
        self.assertEqual(json.loads(target.read_text(encoding="utf-8")), {"ok": True})

    def test_operation_events_use_one_atomic_state_write(self) -> None:
        self.create_workbook("TAS-BATCH-EVENTS")
        self.scan()
        state = self.read_state("TAS-BATCH-EVENTS")
        store = pad.StateStore(self.settings)
        original_atomic_json = pad.atomic_json
        state_writes: list[Path] = []

        def counted_atomic_json(path, value, *args, **kwargs):
            if Path(path).parent == self.settings.states_directory:
                state_writes.append(Path(path))
            return original_atomic_json(path, value, *args, **kwargs)

        events = [
            {"id_caso": f"TAS-EVENT-{number}", "estado": "SIMULADO_OK"}
            for number in range(4)
        ]
        with patch.object(pad, "atomic_json", side_effect=counted_atomic_json):
            store.add_operation_events("TAS-BATCH-EVENTS", state["sha256"], events)
        self.assertEqual(len(state_writes), 1)
        updated = self.read_state("TAS-BATCH-EVENTS")
        simulated = [event for event in updated["events"] if event["event"] == "ALTA_SEGURO_INMUEBLE_SIMULADA"]
        self.assertEqual(len(simulated), 4)

    def test_persistent_storage_error_records_safe_diagnostic_and_preserves_job(self) -> None:
        self.create_workbook("TAS-DIAGNOSTIC")
        self.scan()
        state = self.read_state("TAS-DIAGNOSTIC")
        store = pad.StateStore(self.settings)
        job_path = store.job_path("TAS-DIAGNOSTIC", state["sha256"])
        error = pad.FileOperationError(
            "atomic_json.replace",
            store.state_path("TAS-DIAGNOSTIC", state["sha256"]),
            5,
            PermissionError("synthetic persistent lock"),
        )
        with patch.object(pad.StateStore, "add_operation_events", side_effect=error):
            result = pad.Mock3270Simulator(self.settings).run_job(job_path)
        self.assertEqual(result, ("ERROR_TECNICO", "error"))
        failed = self.read_state("TAS-DIAGNOSTIC")
        self.assertEqual(failed["technical_error"], "PermissionError")
        diagnostic = failed["technical_diagnostic"]
        self.assertEqual(diagnostic["operation"], "atomic_json.replace")
        self.assertEqual(diagnostic["cause_type"], "PermissionError")
        self.assertEqual(diagnostic["retry_attempts"], 5)
        self.assertTrue(diagnostic["affected_path"].startswith("estado\\lotes"))
        self.assertIn("FileOperationError", diagnostic["traceback"])
        self.assertTrue((self.settings.errors_directory / job_path.name).exists())


if __name__ == "__main__":
    unittest.main()
