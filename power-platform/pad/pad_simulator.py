"""Monitor local y simulador seguro para el piloto de tasaciones con PAD.

Este módulo no automatiza ni abre IBM Personal Communications, PCOMM, Excel ni
ninguna sesión bancaria. PAD consume los trabajos JSON creados en ``pendientes``
y ejecuta ``run-job``; este comando deja evidencia local de una operación
simulada contra ``IBM_3270_MOCK``.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import os
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import shutil
import tempfile
import time
import traceback
import uuid
from typing import Any, Callable, Iterator

from openpyxl import load_workbook


FILENAME_PATTERN = re.compile(r"^Resultado_Final_(TAS-[A-Za-z0-9-]+)\.xlsx$", re.IGNORECASE)
TERMINAL_STATES = {
    "SIMULADO_FINALIZADO",
    "SIMULADO_CON_INCIDENCIAS",
    "REVISION_MANUAL",
    "ERROR_ESTRUCTURAL",
    "ERROR_TECNICO",
    "ERROR_ORIGEN_CAMBIO",
    "ENTREGADO",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]", "_", value)


class FileOperationError(RuntimeError):
    """Fallo de almacenamiento con contexto seguro para la bitácora local."""

    def __init__(self, operation: str, path: Path, attempts: int, cause: PermissionError):
        self.operation = operation
        self.path = path
        self.attempts = attempts
        self.cause = cause
        super().__init__(f"{operation} no pudo completar la operación sobre {path.name} tras {attempts} intentos")


def retry_permission_error(
    operation: str,
    path: Path,
    callback: Callable[[], None],
    attempts: int = 5,
    base_delay_seconds: float = 0.1,
) -> None:
    """Reintenta bloqueos transitorios de Windows/OneDrive, con límite estricto."""
    last_error: PermissionError | None = None
    for attempt in range(1, attempts + 1):
        try:
            callback()
            return
        except PermissionError as error:
            last_error = error
            if attempt == attempts:
                break
            time.sleep(base_delay_seconds * (2 ** (attempt - 1)))
    assert last_error is not None
    raise FileOperationError(operation, path, attempts, last_error) from last_error


def atomic_json(
    path: Path,
    value: dict[str, Any],
    attempts: int = 5,
    base_delay_seconds: float = 0.1,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    with tempfile.NamedTemporaryFile(
        mode="w", encoding="utf-8", newline="\n", dir=path.parent, delete=False
    ) as temporary:
        json.dump(value, temporary, ensure_ascii=False, indent=2, sort_keys=True)
        temporary.flush()
        os.fsync(temporary.fileno())
        temporary_path = Path(temporary.name)
    try:
        retry_permission_error(
            "atomic_json.replace", path, lambda: os.replace(temporary_path, path), attempts, base_delay_seconds
        )
    finally:
        if temporary_path is not None and temporary_path.exists():
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                # El diagnóstico principal conserva la operación afectada; nunca se borra evidencia ajena.
                pass


def read_json(path: Path) -> dict[str, Any]:
    # Windows PowerShell can write a UTF-8 BOM while Python tools normally omit
    # it. Both encodings are accepted for local configuration and state files.
    with path.open(encoding="utf-8-sig") as source:
        return json.load(source)


def append_jsonl(
    path: Path,
    events: dict[str, Any] | list[dict[str, Any]],
    attempts: int = 5,
    base_delay_seconds: float = 0.1,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    entries = events if isinstance(events, list) else [events]

    def append() -> None:
        with path.open("a", encoding="utf-8", newline="\n") as sink:
            for event in entries:
                sink.write(json.dumps(event, ensure_ascii=False, sort_keys=True) + "\n")

    retry_permission_error("append_jsonl.write", path, append, attempts, base_delay_seconds)


def move_file(
    source: Path, destination: Path, attempts: int = 5, base_delay_seconds: float = 0.1
) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    retry_permission_error(
        "move_file.replace", destination, lambda: os.replace(source, destination), attempts, base_delay_seconds
    )


def atomic_csv(
    path: Path,
    fieldnames: tuple[str, ...],
    rows: list[dict[str, Any]],
    attempts: int = 5,
    base_delay_seconds: float = 0.1,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    with tempfile.NamedTemporaryFile(
        mode="w", encoding="utf-8-sig", newline="", dir=path.parent, delete=False
    ) as temporary:
        writer = csv.DictWriter(temporary, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
        temporary.flush()
        os.fsync(temporary.fileno())
        temporary_path = Path(temporary.name)
    try:
        retry_permission_error(
            "atomic_csv.replace", path, lambda: os.replace(temporary_path, path), attempts, base_delay_seconds
        )
    finally:
        if temporary_path is not None and temporary_path.exists():
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_contract() -> Any:
    """Carga el contrato versionado en vez de repetir sus encabezados aquí."""
    verifier = Path(__file__).resolve().parents[2] / "tools" / "verify_workbook.py"
    specification = importlib.util.spec_from_file_location("autotasacion_workbook_contract", verifier)
    if specification is None or specification.loader is None:
        raise RuntimeError("No se pudo cargar tools/verify_workbook.py")
    module = importlib.util.module_from_spec(specification)
    specification.loader.exec_module(module)
    return module


@dataclass(frozen=True)
class Settings:
    config_path: Path
    mode: str
    input_directory: Path
    runtime_directory: Path
    poll_seconds: float
    stability_checks: int
    stability_interval_seconds: float
    processing_lease_seconds: int
    suspicious_floors_threshold: int
    suspicious_basements_threshold: int
    file_retry_attempts: int
    file_retry_base_seconds: float
    synthetic_failure_case_ids: frozenset[str]

    @classmethod
    def load(cls, config_path: Path) -> "Settings":
        raw = read_json(config_path)
        if raw.get("mode") != "SIMULACION":
            raise ValueError("Este piloto solo admite mode=SIMULACION")
        root = config_path.parent

        location_profile = raw.get("location_profile")
        if location_profile:
            locations_reference = Path(
                str(raw.get("locations_config_path", "../../config/automation-locations.json"))
            )
            locations_path = (
                locations_reference
                if locations_reference.is_absolute()
                else (root / locations_reference).resolve()
            )
            locations = read_json(locations_path)
            try:
                profile = locations["pad"]["profiles"][str(location_profile)]
                input_value = profile["input_directory"]
                runtime_value = profile["runtime_directory"]
            except (KeyError, TypeError) as error:
                raise ValueError(
                    f"No existe el perfil PAD '{location_profile}' en {locations_path}"
                ) from error
            raw = {**raw, "input_directory": input_value, "runtime_directory": runtime_value}

        def to_path(key: str) -> Path:
            if key not in raw or not str(raw[key]).strip():
                raise ValueError(f"Falta {key} en la configuración PAD")
            candidate = Path(raw[key])
            return candidate if candidate.is_absolute() else (root / candidate).resolve()

        input_directory = to_path("input_directory")
        runtime_directory = to_path("runtime_directory")
        try:
            runtime_directory.relative_to(input_directory)
        except ValueError:
            pass
        else:
            raise ValueError(
                "runtime_directory debe estar fuera de input_directory para no persistir "
                "estado técnico dentro de OneDrive compartido"
            )

        return cls(
            config_path=config_path.resolve(),
            mode=raw["mode"],
            input_directory=input_directory,
            runtime_directory=runtime_directory,
            poll_seconds=float(raw.get("poll_seconds", 5)),
            stability_checks=max(1, int(raw.get("stability_checks", 3))),
            stability_interval_seconds=max(0, float(raw.get("stability_interval_seconds", 2))),
            processing_lease_seconds=max(60, int(raw.get("processing_lease_seconds", 900))),
            suspicious_floors_threshold=max(1, int(raw.get("suspicious_floors_threshold", 100))),
            suspicious_basements_threshold=max(1, int(raw.get("suspicious_basements_threshold", 20))),
            file_retry_attempts=max(1, int(raw.get("file_retry_attempts", 5))),
            file_retry_base_seconds=max(0, float(raw.get("file_retry_base_seconds", 0.1))),
            synthetic_failure_case_ids=frozenset(
                str(item) for item in raw.get("synthetic_failure_case_ids", [])
            ),
        )

    @property
    def states_directory(self) -> Path:
        return self.runtime_directory / "estado" / "lotes"

    @property
    def history_directory(self) -> Path:
        return self.runtime_directory / "estado" / "historial"

    @property
    def locks_directory(self) -> Path:
        return self.runtime_directory / "estado" / "bloqueos"

    @property
    def pending_directory(self) -> Path:
        return self.runtime_directory / "estado" / "pendientes"

    @property
    def completed_jobs_directory(self) -> Path:
        return self.runtime_directory / "estado" / "procesados"

    @property
    def processing_directory(self) -> Path:
        return self.runtime_directory / "procesando"

    @property
    def results_directory(self) -> Path:
        return self.runtime_directory / "resultados"

    @property
    def logs_directory(self) -> Path:
        return self.runtime_directory / "logs"

    @property
    def errors_directory(self) -> Path:
        return self.runtime_directory / "errores"

    def ensure_layout(self) -> None:
        for directory in (
            self.states_directory,
            self.history_directory,
            self.locks_directory,
            self.pending_directory,
            self.completed_jobs_directory,
            self.processing_directory,
            self.results_directory,
            self.logs_directory,
            self.errors_directory,
        ):
            directory.mkdir(parents=True, exist_ok=True)


class StateStore:
    """Estado persistente, bloqueo por lote e idempotencia `(id_lote, SHA-256)`."""

    def __init__(self, settings: Settings):
        self.settings = settings
        settings.ensure_layout()

    def state_path(self, lot_id: str, digest: str) -> Path:
        return self.settings.states_directory / f"{safe_name(lot_id)}_{digest}.json"

    def job_path(self, lot_id: str, digest: str) -> Path:
        return self.settings.pending_directory / f"{safe_name(lot_id)}_{digest}.json"

    def log_path(self, lot_id: str, digest: str) -> Path:
        return self.settings.logs_directory / f"{safe_name(lot_id)}_{digest}.jsonl"

    def result_prefix(self, lot_id: str, digest: str, execution_id: str | None = None) -> str:
        prefix = f"{safe_name(lot_id)}_{digest[:12]}"
        return f"{prefix}_{safe_name(execution_id)[:8]}" if execution_id else prefix

    def _atomic_json(self, path: Path, value: dict[str, Any]) -> None:
        atomic_json(
            path,
            value,
            attempts=self.settings.file_retry_attempts,
            base_delay_seconds=self.settings.file_retry_base_seconds,
        )

    def _append_jsonl(self, path: Path, events: dict[str, Any] | list[dict[str, Any]]) -> None:
        append_jsonl(
            path,
            events,
            attempts=self.settings.file_retry_attempts,
            base_delay_seconds=self.settings.file_retry_base_seconds,
        )

    def diagnostic(self, error: Exception) -> dict[str, Any]:
        """No incluye datos del predio; solo operación, ruta relativa y traceback."""
        if isinstance(error, FileOperationError):
            try:
                affected_path = str(error.path.resolve().relative_to(self.settings.runtime_directory.resolve()))
            except ValueError:
                affected_path = error.path.name
            operation = error.operation
            retry_attempts = error.attempts
            cause_type = type(error.cause).__name__
        else:
            affected_path = None
            operation = "simulacion"
            retry_attempts = 0
            cause_type = type(error).__name__
        return {
            "error_type": type(error).__name__,
            "cause_type": cause_type,
            "operation": operation,
            "affected_path": affected_path,
            "retry_attempts": retry_attempts,
            "traceback": traceback.format_exc(limit=12),
        }

    @staticmethod
    def technical_error_name(error: Exception) -> str:
        return type(error.cause).__name__ if isinstance(error, FileOperationError) else type(error).__name__

    def lot_records(self, lot_id: str) -> list[dict[str, Any]]:
        prefix = f"{safe_name(lot_id)}_"
        records: list[dict[str, Any]] = []
        for path in self.settings.states_directory.glob(f"{prefix}*.json"):
            try:
                records.append(read_json(path))
            except (OSError, json.JSONDecodeError):
                continue
        return records

    @contextmanager
    def lot_lock(self, lot_id: str) -> Iterator[None]:
        """Bloqueo atómico local; los bloqueos expirados permiten recuperación."""
        lock_path = self.settings.locks_directory / f"{safe_name(lot_id)}.lock"
        acquired = False
        for _ in range(3):
            try:
                descriptor = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
                with os.fdopen(descriptor, "w", encoding="utf-8") as lock:
                    json.dump({"pid": os.getpid(), "created_at": utc_now()}, lock)
                acquired = True
                break
            except FileExistsError:
                try:
                    age_seconds = time.time() - lock_path.stat().st_mtime
                    if age_seconds > self.settings.processing_lease_seconds:
                        lock_path.unlink(missing_ok=True)
                        continue
                except OSError:
                    continue
                break
        if not acquired:
            raise RuntimeError(f"El lote {lot_id} está reservado por otra ejecución local")
        try:
            yield
        finally:
            try:
                lock_path.unlink(missing_ok=True)
            except OSError:
                # Un bloqueo temporal del sincronizador no invalida una operación ya persistida.
                # La detección de lease expirado permite recuperar este lock en una ejecución posterior.
                pass

    def event(self, record: dict[str, Any], name: str, **detail: Any) -> None:
        event = {"at": utc_now(), "event": name, **detail}
        record.setdefault("events", []).append(event)
        self._append_jsonl(self.log_path(record["id_lote"], record["sha256"]), {
            "at": event["at"],
            "event": name,
            "id_lote": record["id_lote"],
            "sha256": record["sha256"],
            **detail,
        })

    def save(self, record: dict[str, Any]) -> None:
        record["updated_at"] = utc_now()
        self._atomic_json(self.state_path(record["id_lote"], record["sha256"]), record)

    def register(
        self, lot_id: str, digest: str, source: Path, working_copy: Path
    ) -> tuple[dict[str, Any], str]:
        """Registra una entrada nueva o informa duplicado/cambio de origen."""
        with self.lot_lock(lot_id):
            same_hash: dict[str, Any] | None = None
            other_hash = False
            for record in self.lot_records(lot_id):
                if record.get("sha256") == digest:
                    same_hash = record
                else:
                    other_hash = True
            if same_hash is not None:
                self.event(same_hash, "EVENTO_DUPLICADO_IGNORADO")
                self.save(same_hash)
                return same_hash, "duplicate"
            if other_hash:
                changed = {
                    "schema_version": 1,
                    "execution_id": str(uuid.uuid4()),
                    "id_lote": lot_id,
                    "sha256": digest,
                    "source_name": source.name,
                    "source_path": str(source),
                    "working_copy": str(working_copy),
                    "status": "ERROR_ORIGEN_CAMBIO",
                    "created_at": utc_now(),
                    "updated_at": utc_now(),
                    "events": [],
                    "validation": {},
                }
                self.event(changed, "ORIGEN_CAMBIO_DETECTADO")
                self.save(changed)
                return changed, "origin_changed"
            record = {
                "schema_version": 1,
                "execution_id": str(uuid.uuid4()),
                "id_lote": lot_id,
                "sha256": digest,
                "source_name": source.name,
                "source_path": str(source),
                "working_copy": str(working_copy),
                "status": "DETECTADO",
                "created_at": utc_now(),
                "updated_at": utc_now(),
                "events": [],
                "validation": {},
                "attempts": 0,
            }
            self.event(record, "LOTE_DETECTADO")
            self.save(record)
            return record, "new"

    def update(self, record: dict[str, Any], status: str | None = None, **detail: Any) -> dict[str, Any]:
        with self.lot_lock(record["id_lote"]):
            current = read_json(self.state_path(record["id_lote"], record["sha256"]))
            if status is not None and current["status"] != status:
                self.event(current, "ESTADO_ACTUALIZADO", anterior=current["status"], nuevo=status)
                current["status"] = status
            current.update(detail)
            self.save(current)
            return current

    def add_operation_events(
        self, lot_id: str, digest: str, operation_events: list[dict[str, str]]
    ) -> dict[str, Any]:
        """Guarda el lote de eventos con una sola escritura atómica del estado."""
        with self.lot_lock(lot_id):
            current = read_json(self.state_path(lot_id, digest))
            timestamp = utc_now()
            state_events = [
                {"at": timestamp, "event": "ALTA_SEGURO_INMUEBLE_SIMULADA", **event}
                for event in operation_events
            ]
            log_events = [
                {
                    "at": event["at"],
                    "event": event["event"],
                    "id_lote": lot_id,
                    "sha256": digest,
                    "id_caso": event["id_caso"],
                    "estado": event["estado"],
                }
                for event in state_events
            ]
            self._append_jsonl(self.log_path(lot_id, digest), log_events)
            current.setdefault("events", []).extend(state_events)
            self.save(current)
            return current

    def claim_simulation(self, lot_id: str, digest: str) -> tuple[dict[str, Any], str]:
        """Reserva el lote de forma atómica antes de una simulación PAD."""
        with self.lot_lock(lot_id):
            record = read_json(self.state_path(lot_id, digest))
            status = record["status"]
            if status in TERMINAL_STATES:
                self.event(record, "EJECUCION_IDEMPOTENTE_IGNORADA", estado=status)
                self.save(record)
                return record, "terminal"
            if status == "VALIDADO":
                record["status"] = "SIMULANDO_3270"
                record["attempts"] = int(record.get("attempts", 0)) + 1
                record["lease_until"] = (
                    datetime.now(timezone.utc) + timedelta(seconds=self.settings.processing_lease_seconds)
                ).replace(microsecond=0).isoformat()
                self.event(record, "SIMULACION_RECLAMADA", intento=record["attempts"])
                self.save(record)
                return record, "claimed"
            if status == "SIMULANDO_3270":
                updated = datetime.fromisoformat(record["updated_at"])
                expired = datetime.now(timezone.utc) - updated > timedelta(
                    seconds=self.settings.processing_lease_seconds
                )
                if expired:
                    record["attempts"] = int(record.get("attempts", 0)) + 1
                    record["lease_until"] = (
                        datetime.now(timezone.utc) + timedelta(seconds=self.settings.processing_lease_seconds)
                    ).replace(microsecond=0).isoformat()
                    self.event(record, "SIMULACION_RECUPERADA", intento=record["attempts"])
                    self.save(record)
                    return record, "recovered"
                return record, "busy"
            return record, "not_ready"

    def reprocess(self, lot_id: str, digest: str, reason: str) -> dict[str, Any]:
        """Acción explícita y auditada; nunca la invoca el monitor automáticamente."""
        with self.lot_lock(lot_id):
            current_path = self.state_path(lot_id, digest)
            record = read_json(current_path)
            if record["status"] not in TERMINAL_STATES:
                raise RuntimeError("Solo se puede reprocesar un lote terminal")
            archive = self.settings.history_directory / (
                f"{current_path.stem}_{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}.json"
            )
            shutil.copy2(current_path, archive)
            fresh = dict(record)
            fresh["execution_id"] = str(uuid.uuid4())
            fresh["status"] = "VALIDADO"
            fresh["attempts"] = 0
            fresh["reprocess_of"] = record["execution_id"]
            fresh["reprocess_reason"] = reason
            fresh["events"] = []
            fresh["reprocess_at"] = utc_now()
            for obsolete_key in (
                "technical_error",
                "technical_diagnostic",
                "simulation",
                "lease_until",
                "postprocessing_warning",
                "report_warning",
                "error_job_warning",
            ):
                fresh.pop(obsolete_key, None)
            self.event(fresh, "REPROCESO_EXPLICITO", motivo=reason)
            self.save(fresh)
            job = {
                "schema_version": 1,
                "id_lote": lot_id,
                "sha256": digest,
                "execution_id": fresh["execution_id"],
                "working_copy": fresh["working_copy"],
                "case_ids": fresh.get("validation", {}).get("valid_case_ids", []),
                "manual_count": fresh.get("validation", {}).get("manual_count", 0),
                "revision_ia_count": fresh.get("validation", {}).get("revision_ia_count", 0),
                "created_at": utc_now(),
            }
            self._atomic_json(self.job_path(lot_id, digest), job)
            return fresh


class WorkbookValidator:
    def __init__(self, settings: Settings):
        self.settings = settings
        self.contract = load_contract()
        self.headers = tuple(self.contract.PARA_PROCESAR_HEADERS)
        self.required = tuple(self.contract.REQUIRED_READY_FIELDS)
        self.header_index = {header: index for index, header in enumerate(self.headers)}
        self.year_header = self.headers[-2]

    def validate(self, path: Path) -> dict[str, Any]:
        try:
            workbook = load_workbook(path, data_only=True)
        except Exception as error:
            return {"structural_issues": [f"XLSX_NO_LEIBLE:{type(error).__name__}"], "rows": []}
        try:
            structural = self._structural_issues(workbook)
            if structural:
                return {"structural_issues": structural, "rows": []}
            ready = workbook["PARA_PROCESAR"]
            rows = [self._classify_row(number, values) for number, values in enumerate(
                ready.iter_rows(min_row=2, values_only=True), start=2
            ) if not all(blank(value) for value in values)]
            revision = workbook["REVISION_IA"]
            revision_count = sum(
                1 for values in revision.iter_rows(min_row=2, values_only=True)
                if not all(blank(value) for value in values)
            )
            return {"structural_issues": [], "rows": rows, "revision_ia_count": revision_count}
        finally:
            workbook.close()

    def _structural_issues(self, workbook: Any) -> list[str]:
        issues: list[str] = []
        if tuple(workbook.sheetnames) != tuple(self.contract.SHEETS):
            issues.append("HOJAS_NO_COINCIDEN_CON_CONTRATO")
        for sheet, table in self.contract.TABLES.items():
            if sheet not in workbook.sheetnames:
                continue
            if table not in workbook[sheet].tables:
                issues.append(f"FALTA_TABLA:{table}")
        if "PARA_PROCESAR" not in workbook.sheetnames:
            return issues
        headers = tuple(cell.value for cell in workbook["PARA_PROCESAR"][1])
        if headers != self.headers:
            issues.append("ENCABEZADOS_PARA_PROCESAR_NO_COINCIDEN")
        return issues

    def _classify_row(self, row_number: int, values: tuple[Any, ...]) -> dict[str, Any]:
        issues: list[str] = []
        warnings: list[str] = []
        lookup = {header: values[index] if index < len(values) else None for header, index in self.header_index.items()}
        for header in self.required:
            if blank(lookup[header]):
                issues.append(f"CAMPO_REQUERIDO_AUSENTE:{header}")
        case_id = "" if blank(lookup["ID_CASO"]) else str(lookup["ID_CASO"]).strip()
        if case_id and not case_id.startswith("TAS-"):
            issues.append("ID_CASO_INVALIDO")
        amount = lookup["IMPORTE"]
        if not blank(amount):
            normalized_amount = str(amount).strip()
            if not self.contract.IMPORTE_FORMAT.fullmatch(normalized_amount):
                issues.append("IMPORTE_FORMATO_INVALIDO")
            else:
                try:
                    if Decimal(normalized_amount.replace(",", "")) <= Decimal("0"):
                        issues.append("IMPORTE_NO_POSITIVO")
                except InvalidOperation:
                    issues.append("IMPORTE_FORMATO_INVALIDO")
        currency = "" if blank(lookup["MONEDA"]) else str(lookup["MONEDA"]).strip().upper()
        if currency and currency not in {"PEN", "USD"}:
            issues.append("MONEDA_NO_ADMITIDA")
        self._integer_check(lookup["PISOS"], "PISOS", issues, warnings, self.settings.suspicious_floors_threshold)
        self._integer_check(lookup["SOTANOS"], "SOTANOS", issues, warnings, self.settings.suspicious_basements_threshold)
        self._year_check(lookup[self.year_header], issues)
        classification = "REGISTRO_VALIDO" if not issues and not warnings else "REVISION_MANUAL"
        return {
            "row_number": row_number,
            "id_caso": case_id or None,
            "classification": classification,
            "issues": issues,
            "warnings": warnings,
        }

    @staticmethod
    def _as_integer(value: Any) -> int | None:
        if isinstance(value, bool) or blank(value):
            return None
        if isinstance(value, int):
            return value
        if isinstance(value, float) and value.is_integer():
            return int(value)
        if isinstance(value, str) and re.fullmatch(r"\d+", value.strip()):
            return int(value.strip())
        return None

    def _integer_check(
        self, value: Any, field: str, issues: list[str], warnings: list[str], warning_threshold: int
    ) -> None:
        if blank(value):
            return
        number = self._as_integer(value)
        if number is None or number < 0:
            issues.append(f"{field}_INVALIDO")
        elif number > warning_threshold:
            warnings.append(f"{field}_SOSPECHOSO")

    def _year_check(self, value: Any, issues: list[str]) -> None:
        if blank(value):
            return
        year = self._as_integer(value)
        if year is None or year < 1900 or year > datetime.now().year:
            issues.append("ANIO_INVALIDO")


class LocalMonitor:
    def __init__(self, settings: Settings):
        self.settings = settings
        self.store = StateStore(settings)
        self.validator = WorkbookValidator(settings)

    def wait_until_stable(self, path: Path) -> bool:
        previous: tuple[int, int] | None = None
        matches = 0
        for index in range(self.settings.stability_checks):
            try:
                stat = path.stat()
            except OSError:
                return False
            sample = (stat.st_size, stat.st_mtime_ns)
            matches = matches + 1 if sample == previous else 1
            previous = sample
            if index < self.settings.stability_checks - 1:
                time.sleep(self.settings.stability_interval_seconds)
        return matches == self.settings.stability_checks and previous is not None and previous[0] > 0

    def scan_once(self) -> list[dict[str, str]]:
        self.settings.input_directory.mkdir(parents=True, exist_ok=True)
        processed: list[dict[str, str]] = []
        for candidate in sorted(self.settings.input_directory.iterdir()):
            if not candidate.is_file() or candidate.name.startswith("~$"):
                continue
            match = FILENAME_PATTERN.fullmatch(candidate.name)
            if match is None:
                continue
            if not self.wait_until_stable(candidate):
                processed.append({"file": candidate.name, "outcome": "AUN_COPIANDO"})
                continue
            processed.append({"file": candidate.name, "outcome": self._register_candidate(candidate, match.group(1))})
        return processed

    def _register_candidate(self, source: Path, lot_id: str) -> str:
        digest = sha256_file(source)
        destination = self.settings.processing_directory / safe_name(lot_id) / f"{digest}.xlsx"
        record, outcome = self.store.register(lot_id, digest, source, destination)
        if outcome != "new":
            return outcome
        try:
            destination.parent.mkdir(parents=True, exist_ok=True)
            temporary_copy = destination.with_suffix(".copying")
            shutil.copy2(source, temporary_copy)
            move_file(
                temporary_copy,
                destination,
                attempts=self.settings.file_retry_attempts,
                base_delay_seconds=self.settings.file_retry_base_seconds,
            )
            record = self.store.update(record, "VALIDANDO")
            validation = self.validator.validate(destination)
            if validation["structural_issues"]:
                record = self.store.update(record, "ERROR_ESTRUCTURAL", validation=validation)
                self._write_validation_report(record)
                return "ERROR_ESTRUCTURAL"
            rows = validation["rows"]
            valid_case_ids = [str(row["id_caso"]) for row in rows if row["classification"] == "REGISTRO_VALIDO"]
            manual_count = sum(1 for row in rows if row["classification"] == "REVISION_MANUAL")
            validation["valid_case_ids"] = valid_case_ids
            validation["manual_count"] = manual_count
            record = self.store.update(record, validation=validation)
            self._write_validation_report(record)
            if not valid_case_ids:
                self.store.update(record, "REVISION_MANUAL")
                return "REVISION_MANUAL"
            job = {
                "schema_version": 1,
                "id_lote": lot_id,
                "sha256": digest,
                "execution_id": record["execution_id"],
                "working_copy": str(destination),
                "case_ids": valid_case_ids,
                "manual_count": manual_count,
                "revision_ia_count": validation["revision_ia_count"],
                "created_at": utc_now(),
            }
            self.store._atomic_json(self.store.job_path(lot_id, digest), job)
            self.store.update(record, "VALIDADO")
            return "VALIDADO"
        except Exception as error:
            self.store.update(
                record,
                "ERROR_TECNICO",
                technical_error=self.store.technical_error_name(error),
                technical_diagnostic=self.store.diagnostic(error),
            )
            return "ERROR_TECNICO"

    def _write_validation_report(self, record: dict[str, Any]) -> None:
        payload = {
            "id_lote": record["id_lote"],
            "sha256": record["sha256"],
            "status": record["status"],
            "validation": record.get("validation", {}),
            "created_at": utc_now(),
        }
        self.store._atomic_json(
            self.settings.results_directory / f"{self.store.result_prefix(record['id_lote'], record['sha256'], record['execution_id'])}_validacion.json",
            payload,
        )


class Mock3270Simulator:
    """Ejecuta exclusivamente la evidencia local de la simulación."""

    def __init__(self, settings: Settings):
        self.settings = settings
        self.store = StateStore(settings)

    def run_job(self, job_path: Path) -> tuple[str, str]:
        if job_path.parent.resolve() != self.settings.pending_directory.resolve():
            raise ValueError("El trabajo debe provenir de estado/pendientes")
        job = read_json(job_path)
        lot_id, digest = str(job["id_lote"]), str(job["sha256"])
        record, outcome = self.store.claim_simulation(lot_id, digest)
        if outcome != "claimed" and outcome != "recovered":
            return record["status"], outcome
        rows: list[dict[str, Any]] = []
        operation_events: list[dict[str, str]] = []
        started = time.perf_counter()
        finalized = False
        try:
            for case_id in job["case_ids"]:
                time.sleep(0.001)
                status = "SIMULADO_OK"
                code = "ALTA_SEGURO_INMUEBLE_SIMULADA"
                if str(case_id) in self.settings.synthetic_failure_case_ids:
                    status = "SIMULADO_ERROR"
                    code = "FALLO_SIMULADO_CONFIGURADO"
                elapsed_ms = round((time.perf_counter() - started) * 1000, 3)
                rows.append({
                    "ID_LOTE": lot_id,
                    "ID_CASO": str(case_id),
                    "OPERACION": "ALTA_SEGURO_INMUEBLE_SIMULADA",
                    "DESTINO": "IBM_3270_MOCK",
                    "ESTADO": status,
                    "SIMULACION": "SI",
                    "TIEMPO_MS": elapsed_ms,
                    "CODIGO": code,
                })
                operation_events.append({"id_caso": str(case_id), "estado": status})
            record = self.store.add_operation_events(lot_id, digest, operation_events)
            failures = sum(1 for row in rows if row["ESTADO"] != "SIMULADO_OK")
            incidents = failures + int(job.get("manual_count", 0)) + int(job.get("revision_ia_count", 0))
            final_status = "SIMULADO_CON_INCIDENCIAS" if incidents else "SIMULADO_FINALIZADO"
            record = self.store.update(
                record,
                final_status,
                simulation={
                    "processed_count": len(rows),
                    "failed_count": failures,
                    "manual_count": int(job.get("manual_count", 0)),
                    "revision_ia_count": int(job.get("revision_ia_count", 0)),
                    "target": "IBM_3270_MOCK",
                    "mode": "SIMULACION",
                },
            )
            finalized = True
            self._write_reports(record, rows)
            destination = self.settings.completed_jobs_directory / job_path.name
            move_file(
                job_path,
                destination,
                attempts=self.settings.file_retry_attempts,
                base_delay_seconds=self.settings.file_retry_base_seconds,
            )
            return record["status"], outcome
        except Exception as error:
            diagnostic = self.store.diagnostic(error)
            if finalized:
                # El resultado ya es terminal. Un bloqueo al guardar/mover evidencia no relanza el lote.
                try:
                    self.store.update(record, postprocessing_warning=diagnostic)
                except Exception:
                    pass
                return record["status"], "cleanup_pending"
            try:
                record = self.store.update(
                    record,
                    "ERROR_TECNICO",
                    technical_error=self.store.technical_error_name(error),
                    technical_diagnostic=diagnostic,
                )
            except Exception:
                return "ERROR_TECNICO", "diagnostic_persistence_failed"
            try:
                self._write_reports(record, rows)
            except Exception as report_error:
                try:
                    self.store.update(
                        record,
                        report_warning=self.store.diagnostic(report_error),
                    )
                except Exception:
                    pass
            error_job = self.settings.errors_directory / job_path.name
            if job_path.exists():
                try:
                    move_file(
                        job_path,
                        error_job,
                        attempts=self.settings.file_retry_attempts,
                        base_delay_seconds=self.settings.file_retry_base_seconds,
                    )
                except Exception as move_error:
                    try:
                        self.store.update(record, error_job_warning=self.store.diagnostic(move_error))
                    except Exception:
                        pass
            return record["status"], "error"

    def _write_reports(self, record: dict[str, Any], rows: list[dict[str, Any]]) -> None:
        prefix = self.store.result_prefix(record["id_lote"], record["sha256"], record["execution_id"])
        csv_path = self.settings.results_directory / f"{prefix}_operaciones.csv"
        atomic_csv(
            csv_path,
            (
                "ID_LOTE", "ID_CASO", "OPERACION", "DESTINO", "ESTADO", "SIMULACION", "TIEMPO_MS", "CODIGO"
            ),
            rows,
            attempts=self.settings.file_retry_attempts,
            base_delay_seconds=self.settings.file_retry_base_seconds,
        )
        self.store._atomic_json(self.settings.results_directory / f"{prefix}_resumen.json", {
            "id_lote": record["id_lote"],
            "sha256": record["sha256"],
            "status": record["status"],
            "simulation": record.get("simulation", {}),
            "created_at": utc_now(),
        })


def parse_config(value: str) -> Settings:
    return Settings.load(Path(value))


def command_monitor(settings: Settings, once: bool = False) -> int:
    monitor = LocalMonitor(settings)
    while True:
        outcomes = monitor.scan_once()
        for outcome in outcomes:
            print(json.dumps(outcome, ensure_ascii=False))
        if once:
            return 0
        time.sleep(settings.poll_seconds)


def command_run_job(settings: Settings, job_path: Path) -> int:
    status, outcome = Mock3270Simulator(settings).run_job(job_path)
    print(json.dumps({"status": status, "outcome": outcome}, ensure_ascii=False))
    return 0 if status.startswith("SIMULADO") or outcome in {"terminal", "busy"} else 1


def run_cycle(settings: Settings) -> dict[str, Any]:
    """Ejecuta una vuelta local: detecta XLSX y consume solo trabajos pendientes.

    PAD invoca este punto de entrada una vez. No mantiene un bucle propio, no
    abre Excel y no interactúa con 3270. La reserva atómica del simulador sigue
    siendo la única autoridad para impedir una doble ejecución de un lote.
    """
    detection_results = LocalMonitor(settings).scan_once()
    simulator = Mock3270Simulator(settings)
    job_results: list[dict[str, str]] = []
    errors = 0

    for job_path in sorted(settings.pending_directory.glob("*.json")):
        try:
            status, outcome = simulator.run_job(job_path)
        except Exception:
            # El detalle técnico queda en el estado del lote cuando el trabajo
            # pudo reclamarse. La salida de consola no incluye datos del Excel.
            status, outcome = "ERROR_TECNICO", "unhandled_job_error"
        job_results.append({"status": status, "outcome": outcome})
        if status == "ERROR_TECNICO" or outcome in {"error", "unhandled_job_error", "diagnostic_persistence_failed"}:
            errors += 1

    detected_outcomes: dict[str, int] = {}
    for result in detection_results:
        outcome = result["outcome"]
        detected_outcomes[outcome] = detected_outcomes.get(outcome, 0) + 1
    job_outcomes: dict[str, int] = {}
    for result in job_results:
        key = f"{result['status']}:{result['outcome']}"
        job_outcomes[key] = job_outcomes.get(key, 0) + 1

    return {
        "mode": "SIMULACION",
        "detected_files": len(detection_results),
        "detection_outcomes": detected_outcomes,
        "pending_jobs_seen": len(job_results),
        "job_outcomes": job_outcomes,
        "technical_errors": errors,
    }


def command_run_cycle(settings: Settings) -> int:
    summary = run_cycle(settings)
    print(json.dumps(summary, ensure_ascii=False, sort_keys=True))
    return 0 if summary["technical_errors"] == 0 else 1


def command_reprocess(settings: Settings, lot_id: str, digest: str, reason: str) -> int:
    record = StateStore(settings).reprocess(lot_id, digest, reason)
    print(json.dumps({"status": record["status"], "execution_id": record["execution_id"]}, ensure_ascii=False))
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Piloto local PAD/IBM 3270 en modo simulación")
    parser.add_argument("--config", required=True, help="Ruta del JSON de configuración local")
    subcommands = parser.add_subparsers(dest="command", required=True)
    monitor_parser = subcommands.add_parser("monitor")
    monitor_parser.add_argument("--once", action="store_true")
    job_parser = subcommands.add_parser("run-job")
    job_parser.add_argument("--job", required=True, type=Path)
    subcommands.add_parser(
        "run-cycle",
        help="Detecta XLSX y ejecuta una vez todos los trabajos locales pendientes en modo simulación.",
    )
    reprocess_parser = subcommands.add_parser("reprocess")
    reprocess_parser.add_argument("--lot", required=True)
    reprocess_parser.add_argument("--sha256", required=True)
    reprocess_parser.add_argument("--reason", required=True)
    arguments = parser.parse_args()
    settings = parse_config(arguments.config)
    if arguments.command == "monitor":
        return command_monitor(settings, arguments.once)
    if arguments.command == "run-job":
        return command_run_job(settings, arguments.job)
    if arguments.command == "run-cycle":
        return command_run_cycle(settings)
    return command_reprocess(settings, arguments.lot, arguments.sha256, arguments.reason)


if __name__ == "__main__":
    raise SystemExit(main())
