"""Prepara y ejecuta el último tramo local: Resultado_Final -> MASIVO -> 3270.

Este módulo no modifica el XLSX de Cloud ni la plantilla maestra. Copia la
plantilla XLSM al runtime, carga solo PARA_PROCESAR y ejecuta el procedimiento
VBA versionado que opera la sesión PCOMM ya abierta.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

SOURCE_HEADERS = (
    "PRESTAMO", "TIPO DE INMUEBLE", "VALOR DEL BIEN", "MONEDA", "IMPORTE", "COL_F",
    "DIRECCION", "DIRECCION1", "EXTERIOR", "INTERIOR", "REFERENCIA", "COL_L",
    "UBICACION", "UBICACION1", "MUNICIPIO", "DIST_COD", "COL_Q", "PROV_COD", "COL_S",
    "DEPT_COD", "CLASE", "PISOS", "SOTANOS", "AÑO", "ID_CASO",
)
REQUIRED_SOURCE_FIELDS = (
    "PRESTAMO", "TIPO DE INMUEBLE", "MONEDA", "IMPORTE", "DIRECCION", "DIRECCION1",
    "EXTERIOR", "MUNICIPIO", "DIST_COD", "PROV_COD", "DEPT_COD", "PISOS", "SOTANOS", "AÑO",
)
READ_ONLY_MACRO_PROCEDURE = "AutoTasacion3270.ValidarSesionSoloLectura"
NAVIGATION_ONLY_MACRO_PROCEDURE = "AutoTasacion3270.NavegarSesionSinCambios"


class Macro3270Error(RuntimeError):
    """Error visible para PAD sin exponer contenido de la tasación."""


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]", "_", value)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def atomic_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def as_code(value: Any) -> str:
    text = clean(value)
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def parse_amount(value: Any) -> float:
    raw = clean(value).replace(",", "")
    try:
        amount = Decimal(raw)
    except (InvalidOperation, ValueError) as error:
        raise Macro3270Error("IMPORTE_INVALIDO_EN_PARA_PROCESAR") from error
    if amount <= 0:
        raise Macro3270Error("IMPORTE_NO_POSITIVO_EN_PARA_PROCESAR")
    return float(amount)


@dataclass(frozen=True)
class Settings:
    config_path: Path
    input_directory: Path
    runtime_directory: Path
    template_path: Path
    macro_procedure: str
    session_name: str
    excel_visible: bool
    stability_checks: int
    stability_interval_seconds: float
    output_directory: Path | None

    @property
    def work_directory(self) -> Path:
        return self.runtime_directory / "trabajo"

    @property
    def state_directory(self) -> Path:
        return self.runtime_directory / "estado" / "macro_3270"

    @property
    def results_directory(self) -> Path:
        return self.runtime_directory / "resultados"


def resolve_path(config_path: Path, raw: str) -> Path:
    path = Path(raw).expanduser()
    return path if path.is_absolute() else (config_path.parent / path).resolve()


def load_settings(config_path: Path) -> Settings:
    payload = json.loads(config_path.read_text(encoding="utf-8-sig"))
    if payload.get("mode") != "MACRO_3270":
        raise Macro3270Error("El modo de configuración debe ser MACRO_3270.")

    if "input_directory" not in payload or "runtime_directory" not in payload:
        locations_path = resolve_path(config_path, payload.get("locations_config_path", "../../config/automation-locations.json"))
        locations = json.loads(locations_path.read_text(encoding="utf-8-sig"))
        profile = payload.get("location_profile", "production")
        try:
            profile_values = locations["pad"]["profiles"][profile]
        except KeyError as error:
            raise Macro3270Error("No se pudo resolver el perfil de ubicaciones PAD.") from error
        payload = {**profile_values, **payload}

    required = ("input_directory", "runtime_directory", "macro_template_path")
    missing = [key for key in required if not clean(payload.get(key))]
    if missing:
        raise Macro3270Error("Faltan campos de configuración: " + ", ".join(missing))

    settings = Settings(
        config_path=config_path,
        input_directory=resolve_path(config_path, payload["input_directory"]),
        runtime_directory=resolve_path(config_path, payload["runtime_directory"]),
        template_path=resolve_path(config_path, payload["macro_template_path"]),
        macro_procedure=payload.get("macro_procedure", "AutoTasacion3270.EjecutarLote3270"),
        session_name=payload.get("session_name", "A"),
        excel_visible=bool(payload.get("excel_visible", True)),
        stability_checks=int(payload.get("stability_checks", 3)),
        stability_interval_seconds=float(payload.get("stability_interval_seconds", 2)),
        output_directory=(resolve_path(config_path, payload["output_directory"]) if clean(payload.get("output_directory")) else None),
    )
    if not settings.input_directory.exists():
        raise Macro3270Error("No existe input_directory configurado para PAD.")
    if not settings.template_path.is_file() or settings.template_path.suffix.lower() != ".xlsm":
        raise Macro3270Error("No existe una plantilla XLSM aprobada en macro_template_path.")
    return settings


def wait_stable(path: Path, checks: int, interval: float) -> None:
    observed: tuple[int, int] | None = None
    consecutive = 0
    while consecutive < checks:
        if not path.is_file():
            raise Macro3270Error("RESULTADO_FINAL_NO_ENCONTRADO")
        current = (path.stat().st_size, path.stat().st_mtime_ns)
        if current == observed:
            consecutive += 1
        else:
            observed = current
            consecutive = 1
        if consecutive < checks:
            time.sleep(interval)


def result_workbook(settings: Settings, lot_id: str) -> Path:
    candidate = settings.input_directory / f"Resultado_Final_{lot_id}.xlsx"
    if candidate.is_file():
        return candidate
    raise Macro3270Error("RESULTADO_FINAL_NO_ENCONTRADO")


def read_ready_rows(source: Path) -> list[dict[str, Any]]:
    # Se conserva el libro cerrado enseguida. ``read_only=False`` permite leer
    # de forma fiable las tablas de Excel que componen el contrato Cloud.
    workbook = load_workbook(source, read_only=False, data_only=True)
    try:
        if "PARA_PROCESAR" not in workbook.sheetnames:
            raise Macro3270Error("FALTA_HOJA_PARA_PROCESAR")
        sheet = workbook["PARA_PROCESAR"]
        if "tblParaProcesar" not in sheet.tables:
            raise Macro3270Error("FALTA_TABLA_PARA_PROCESAR")
        table = sheet.tables["tblParaProcesar"]
        minimum_column, minimum_row, maximum_column, maximum_row = range_boundaries(table.ref)
        headers = tuple(sheet.cell(minimum_row, col).value for col in range(minimum_column, maximum_column + 1))
        if headers != SOURCE_HEADERS:
            raise Macro3270Error("ENCABEZADOS_PARA_PROCESAR_NO_COINCIDEN")
        rows: list[dict[str, Any]] = []
        for row_number in range(minimum_row + 1, maximum_row + 1):
            values = [sheet.cell(row_number, col).value for col in range(minimum_column, maximum_column + 1)]
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            row = dict(zip(SOURCE_HEADERS, values, strict=True))
            validate_ready_row(row)
            rows.append(row)
        if not rows:
            raise Macro3270Error("SIN_FILAS_PARA_PROCESAR")
        return rows
    finally:
        workbook.close()


def validate_ready_row(row: dict[str, Any]) -> None:
    loan = as_code(row["PRESTAMO"])
    if not re.fullmatch(r"\d{20}", loan):
        raise Macro3270Error("PRESTAMO_INVALIDO_EN_PARA_PROCESAR")
    missing = [field for field in REQUIRED_SOURCE_FIELDS if clean(row[field]) == ""]
    if missing:
        raise Macro3270Error("CAMPOS_REQUERIDOS_FALTANTES_EN_PARA_PROCESAR:" + ",".join(missing))
    if clean(row["MONEDA"]).upper() not in {"PEN", "USD"}:
        raise Macro3270Error("MONEDA_INVALIDA_EN_PARA_PROCESAR")
    parse_amount(row["IMPORTE"])
    for field, width in (("DIST_COD", 3), ("PROV_COD", 2), ("DEPT_COD", 2)):
        code = as_code(row[field])
        if not code.isdigit() or len(code) != width:
            raise Macro3270Error("CODIGO_TERRITORIAL_INVALIDO_EN_PARA_PROCESAR")


def com_client() -> Any:
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as error:
        raise Macro3270Error("FALTA_PYWIN32_EN_ENTORNO_PAD") from error
    return win32com.client


def clear_non_formula_cells(sheet: Any, end_row: int) -> None:
    for row in range(2, end_row + 1):
        for column in range(1, 26):
            cell = sheet.Cells(row, column)
            if not cell.HasFormula:
                cell.ClearContents()


def set_value(sheet: Any, row: int, column: int, value: Any, number_format: str | None = None) -> None:
    cell = sheet.Cells(row, column)
    if number_format:
        cell.NumberFormat = number_format
    cell.Value = value


def populate_workbook(
    workbook_path: Path,
    lot_id: str,
    source_sha: str,
    rows: list[dict[str, Any]],
    settings: Settings,
    execution_mode: str,
) -> list[dict[str, Any]]:
    client = com_client()
    excel = None
    workbook = None
    phase = "crear_instancia_excel"
    try:
        excel = client.DispatchEx("Excel.Application")
        phase = "configurar_excel"
        excel.Visible = settings.excel_visible
        excel.DisplayAlerts = False
        phase = "abrir_copia_plantilla"
        workbook = excel.Workbooks.Open(str(workbook_path), UpdateLinks=0, ReadOnly=False, IgnoreReadOnlyRecommended=True)
        phase = "obtener_hoja_y_tabla_masivo"
        sheet = workbook.Worksheets("MASIVO")
        table = sheet.ListObjects("Tabla164")
        last_row = len(rows) + 1
        phase = "redimensionar_tabla_masivo"
        table.Resize(sheet.Range(sheet.Cells(1, 1), sheet.Cells(last_row, 25)))
        phase = "limpiar_celdas_de_entrada"
        clear_non_formula_cells(sheet, last_row)

        phase = "crear_traza_pad"
        trace = configure_trace_sheet(workbook, lot_id, source_sha, rows)
        phase = "cargar_filas_para_procesar"
        for index, source in enumerate(rows, start=2):
            set_value(sheet, index, 1, as_code(source["PRESTAMO"]), "@")
            set_value(sheet, index, 2, clean(source["TIPO DE INMUEBLE"]).upper(), "@")
            set_value(sheet, index, 4, clean(source["MONEDA"]).upper(), "@")
            set_value(sheet, index, 5, parse_amount(source["IMPORTE"]), "#,##0.00")
            set_value(sheet, index, 6, clean(source["DIRECCION"]).upper(), "@")
            set_value(sheet, index, 8, clean(source["DIRECCION1"]).upper(), "@")
            set_value(sheet, index, 9, clean(source["EXTERIOR"]).upper(), "@")
            set_value(sheet, index, 10, clean(source["INTERIOR"]).upper(), "@")
            set_value(sheet, index, 11, clean(source["REFERENCIA"]).upper(), "@")
            set_value(sheet, index, 12, clean(source["UBICACION"]).upper(), "@")
            set_value(sheet, index, 14, clean(source["UBICACION1"]).upper(), "@")
            set_value(sheet, index, 15, clean(source["MUNICIPIO"]).upper(), "@")
            set_value(sheet, index, 16, as_code(source["DIST_COD"]), "@")
            set_value(sheet, index, 18, as_code(source["PROV_COD"]), "@")
            set_value(sheet, index, 20, as_code(source["DEPT_COD"]), "@")
            set_value(sheet, index, 22, int(float(source["PISOS"])))
            set_value(sheet, index, 23, int(float(source["SOTANOS"])))
            set_value(sheet, index, 24, int(float(source["AÑO"])))
            set_value(sheet, index, 25, None)
            trace.Cells(index, 1).Value = index
            trace.Cells(index, 2).Value = lot_id
            trace.Cells(index, 3).Value = clean(source["ID_CASO"])
            trace.Cells(index, 4).NumberFormat = "@"
            trace.Cells(index, 4).Value = as_code(source["PRESTAMO"])
            trace.Cells(index, 5).Value = "PREPARADO"
            trace.Cells(index, 8).Value = utc_now()
            trace.Cells(index, 10).Value = source_sha

        phase = "recalcular_y_guardar_libro"
        excel.CalculateFullRebuild()
        workbook.Save()
        if execution_mode != "PREPARAR":
            phase = {
                "SOLO_LECTURA": "ejecutar_validacion_solo_lectura",
                "NAVEGACION": "ejecutar_navegacion_sin_cambios",
                "EJECUTAR": "ejecutar_macro_3270",
            }[execution_mode]
            evidence_directory = workbook_path.parent / "evidencias_3270"
            evidence_directory.mkdir(parents=True, exist_ok=True)
            macro_procedure = {
                "SOLO_LECTURA": READ_ONLY_MACRO_PROCEDURE,
                "NAVEGACION": NAVIGATION_ONLY_MACRO_PROCEDURE,
                "EJECUTAR": settings.macro_procedure,
            }[execution_mode]
            excel.Run(f"'{workbook.Name}'!{macro_procedure}", 2, last_row, str(evidence_directory), settings.session_name)
            phase = {
                "SOLO_LECTURA": "guardar_resultado_validacion",
                "NAVEGACION": "guardar_resultado_navegacion",
                "EJECUTAR": "guardar_resultado_macro_3270",
            }[execution_mode]
            excel.CalculateFull()
            workbook.Save()

        phase = "leer_resultados_macro_3270"
        output: list[dict[str, Any]] = []
        for index, source in enumerate(rows, start=2):
            result = clean(sheet.Cells(index, 25).Value)
            trace_status = clean(trace.Cells(index, 6).Value) or ("PENDIENTE" if execution_mode == "PREPARAR" else "SIN_RESULTADO")
            output.append({
                "id_caso": clean(source["ID_CASO"]),
                "prestamo": as_code(source["PRESTAMO"]),
                "fila_masivo": index,
                "estado": trace_status,
                "resultado_3270": result,
            })
        return output
    except Macro3270Error:
        raise
    except Exception as error:
        raise Macro3270Error(f"EXCEL_O_MACRO_3270_ERROR:{phase}:{type(error).__name__}") from error
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=True)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass


def configure_trace_sheet(workbook: Any, lot_id: str, source_sha: str, rows: list[dict[str, Any]]) -> Any:
    headers = (
        "FILA_MASIVO", "ID_LOTE", "ID_CASO", "PRESTAMO", "ESTADO_PREPARACION", "ESTADO_3270", "RESULTADO_3270",
        "PREPARADO_EN", "PROCESADO_EN", "SHA256_RESULTADO_CLOUD",
    )
    try:
        workbook.Worksheets("TRAZA_PAD").Delete()
    except Exception:
        pass
    trace = workbook.Worksheets.Add(After=workbook.Worksheets(workbook.Worksheets.Count))
    trace.Name = "TRAZA_PAD"
    for index, header in enumerate(headers, start=1):
        trace.Cells(1, index).Value = header
    trace.Columns("A:J").AutoFit()
    return trace


def state_path(settings: Settings, lot_id: str, source_sha: str) -> Path:
    return settings.state_directory / f"{safe_name(lot_id)}_{source_sha}.json"


def copy_to_output(settings: Settings, workbook_path: Path) -> Path | None:
    if settings.output_directory is None:
        return None
    settings.output_directory.mkdir(parents=True, exist_ok=True)
    target = settings.output_directory / workbook_path.name
    shutil.copy2(workbook_path, target)
    return target


def process_lot(settings: Settings, lot_id: str, execution_mode: str, force: bool = False) -> dict[str, Any]:
    source = result_workbook(settings, lot_id)
    wait_stable(source, settings.stability_checks, settings.stability_interval_seconds)
    source_sha = sha256(source)
    state_file = state_path(settings, lot_id, source_sha)
    if execution_mode == "EJECUTAR" and state_file.is_file() and not force:
        previous = json.loads(state_file.read_text(encoding="utf-8"))
        if previous.get("status") in {"COMPLETADO_3270", "COMPLETADO_3270_CON_INCIDENCIAS"}:
            return {"outcome": "idempotente", **previous}

    rows = read_ready_rows(source)
    work_dir = settings.work_directory / f"{safe_name(lot_id)}_{source_sha[:12]}"
    work_dir.mkdir(parents=True, exist_ok=True)
    working_book = work_dir / f"Resultado_MASIVO_{safe_name(lot_id)}.xlsm"
    shutil.copy2(settings.template_path, working_book)
    status = {
        "PREPARAR": "PREPARADO_PARA_3270",
        "SOLO_LECTURA": "VALIDANDO_LECTURA_3270",
        "NAVEGACION": "NAVEGANDO_SIN_CAMBIOS_3270",
        "EJECUTAR": "PREPARANDO_3270",
    }[execution_mode]
    record: dict[str, Any] = {
        "schema_version": 1,
        "id_lote": lot_id,
        "source_file": source.name,
        "source_sha256": source_sha,
        "working_book": str(working_book),
        "status": status,
        "created_at": utc_now(),
        "updated_at": utc_now(),
        "filas_para_procesar": len(rows),
        "revision_ia_excluida": True,
        "modo_ejecucion": execution_mode,
    }
    atomic_json(state_file, record)

    try:
        results = populate_workbook(working_book, lot_id, source_sha, rows, settings, execution_mode)
        published = copy_to_output(settings, working_book) if execution_mode != "PREPARAR" else None
        completed_without_incidents = all(item["estado"] == "COMPLETADO" for item in results)
        read_only_validated = all(item["estado"] == "LECTURA_3270_OK" for item in results)
        navigation_validated = all(item["estado"] == "NAVEGACION_3270_OK" for item in results)
        if execution_mode == "EJECUTAR":
            final_status = "COMPLETADO_3270" if completed_without_incidents else "COMPLETADO_3270_CON_INCIDENCIAS"
        elif execution_mode == "SOLO_LECTURA":
            final_status = "VALIDADO_LECTURA_3270" if read_only_validated else "VALIDADO_LECTURA_3270_CON_INCIDENCIAS"
        elif execution_mode == "NAVEGACION":
            final_status = "NAVEGADO_SIN_CAMBIOS_3270" if navigation_validated else "NAVEGADO_SIN_CAMBIOS_3270_CON_INCIDENCIAS"
        else:
            final_status = "PREPARADO_PARA_3270"
        record.update({
            "status": final_status,
            "updated_at": utc_now(),
            "results": results,
            "published_workbook": str(published) if published else None,
        })
        atomic_json(state_file, record)
        result_file = settings.results_directory / f"{safe_name(lot_id)}_{source_sha[:12]}_macro_3270.json"
        atomic_json(result_file, record)
        outcome = {
            "PREPARAR": "preparado",
            "SOLO_LECTURA": "lectura_validada",
            "NAVEGACION": "navegacion_validada",
            "EJECUTAR": "ejecutado",
        }[execution_mode]
        return {"outcome": outcome, **record}
    except Exception as error:
        record.update({
            "status": "ERROR_MACRO_3270",
            "updated_at": utc_now(),
            "technical_error": str(error),
        })
        atomic_json(state_file, record)
        raise


def main() -> int:
    parser = argparse.ArgumentParser(description="Preparador y ejecutor MASIVO/3270 para PAD")
    parser.add_argument("--config", required=True, type=Path)
    commands = parser.add_subparsers(dest="command", required=True)
    for name in ("prepare-lot", "read-lot", "navigate-lot", "run-lot"):
        command = commands.add_parser(name)
        command.add_argument("--lot", required=True)
        command.add_argument("--force", action="store_true")
    arguments = parser.parse_args()
    try:
        settings = load_settings(arguments.config.resolve())
        execution_mode = {
            "prepare-lot": "PREPARAR",
            "read-lot": "SOLO_LECTURA",
            "navigate-lot": "NAVEGACION",
            "run-lot": "EJECUTAR",
        }[arguments.command]
        result = process_lot(settings, arguments.lot, execution_mode=execution_mode, force=arguments.force)
        print(json.dumps({"id_lote": result["id_lote"], "estado": result["status"], "outcome": result["outcome"]}, ensure_ascii=False))
        return 0
    except Macro3270Error as error:
        print(json.dumps({"estado": "ERROR_MACRO_3270", "codigo": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
