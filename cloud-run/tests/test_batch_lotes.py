import io
import os
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

import fitz
from flask import Flask, request
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from batch_api import BatchControlApi, handle_request
from batch_lotes import (
    BATCH_COMPLETED,
    BATCH_DELIVERED,
    BATCH_FAILED,
    BATCH_PROCESSING,
    BATCH_READY,
    BATCH_SOURCE_CHANGED,
    FILE_UPLOADED,
    MAX_BATCH_TOTAL_BYTES,
    MAX_PDF_BYTES,
    BatchError,
    InMemoryBatchStore,
    build_manifest,
    manifest_idempotency_key,
    public_batch_status,
    update_manifest,
)
from batch_worker import process_batch_gcs
from service import CONTROL_COLUMNS, PARA_PROCESAR_COLUMNS, REVIEW_COLUMNS, process_pdf_entries


class FakeBlob:
    def __init__(self, name):
        self.name = name
        self.data = None
        self.size = None
        self.content_type = None

    def exists(self):
        return self.data is not None

    def reload(self):
        if self.data is not None:
            self.size = len(self.data)

    def download_as_bytes(self):
        return self.data

    def upload_from_file(self, source, rewind=False, content_type=None, **_kwargs):
        if rewind:
            source.seek(0)
        self.data = source.read()
        self.size = len(self.data)
        self.content_type = content_type


class FakeBucket:
    def __init__(self):
        self.blobs = {}

    def blob(self, name):
        if name not in self.blobs:
            self.blobs[name] = FakeBlob(name)
        return self.blobs[name]


class FakeStorageClient:
    def __init__(self):
        self.buckets = {}

    def bucket(self, name):
        if name not in self.buckets:
            self.buckets[name] = FakeBucket()
        return self.buckets[name]


def make_pdf(text="Tasación sintética"):
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text, fontsize=10)
    content = document.tobytes()
    document.close()
    return content


def file_item(index, *, etag=None, size=1024, name=None):
    return {
        "item_id": f"onedrive-item-{index}",
        "nombre": name or f"tasacion-{index}.pdf",
        "etag": etag or f"etag-{index}",
        "tamano_bytes": size,
    }


def payload(files):
    return {"carpeta_origen": "/auto-tasaciones/PDFs", "archivos": files}


class BatchManifestTests(unittest.TestCase):
    def test_allows_three_hundred_pdfs_and_rejects_three_hundred_one(self):
        manifest = build_manifest(payload([file_item(index) for index in range(300)]))
        self.assertEqual(manifest["total_pdfs"], 300)
        self.assertEqual(manifest["estado"], "RECIBIDO")
        with self.assertRaisesRegex(BatchError, "máximo de 300"):
            build_manifest(payload([file_item(index) for index in range(301)]))

    def test_rejects_total_over_two_gibibytes(self):
        files = [file_item(1, size=MAX_BATCH_TOTAL_BYTES), file_item(2, size=1)]
        with self.assertRaisesRegex(BatchError, "tamaño total"):
            build_manifest(payload(files))

    def test_individual_pdf_over_limit_is_recorded_as_failed_without_ticket(self):
        manifest = build_manifest(payload([file_item(1, size=MAX_PDF_BYTES + 1)]))
        self.assertEqual(manifest["estado"], BATCH_FAILED)
        self.assertEqual(manifest["archivos"][0]["estado"], BATCH_FAILED)
        self.assertIn("individual", manifest["mensaje"])

    def test_rejects_non_pdf_manifest_entry(self):
        with self.assertRaisesRegex(BatchError, "PDF"):
            build_manifest(payload([file_item(1, name="tasacion.exe")]))

    def test_idempotency_is_order_independent_and_etag_sensitive(self):
        first = build_manifest(payload([file_item(1), file_item(2)]))
        reordered = build_manifest(payload([file_item(2), file_item(1)]))
        changed = build_manifest(payload([file_item(1, etag="etag-nuevo"), file_item(2)]))
        self.assertEqual(first["clave_idempotencia"], reordered["clave_idempotencia"])
        self.assertEqual(first["id_lote"], reordered["id_lote"])
        self.assertNotEqual(first["clave_idempotencia"], changed["clave_idempotencia"])
        self.assertNotIn("drive_id", first)

    def test_records_an_immutable_start_timestamp_when_the_batch_is_registered(self):
        with patch("batch_lotes.utc_timestamp", return_value="2026-10-08T12:00:00Z"):
            manifest = build_manifest(payload([file_item(1)]))

        self.assertEqual(manifest["fecha_inicio"], "2026-10-08T12:00:00Z")
        self.assertEqual(manifest["fecha_fin"], "")

    def test_initial_validation_failure_has_a_finished_lifecycle(self):
        with patch("batch_lotes.utc_timestamp", return_value="2026-10-08T12:00:00Z"):
            manifest = build_manifest(payload([file_item(1, size=MAX_PDF_BYTES + 1)]))

        self.assertEqual(manifest["estado"], BATCH_FAILED)
        self.assertEqual(manifest["fecha_inicio"], "2026-10-08T12:00:00Z")
        self.assertEqual(manifest["fecha_fin"], "2026-10-08T12:00:00Z")


class BatchControlApiTests(unittest.TestCase):
    def setUp(self):
        self.store = InMemoryBatchStore()
        self.storage = FakeStorageClient()
        self.api = BatchControlApi(
            self.store,
            upload_bucket="tasaciones-prueba",
            storage_client=self.storage,
            signed_url_factory=lambda name: f"https://signed.invalid/{name}",
            job_launcher=lambda _manifest: "operations/prueba",
        )

    def _create_one_file_batch(self, *, index=1):
        content = make_pdf()
        manifest, created = self.api.create_batch(payload([file_item(index, size=len(content))]))
        self.assertTrue(created)
        return manifest, content

    def _processing_one_file_batch(self, *, index=1):
        """Build a batch through its valid transitions up to EN_PROCESO."""
        manifest, content = self._create_one_file_batch(index=index)
        item = manifest["archivos"][0]
        self.api.upload_ticket(manifest["id_lote"], item["id_archivo"])
        blob = self.storage.bucket("tasaciones-prueba").blob(item["objeto_gcs"])
        blob.data = content
        blob.content_type = "application/pdf"
        self.api.confirm_upload(manifest["id_lote"], item["id_archivo"], {"etag_confirmado": item["etag"]})
        self.api.start_batch(manifest["id_lote"])
        return self.store.get(manifest["id_lote"])

    def test_identical_manifest_reuses_same_batch(self):
        first, created = self.api.create_batch(payload([file_item(1)]))
        second, second_created = self.api.create_batch(payload([file_item(1)]))
        self.assertTrue(created)
        self.assertFalse(second_created)
        self.assertEqual(first["id_lote"], second["id_lote"])

    def test_does_not_start_partial_batch(self):
        manifest, _ = self._create_one_file_batch()
        with self.assertRaisesRegex(BatchError, "todos los PDFs confirmados"):
            self.api.start_batch(manifest["id_lote"])

    def test_upload_ticket_and_confirmation_moves_batch_to_ready(self):
        manifest, content = self._create_one_file_batch()
        item = manifest["archivos"][0]
        ticket = self.api.upload_ticket(manifest["id_lote"], item["id_archivo"])
        self.assertEqual(ticket["encabezados_carga"], {"Content-Type": "application/pdf"})
        self.assertNotIn("objeto", ticket)
        blob = self.storage.bucket("tasaciones-prueba").blob(item["objeto_gcs"])
        blob.data = content
        blob.content_type = "application/pdf"
        status = self.api.confirm_upload(manifest["id_lote"], item["id_archivo"], {"etag_confirmado": item["etag"]})
        self.assertEqual(status["estado"], BATCH_READY)
        self.assertEqual(status["pdfs_cargados"], 1)
        started = self.api.start_batch(manifest["id_lote"])
        self.assertEqual(started["estado"], BATCH_PROCESSING)

    def test_http_api_registers_and_reads_batch_status(self):
        app = Flask(__name__)
        with app.test_request_context("/v1/lotes", method="POST", json=payload([file_item(1)])):
            response = handle_request(request, self.api)
        self.assertEqual(response.status_code, 201)
        body = response.get_json()
        self.assertTrue(body["id_lote"].startswith("TAS-"))
        self.assertEqual(body["archivos"][0]["item_id"], "onedrive-item-1")
        self.assertTrue(body["archivos"][0]["id_archivo"].startswith("ARC-"))
        with app.test_request_context(f"/v1/lotes/{body['id_lote']}", method="GET"):
            response = handle_request(request, self.api)
        self.assertEqual(response.status_code, 200)
        status = response.get_json()
        self.assertEqual(status["estado"], "RECIBIDO")
        self.assertTrue(status["fecha_inicio"])
        self.assertEqual(status["fecha_fin"], "")
        self.assertGreaterEqual(status["duracion_segundos"], 0)
        self.assertEqual(status["archivos"], [{
            "id_archivo": body["archivos"][0]["id_archivo"],
            "item_id": "onedrive-item-1",
            "nombre": "tasacion-1.pdf",
            "etag": "etag-1",
            "tamano_bytes": 1024,
            "estado": "PENDIENTE",
        }])
        self.assertNotIn("objeto_gcs", status["archivos"][0])

    def test_delivery_freezes_the_persisted_duration(self):
        with patch("batch_lotes.utc_timestamp", return_value="2026-10-08T12:00:00Z"):
            manifest = self._processing_one_file_batch()

        def mark_completed(current):
            current["estado"] = BATCH_COMPLETED
            current["resultado_disponible"] = True

        with patch("batch_lotes.utc_timestamp", return_value="2026-10-08T12:03:42Z"):
            update_manifest(self.store, manifest["id_lote"], mark_completed)
            self.api.confirm_delivery(manifest["id_lote"])
        delivered = self.store.get(manifest["id_lote"])
        self.assertEqual(delivered["fecha_fin"], "2026-10-08T12:03:42Z")
        self.assertEqual(public_batch_status(delivered)["duracion_segundos"], 222)

        with patch("batch_lotes.utc_timestamp", return_value="2026-10-08T12:10:00Z"):
            update_manifest(self.store, manifest["id_lote"], lambda _current: None)
        self.assertEqual(
            self.store.get(manifest["id_lote"])["fecha_fin"],
            "2026-10-08T12:03:42Z",
        )

    def test_failure_freezes_the_persisted_duration(self):
        with patch("batch_lotes.utc_timestamp", return_value="2026-10-08T12:00:00Z"):
            manifest, _ = self.api.create_batch(payload([file_item(1)]))

        def mark_failed(current):
            current["estado"] = BATCH_FAILED
            current["mensaje"] = "Fallo de prueba"

        with patch("batch_lotes.utc_timestamp", return_value="2026-10-08T12:01:15Z"):
            update_manifest(self.store, manifest["id_lote"], mark_failed)
        failed = self.store.get(manifest["id_lote"])
        self.assertEqual(failed["fecha_fin"], "2026-10-08T12:01:15Z")
        self.assertEqual(public_batch_status(failed)["duracion_segundos"], 75)

    def test_http_api_records_and_rejects_oversized_pdf(self):
        app = Flask(__name__)
        with app.test_request_context(
            "/v1/lotes", method="POST",
            json=payload([file_item(8, size=MAX_PDF_BYTES + 1)]),
        ):
            response = handle_request(request, self.api)
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.get_json()["estado"], BATCH_FAILED)

    def test_http_api_rejects_invalid_json_body(self):
        app = Flask(__name__)
        with app.test_request_context("/v1/lotes", method="POST", data="{", content_type="application/json"):
            response = handle_request(request, self.api)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["codigo"], "MANIFIESTO_INVALIDO")

    def test_http_api_requires_configured_control_token(self):
        app = Flask(__name__)
        previous = os.environ.get("BATCH_CONTROL_API_TOKEN")
        os.environ["BATCH_CONTROL_API_TOKEN"] = "secreto-sintetico\n"
        try:
            with app.test_request_context("/v1/lotes", method="POST", json=payload([file_item(9)])):
                denied = handle_request(request, self.api, require_authentication=True)
            self.assertEqual(denied.status_code, 401)
            with app.test_request_context(
                "/v1/lotes", method="POST", json=payload([file_item(9)]),
                headers={"X-Batch-Control-Token": "secreto-sintetico"},
            ):
                accepted = handle_request(request, self.api, require_authentication=True)
            self.assertEqual(accepted.status_code, 201)
        finally:
            if previous is None:
                os.environ.pop("BATCH_CONTROL_API_TOKEN", None)
            else:
                os.environ["BATCH_CONTROL_API_TOKEN"] = previous

    def test_result_ticket_is_available_only_after_completion(self):
        manifest = self._processing_one_file_batch(index=2)
        with self.assertRaisesRegex(BatchError, "resultado del lote"):
            self.api.result_ticket(manifest["id_lote"])

        result = self.storage.bucket("tasaciones-prueba").blob(manifest["resultado_objeto"])
        result.data = b"xlsx-sintetico"

        def mark_completed(current):
            current["estado"] = BATCH_COMPLETED

        update_manifest(self.store, manifest["id_lote"], mark_completed)
        with self.assertRaisesRegex(BatchError, "resultado del lote"):
            self.api.result_ticket(manifest["id_lote"])

        def mark_result_available(current):
            current["resultado_disponible"] = True

        update_manifest(self.store, manifest["id_lote"], mark_result_available)

        ticket = self.api.result_ticket(manifest["id_lote"])
        self.assertEqual(ticket["id_lote"], manifest["id_lote"])
        self.assertIn(manifest["resultado_objeto"], ticket["url_descarga"])
        self.assertNotIn("objeto", ticket)

    def test_confirm_delivery_only_changes_a_completed_available_result(self):
        manifest, _ = self._create_one_file_batch()
        with self.assertRaisesRegex(BatchError, "resultado del lote"):
            self.api.confirm_delivery(manifest["id_lote"])
        self.assertEqual(self.store.get(manifest["id_lote"])["estado"], "RECIBIDO")

        manifest = self._processing_one_file_batch(index=2)
        result = self.storage.bucket("tasaciones-prueba").blob(manifest["resultado_objeto"])
        result.data = b"xlsx-sintetico"

        def mark_completed(current):
            current["estado"] = BATCH_COMPLETED
            current["resultado_disponible"] = True

        update_manifest(self.store, manifest["id_lote"], mark_completed)
        delivered = self.api.confirm_delivery(manifest["id_lote"])
        self.assertEqual(delivered["estado"], BATCH_DELIVERED)
        self.assertEqual(self.store.get(manifest["id_lote"])["estado"], BATCH_DELIVERED)

    def test_manifest_rejects_backward_state_transitions(self):
        manifest = self._processing_one_file_batch()

        def mark_completed(current):
            current["estado"] = BATCH_COMPLETED
            current["resultado_disponible"] = True

        update_manifest(self.store, manifest["id_lote"], mark_completed)
        with self.assertRaisesRegex(BatchError, "transición de estado"):
            update_manifest(
                self.store,
                manifest["id_lote"],
                lambda current: current.update({"estado": BATCH_PROCESSING}),
            )

        self.api.confirm_delivery(manifest["id_lote"])
        with self.assertRaisesRegex(BatchError, "transición de estado"):
            update_manifest(
                self.store,
                manifest["id_lote"],
                lambda current: current.update({"estado": BATCH_COMPLETED}),
            )

    def test_failed_job_launch_marks_the_batch_failed_instead_of_leaving_processing(self):
        manifest, content = self._create_one_file_batch()
        item = manifest["archivos"][0]
        self.api.upload_ticket(manifest["id_lote"], item["id_archivo"])
        blob = self.storage.bucket("tasaciones-prueba").blob(item["objeto_gcs"])
        blob.data = content
        blob.content_type = "application/pdf"
        self.api.confirm_upload(manifest["id_lote"], item["id_archivo"], {"etag_confirmado": item["etag"]})
        failing_api = BatchControlApi(
            self.store,
            upload_bucket="tasaciones-prueba",
            storage_client=self.storage,
            signed_url_factory=lambda _name: "https://signed.invalid/upload",
            job_launcher=lambda _manifest: (_ for _ in ()).throw(
                BatchError("Job no disponible", code="JOB_NO_INICIADO", status=503)
            ),
        )

        with self.assertRaisesRegex(BatchError, "Job no disponible"):
            failing_api.start_batch(manifest["id_lote"])
        saved = self.store.get(manifest["id_lote"])
        self.assertEqual(saved["estado"], BATCH_FAILED)
        self.assertEqual(saved["mensaje"], "No fue posible iniciar el Cloud Run Job")

    def test_missing_gcs_object_marks_batch_failed(self):
        manifest, _ = self._create_one_file_batch()
        item = manifest["archivos"][0]
        self.api.upload_ticket(manifest["id_lote"], item["id_archivo"])
        with self.assertRaisesRegex(BatchError, "no existe"):
            self.api.confirm_upload(manifest["id_lote"], item["id_archivo"], {"etag_confirmado": item["etag"]})
        self.assertEqual(self.store.get(manifest["id_lote"])["estado"], BATCH_FAILED)

    def test_changed_etag_marks_batch_as_source_changed(self):
        manifest, _ = self._create_one_file_batch()
        item = manifest["archivos"][0]
        self.api.upload_ticket(manifest["id_lote"], item["id_archivo"])
        with self.assertRaisesRegex(BatchError, "cambió"):
            self.api.confirm_upload(manifest["id_lote"], item["id_archivo"], {"etag_confirmado": "etag-diferente"})
        saved = self.store.get(manifest["id_lote"])
        self.assertEqual(saved["estado"], BATCH_SOURCE_CHANGED)
        self.assertEqual(saved["archivos"][0]["estado"], BATCH_SOURCE_CHANGED)

    def test_wrong_gcs_size_marks_batch_failed(self):
        manifest, _ = self._create_one_file_batch()
        item = manifest["archivos"][0]
        self.api.upload_ticket(manifest["id_lote"], item["id_archivo"])
        blob = self.storage.bucket("tasaciones-prueba").blob(item["objeto_gcs"])
        blob.data = b"%PDF-incompleto"
        blob.content_type = "application/pdf"
        with self.assertRaisesRegex(BatchError, "tamaño"):
            self.api.confirm_upload(manifest["id_lote"], item["id_archivo"], {"etag_confirmado": item["etag"]})
        self.assertEqual(self.store.get(manifest["id_lote"])["estado"], BATCH_FAILED)


class OrchestrationClaimTests(unittest.TestCase):
    """Pruebas del claim que protege al trigger de eventos duplicados."""

    def setUp(self):
        self.store = InMemoryBatchStore()
        self.storage = FakeStorageClient()
        self.job_calls = []
        self.api = BatchControlApi(
            self.store,
            upload_bucket="tasaciones-prueba",
            storage_client=self.storage,
            signed_url_factory=lambda name: f"https://signed.invalid/{name}",
            job_launcher=lambda manifest: self.job_calls.append(manifest["id_lote"]) or "operations/prueba",
        )

    def _batch(self, *indices):
        # These orchestration tests exercise manifest transitions, not PDF
        # parsing. A fixed payload keeps the GCS-size verification deterministic.
        files = [file_item(index, size=len(b"%PDF-orchestration-test")) for index in indices]
        manifest, created = self.api.create_batch(payload(files))
        self.assertTrue(created)
        return manifest

    def _confirm_all_files(self, manifest):
        bucket = self.storage.bucket("tasaciones-prueba")
        for item in manifest["archivos"]:
            content = b"%PDF-orchestration-test"
            self.api.upload_ticket(manifest["id_lote"], item["id_archivo"])
            blob = bucket.blob(item["objeto_gcs"])
            blob.data = content
            blob.content_type = "application/pdf"
            self.api.confirm_upload(manifest["id_lote"], item["id_archivo"], {"etag_confirmado": item["etag"]})

    def test_one_pdf_claim_is_atomic_and_duplicate_event_is_occupied(self):
        manifest = self._batch(1)
        first = self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-a"})
        duplicate = self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-b"})
        self.assertTrue(first["reclamado"])
        self.assertEqual(first["resultado"], "RECLAMADO")
        self.assertFalse(duplicate["reclamado"])
        self.assertEqual(duplicate["resultado"], "OCUPADO")
        persisted = self.store.get(manifest["id_lote"])["orquestacion"]
        self.assertEqual(persisted["intentos"], 1)
        self.assertEqual(persisted["propietario"], "run-a")

    def test_repeated_execution_id_reuses_claim_without_new_attempt(self):
        manifest = self._batch(1)
        self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-a"})
        repeated = self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-a"})
        self.assertTrue(repeated["reclamado"])
        self.assertEqual(repeated["resultado"], "REUTILIZADO")
        self.assertEqual(self.store.get(manifest["id_lote"])["orquestacion"]["intentos"], 1)

    def test_http_claim_and_renew_routes_require_the_same_execution_owner(self):
        manifest = self._batch(1)
        app = Flask(__name__)
        with app.test_request_context(
            f"/v1/lotes/{manifest['id_lote']}/orquestacion/reclamar",
            method="POST",
            json={"id_ejecucion": "run-a"},
        ):
            claimed = handle_request(request, self.api)
        self.assertEqual(claimed.status_code, 200)
        self.assertTrue(claimed.get_json()["reclamado"])
        with app.test_request_context(
            f"/v1/lotes/{manifest['id_lote']}/orquestacion/renovar",
            method="POST",
            json={"id_ejecucion": "run-b"},
        ):
            denied = handle_request(request, self.api)
        self.assertEqual(denied.status_code, 409)
        self.assertEqual(denied.get_json()["codigo"], "ORQUESTACION_NO_PROPIETARIA")

    def test_multiple_pdfs_and_double_start_launch_exactly_one_job(self):
        manifest = self._batch(1, 2)
        claim = self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-a"})
        self.assertTrue(claim["reclamado"])
        self._confirm_all_files(manifest)
        first = self.api.start_batch(manifest["id_lote"])
        repeated = self.api.start_batch(manifest["id_lote"])
        self.assertEqual(first["estado"], BATCH_PROCESSING)
        self.assertEqual(repeated["estado"], BATCH_PROCESSING)
        self.assertEqual(self.job_calls, [manifest["id_lote"]])

    def test_expired_or_released_claim_can_be_recovered_by_a_new_execution(self):
        manifest = self._batch(1)
        with patch("batch_api.utc_timestamp", return_value="2026-10-09T10:00:00Z"):
            self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-a"})
        with patch("batch_api.utc_timestamp", return_value="2026-10-09T10:01:00Z"):
            self.assertEqual(
                self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-b"})["resultado"],
                "OCUPADO",
            )
        with patch("batch_api.utc_timestamp", return_value="2026-10-09T13:01:00Z"):
            recovered = self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-b"})
        self.assertTrue(recovered["reclamado"])
        self.assertEqual(recovered["resultado"], "RECLAMADO")
        self.assertEqual(self.store.get(manifest["id_lote"])["orquestacion"]["intentos"], 2)

    def test_release_records_retry_reason_without_demoting_batch_state(self):
        manifest = self._batch(1)
        self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-a"})
        released = self.api.release_orchestration(
            manifest["id_lote"],
            {"id_ejecucion": "run-a", "motivo": "TIEMPO_DE_ESPERA_AGOTADO"},
        )
        self.assertTrue(released["liberado"])
        saved = self.store.get(manifest["id_lote"])
        self.assertEqual(saved["estado"], "RECIBIDO")
        self.assertEqual(saved["orquestacion"]["fase"], "REINTENTO_REQUERIDO")
        self.assertIn("TIEMPO_DE_ESPERA_AGOTADO", saved["mensaje"])
        self.assertTrue(self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-b"})["reclamado"])

    def test_interrupted_upload_reuses_verified_gcs_object_then_confirms(self):
        manifest = self._batch(1)
        item = manifest["archivos"][0]
        self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-a"})
        first_ticket = self.api.upload_ticket(manifest["id_lote"], item["id_archivo"])
        self.assertTrue(first_ticket["requiere_carga"])
        blob = self.storage.bucket("tasaciones-prueba").blob(item["objeto_gcs"])
        blob.data = b"%PDF-orchestration-test"
        blob.content_type = "application/pdf"
        recovered_ticket = self.api.upload_ticket(manifest["id_lote"], item["id_archivo"])
        self.assertFalse(recovered_ticket["requiere_carga"])
        self.assertIn(item["objeto_gcs"], recovered_ticket["url_carga"])
        status = self.api.confirm_upload(manifest["id_lote"], item["id_archivo"], {"etag_confirmado": item["etag"]})
        self.assertEqual(status["estado"], BATCH_READY)

    def test_double_delivery_is_idempotent_and_already_delivered_is_not_claimed(self):
        manifest = self._batch(1)
        self._confirm_all_files(manifest)
        self.api.start_batch(manifest["id_lote"])

        def completed(current):
            current["estado"] = BATCH_COMPLETED
            current["resultado_disponible"] = True

        update_manifest(self.store, manifest["id_lote"], completed)
        self.assertEqual(self.api.confirm_delivery(manifest["id_lote"])["estado"], BATCH_DELIVERED)
        self.assertEqual(self.api.confirm_delivery(manifest["id_lote"])["estado"], BATCH_DELIVERED)
        already_delivered = self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-after"})
        self.assertFalse(already_delivered["reclamado"])
        self.assertEqual(already_delivered["resultado"], "TERMINAL")

    def test_upload_and_job_failures_remain_terminal_for_orchestrator_retries(self):
        manifest = self._batch(1)
        item = manifest["archivos"][0]
        self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-a"})
        self.api.upload_ticket(manifest["id_lote"], item["id_archivo"])
        with self.assertRaisesRegex(BatchError, "no existe"):
            self.api.confirm_upload(manifest["id_lote"], item["id_archivo"], {"etag_confirmado": item["etag"]})
        retry = self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-b"})
        self.assertEqual(self.store.get(manifest["id_lote"])["estado"], BATCH_FAILED)
        self.assertFalse(retry["reclamado"])
        self.assertEqual(retry["resultado"], "TERMINAL")

    def test_job_failure_leaves_a_terminal_manifest_for_a_duplicate_event(self):
        manifest = self._batch(2)
        self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-a"})
        self._confirm_all_files(manifest)
        failing_api = BatchControlApi(
            self.store,
            upload_bucket="tasaciones-prueba",
            storage_client=self.storage,
            signed_url_factory=lambda _name: "https://signed.invalid/upload",
            job_launcher=lambda _manifest: (_ for _ in ()).throw(
                BatchError("Job no disponible", code="JOB_NO_INICIADO", status=503)
            ),
        )
        with self.assertRaisesRegex(BatchError, "Job no disponible"):
            failing_api.start_batch(manifest["id_lote"])
        duplicate = self.api.claim_orchestration(manifest["id_lote"], {"id_ejecucion": "run-b"})
        self.assertEqual(self.store.get(manifest["id_lote"])["estado"], BATCH_FAILED)
        self.assertEqual(duplicate["resultado"], "TERMINAL")


class BatchWorkerTests(unittest.TestCase):
    def setUp(self):
        self.store = InMemoryBatchStore()
        self.storage = FakeStorageClient()
        self.api = BatchControlApi(
            self.store,
            upload_bucket="tasaciones-prueba",
            storage_client=self.storage,
            signed_url_factory=lambda _name: "https://signed.invalid/upload",
            job_launcher=lambda _manifest: "operations/prueba",
        )

    def _confirmed_batch(self):
        first_pdf = make_pdf("PDF sintético uno")
        second_pdf = make_pdf("PDF sintético dos")
        manifest, _ = self.api.create_batch(payload([
            file_item(1, size=len(first_pdf)), file_item(2, size=len(second_pdf)),
        ]))
        bucket = self.storage.bucket("tasaciones-prueba")
        for item, content in zip(manifest["archivos"], (first_pdf, second_pdf)):
            self.api.upload_ticket(manifest["id_lote"], item["id_archivo"])
            blob = bucket.blob(item["objeto_gcs"])
            blob.data = content
            blob.content_type = "application/pdf"
            self.api.confirm_upload(manifest["id_lote"], item["id_archivo"], {"etag_confirmado": item["etag"]})
        self.api.start_batch(manifest["id_lote"])
        return self.store.get(manifest["id_lote"])

    def test_worker_processes_confirmed_pdfs_one_by_one_and_preserves_workbook_contract(self):
        manifest = self._confirmed_batch()
        summary = process_batch_gcs(
            manifest["id_lote"],
            store=self.store,
            storage_client=self.storage,
            bucket_name="tasaciones-prueba",
            expected_input_prefix=f"ingresos/{manifest['id_lote']}/pdfs/",
            expected_output_object=manifest["resultado_objeto"],
            use_environment_reviewer=False,
        )
        self.assertEqual(summary["estado"], BATCH_COMPLETED)
        saved = self.store.get(manifest["id_lote"])
        self.assertEqual(saved["pdfs_procesados"], 2)
        self.assertTrue(saved["resultado_disponible"])
        output = self.storage.bucket("tasaciones-prueba").blob(manifest["resultado_objeto"]).data
        workbook = load_workbook(io.BytesIO(output), data_only=True)
        self.assertEqual(workbook.sheetnames, ["PARA_PROCESAR", "REVISION_IA", "CONTROL"])
        self.assertEqual([cell.value for cell in workbook["PARA_PROCESAR"][1]], PARA_PROCESAR_COLUMNS)
        self.assertEqual([cell.value for cell in workbook["REVISION_IA"][1]], REVIEW_COLUMNS)
        self.assertEqual([cell.value for cell in workbook["CONTROL"][1]], CONTROL_COLUMNS)

    def test_immediate_job_sees_processing_and_completes_the_lifecycle(self):
        """Reproduces a Cloud Run Job starting before the run API returns."""
        observed_states = []

        def immediate_job(manifest):
            observed_states.append(self.store.get(manifest["id_lote"])["estado"])
            process_batch_gcs(
                manifest["id_lote"],
                store=self.store,
                storage_client=self.storage,
                bucket_name="tasaciones-prueba",
                expected_input_prefix=f"ingresos/{manifest['id_lote']}/pdfs/",
                expected_output_object=manifest["resultado_objeto"],
                use_environment_reviewer=False,
            )
            return "operations/inmediata"

        api = BatchControlApi(
            self.store,
            upload_bucket="tasaciones-prueba",
            storage_client=self.storage,
            signed_url_factory=lambda _name: "https://signed.invalid/upload",
            job_launcher=immediate_job,
        )
        content = make_pdf("PDF sintético de carrera")
        manifest, _ = api.create_batch(payload([file_item(7, size=len(content))]))
        item = manifest["archivos"][0]
        self.assertEqual(self.store.get(manifest["id_lote"])["estado"], "RECIBIDO")
        api.upload_ticket(manifest["id_lote"], item["id_archivo"])
        self.assertEqual(self.store.get(manifest["id_lote"])["estado"], "CARGANDO_PDFS")
        blob = self.storage.bucket("tasaciones-prueba").blob(item["objeto_gcs"])
        blob.data = content
        blob.content_type = "application/pdf"
        api.confirm_upload(manifest["id_lote"], item["id_archivo"], {"etag_confirmado": item["etag"]})
        self.assertEqual(self.store.get(manifest["id_lote"])["estado"], BATCH_READY)

        started = api.start_batch(manifest["id_lote"])
        self.assertEqual(observed_states, [BATCH_PROCESSING])
        self.assertEqual(started["estado"], BATCH_COMPLETED)
        self.assertTrue(started["resultado_disponible"])

        app = Flask(__name__)
        with app.test_request_context(f"/v1/lotes/{manifest['id_lote']}", method="GET"):
            response = handle_request(request, api)
        status = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(status["id_lote"], manifest["id_lote"])
        self.assertEqual(status["estado"], BATCH_COMPLETED)
        self.assertTrue(status["resultado_disponible"])

        ticket = api.result_ticket(manifest["id_lote"])
        self.assertEqual(ticket["id_lote"], manifest["id_lote"])
        self.assertEqual(api.confirm_delivery(manifest["id_lote"])["estado"], BATCH_DELIVERED)

    def test_completed_batch_cannot_regress_when_start_is_retried(self):
        manifest = self._confirmed_batch()
        process_batch_gcs(
            manifest["id_lote"],
            store=self.store,
            storage_client=self.storage,
            bucket_name="tasaciones-prueba",
            expected_input_prefix=f"ingresos/{manifest['id_lote']}/pdfs/",
            expected_output_object=manifest["resultado_objeto"],
            use_environment_reviewer=False,
        )
        with self.assertRaisesRegex(BatchError, "todos los PDFs confirmados"):
            self.api.start_batch(manifest["id_lote"])
        saved = self.store.get(manifest["id_lote"])
        self.assertEqual(saved["estado"], BATCH_COMPLETED)
        self.assertTrue(saved["resultado_disponible"])

    def test_stale_worker_cannot_demote_a_delivered_batch(self):
        manifest = self._confirmed_batch()
        process_batch_gcs(
            manifest["id_lote"],
            store=self.store,
            storage_client=self.storage,
            bucket_name="tasaciones-prueba",
            expected_input_prefix=f"ingresos/{manifest['id_lote']}/pdfs/",
            expected_output_object=manifest["resultado_objeto"],
            use_environment_reviewer=False,
        )
        self.api.confirm_delivery(manifest["id_lote"])

        with self.assertRaisesRegex(BatchError, "no está listo"):
            process_batch_gcs(
                manifest["id_lote"],
                store=self.store,
                storage_client=self.storage,
                bucket_name="tasaciones-prueba",
                use_environment_reviewer=False,
            )
        self.assertEqual(self.store.get(manifest["id_lote"])["estado"], BATCH_DELIVERED)

    def test_worker_rejects_manifest_without_all_confirmed_pdfs(self):
        manifest, _ = self.api.create_batch(payload([file_item(1)]))
        # Deliberately corrupt the in-memory fixture; a valid API transition
        # can never start the Job before all PDFs are confirmed.
        self.store._batches[manifest["id_lote"]]["estado"] = BATCH_PROCESSING
        with self.assertRaisesRegex(BatchError, "todos los PDFs confirmados"):
            process_batch_gcs(
                manifest["id_lote"],
                store=self.store,
                storage_client=self.storage,
                bucket_name="tasaciones-prueba",
                use_environment_reviewer=False,
            )

    def test_direct_pdf_entries_keep_distinct_case_ids_for_same_content_and_names(self):
        content = make_pdf()
        _, review_rows, control_rows = process_pdf_entries(
            [("D04.pdf", content), ("D05.pdf", content)],
            use_environment_reviewer=False,
        )
        self.assertEqual(len(review_rows), 2)
        self.assertEqual(len({row["ID_CASO"] for row in control_rows}), 2)


if __name__ == "__main__":
    unittest.main()
