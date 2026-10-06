import io
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import fitz
from flask import Flask, request
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from catalog import (
    lookup_class_code,
    lookup_currency_code,
    lookup_direction_code,
    lookup_location,
    lookup_location_from_text,
    lookup_property_codes,
)
from pdf_extractor import extract_address, extract_pdf, extract_pisos_sotanos
from profile_registry import detect_document_profile, profile_uses_strategy, validate_reference_data
from provider_ocr import extract_braschi_floor_table_ocr
from service import (
    CONTROL_COLUMNS, MACRO_COLUMNS, MAX_ZIP_BYTES, PARA_PROCESAR_COLUMNS, REVIEW_COLUMNS, build_upload_ticket,
    build_workbook, declared_zip_size,
    download_staged_zip, procesar_tasaciones, process_zip, to_macro_row,
)
from tools.validate_correcciones_operador import validate as validate_corrections


def build_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        """Dirección: Avenida Prueba 123, Miraflores, Lima
Tipo de inmueble: Departamento
Valor Comercial
US$ 125,000.00
S/ 475,000.00
Año de construcción: 2018
El edificio consta de 7 pisos y 3 sótanos.
Nro de pisos: 8
Nro de sótanos: 2
Préstamo: 12345678901234567890""",
        fontsize=10,
    )
    content = document.tobytes()
    document.close()
    return content


def build_pdf_with_effective_age() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Edad Efectiva (años)", fontsize=10)
    page.insert_text((72, 88), "12", fontsize=10)
    page.insert_text((72, 110), "Fecha de Expedición 15/08/2025", fontsize=10)
    content = document.tobytes()
    document.close()
    return content


def build_pdf_with_year_in_adjacent_cell() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Año de construcción", fontsize=10)
    page.insert_text((230, 72), "2025", fontsize=10)
    content = document.tobytes()
    document.close()
    return content


def build_pdf_with_minuta_and_table() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Dirección según Minuta", fontsize=10)
    page.insert_text((250, 72), "Avenida Minuta 123, Miraflores, Lima", fontsize=10)
    page.insert_text((72, 96), "Dirección según Inspección ocular", fontsize=10)
    page.insert_text((250, 96), "Calle Inspección 456, Miraflores, Lima", fontsize=10)
    page.insert_text((72, 140), "N° de Pisos", fontsize=10)
    page.insert_text((200, 140), "N° Sótanos", fontsize=10)
    page.insert_text((88, 158), "6", fontsize=10)
    page.insert_text((220, 158), "2", fontsize=10)
    page.insert_text((72, 190), "El edificio consta de 5 pisos y 2 sótanos.", fontsize=10)
    content = document.tobytes()
    document.close()
    return content


def build_pdf_with_pisos_sotanos_pair() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "N° de Pisos/ Sótanos del edificio", fontsize=10)
    page.insert_text((185, 90), "3 / 0", fontsize=10)
    content = document.tobytes()
    document.close()
    return content


def build_pdf_with_azotea_not_counted_as_floor() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "El edificio consta de 5 pisos, mas azotea y 2 sotanos.",
        fontsize=10,
    )
    content = document.tobytes()
    document.close()
    return content


def build_pdf_with_solicitud_construyo_address() -> bytes:
    document = fitz.open()
    primary = document.new_page()
    primary.insert_text((72, 72), "Dirección: Avenida Afilador S/N Mz E Lote 7", fontsize=10)
    appendix = document.new_page()
    appendix.insert_text(
        (72, 72),
        """Solicitud Construyo
2. Dirección del inmueble o ubicación:
AV. AFILADOR S/N MZ. E LT. 7 SECTOR AFILADOR RUPA RUPA - LEONCIO PRADO - HUANUCO
3. Partida electrónica""",
        fontsize=10,
    )
    content = document.tobytes()
    document.close()
    return content


def build_pdf_with_first_construyo() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        """Se trata del primer construyo de una edificación proyectada a 2 pisos y azotea.
Se verificó el terreno sin construcciones.
Fecha de Expedición 29-May-2026
Fecha de Caducidad 29-May-2027""",
        fontsize=10,
    )
    content = document.tobytes()
    document.close()
    return content


def build_pdf_with_positioned_expedition_date() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Fecha de Expedición", fontsize=10)
    page.insert_text((250, 72), "29-May-2026", fontsize=10)
    page.insert_text((72, 90), "Fecha de Caducidad", fontsize=10)
    page.insert_text((250, 90), "29-May-2027", fontsize=10)
    content = document.tobytes()
    document.close()
    return content


def build_pdf_with_geographic_department() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        """Provincia Chiclayo, Departamento Lambayeque.
Se trata de una vivienda unifamiliar.""",
        fontsize=10,
    )
    content = document.tobytes()
    document.close()
    return content


def build_pdf_with_opd_construyo_profile() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        """OP-D
Solicitud Construyo
Tipo de inmueble: Casa
Valor de reposicion
S/ 350,000.00""",
        fontsize=10,
    )
    content = document.tobytes()
    document.close()
    return content


def build_pdf_with_braschi_construyo_floor_table() -> bytes:
    """Plantilla sintética con celdas vacías para ejercitar el OCR local."""
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        """Braschi Tasaciones
Solicitud Construyo
N° de pisos en el edificio
N° de sótanos y/o semisótanos""",
        fontsize=10,
    )
    content = document.tobytes()
    document.close()
    return content


def build_pdf_with_text_braschi_construyo_floor_table() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Braschi Tasaciones\nSolicitud Construyo", fontsize=10)
    page.insert_text((72, 150), "N° de pisos en el edificio", fontsize=10)
    page.insert_text((340, 150), "N° de sótanos y/o semisótanos", fontsize=10)
    page.insert_text((90, 168), "6", fontsize=10)
    page.insert_text((370, 168), "2", fontsize=10)
    content = document.tobytes()
    document.close()
    return content


def build_zip(name: str = "D01.pdf", content: bytes | None = None) -> bytes:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zip_file:
        zip_file.writestr(name, content or build_pdf())
    return archive.getvalue()


def complete_row(extracted: dict) -> tuple[list, list[str]]:
    """Aísla el enrutamiento de las variaciones del parser de direcciones."""
    row = [""] * len(MACRO_COLUMNS)
    values = {
        "TIPO DE INMUEBLE": "DEPARTAMENTO", "VALOR DEL BIEN": "", "MONEDA": "PEN", "IMPORTE": "475,000.00",
        "DIRECCION": "AVENIDA", "DIRECCION1": "PRUEBA", "MUNICIPIO": "WANCHAQ", "DIST_COD": "008",
        "PROV_COD": "01", "DEPT_COD": "08", "CLASE": "2", "PISOS": 7, "SOTANOS": 2,
        "AÑO": extracted.get("Año construccion", ""),
    }
    for field, value in values.items():
        row[MACRO_COLUMNS.index(field)] = value
    return row, []


class FakeReviewer:
    model_name = "gemini-prueba"

    def review(self, _pdf, **_kwargs):
        return {
            "decision": "RESUELTO",
            "reason": "Año identificado en el PDF",
            "corrections": [{
                "field": "Año construccion", "value": "2018", "page": 1,
                "evidence": "Año de construcción: 2018", "rule": "Año explícito",
            }],
        }


class UnsafeReviewer:
    model_name = "gemini-prueba"

    def review(self, _pdf, **_kwargs):
        return {
            "decision": "RESUELTO",
            "reason": "Sin fuente comprobable",
            "corrections": [{"field": "Año construccion", "value": "2018", "page": 1, "evidence": ""}],
        }


class NeverCalledReviewer:
    model_name = "gemini-prueba"

    def review(self, *_args, **_kwargs):
        raise AssertionError("Un conflicto sin regla aprobada no debe enviarse a IA")


class TasacionesServiceTests(unittest.TestCase):
    def test_extracts_core_fields_without_prestamo(self):
        result = extract_pdf(build_pdf(), "D01.pdf")
        self.assertEqual(result["PRESTAMO"], "")
        self.assertEqual(result["Tipo inmueble"], "DEPARTAMENTO")
        self.assertEqual(result["Valor elegido US$"], 125000.0)
        self.assertEqual(result["Año construccion"], 2018)
        self.assertEqual(result["Nro pisos edificio"], 8)
        self.assertNotIn("Préstamo", result["Observacion extraccion"])
        self.assertEqual(result["Perfil plantilla"], "generic-v1")
        self.assertEqual(result["Confianza perfil"], 0.0)

    def test_detects_profile_and_uses_its_declared_anchor_alias(self):
        result = extract_pdf(build_pdf_with_opd_construyo_profile(), "D11.pdf")
        self.assertEqual(result["Perfil plantilla"], "opd-construyo-v1")
        self.assertEqual(result["Version perfil"], "1")
        self.assertGreater(result["Confianza perfil"], 0.0)
        self.assertIn("OP-D", result["Coincidencias perfil"])
        self.assertEqual(result["Valor elegido tipo"], "VALOR DE RECONSTRUCCION")
        self.assertEqual(result["Valor elegido S/"], 350000.0)

    def test_detects_braschi_construyo_profile_and_its_technical_strategy(self):
        profile = detect_document_profile([
            "Braschi Tasaciones", "Solicitud Construyo",
            "N° de pisos en el edificio", "N° de sótanos y/o semisótanos",
        ])
        self.assertEqual(profile.profile_id, "braschi-construyo-v1")
        self.assertEqual(profile.provider_id, "braschi-tasaciones")
        self.assertEqual(profile.provider_source, "PERFIL_TECNICO")
        self.assertTrue(profile_uses_strategy(profile, "braschi_floor_table_ocr"))

    def test_braschi_profile_uses_local_ocr_only_when_text_table_is_missing(self):
        with patch("pdf_extractor.extract_braschi_floor_table_ocr", return_value=(6, 2, 1)) as ocr:
            result = extract_pdf(build_pdf_with_braschi_construyo_floor_table(), "D14.pdf")
        ocr.assert_called_once()
        self.assertEqual(result["Perfil plantilla"], "braschi-construyo-v1")
        self.assertEqual(result["Nro pisos edificio"], 6)
        self.assertEqual(result["Nro sotanos edificio"], 2)
        self.assertEqual(result["Pagina pisos/sotanos"], 1)

    def test_braschi_profile_skips_ocr_when_text_table_is_already_readable(self):
        with patch("pdf_extractor.extract_braschi_floor_table_ocr") as ocr:
            result = extract_pdf(build_pdf_with_text_braschi_construyo_floor_table(), "D15.pdf")
        ocr.assert_not_called()
        self.assertEqual(result["Nro pisos edificio"], 6)
        self.assertEqual(result["Nro sotanos edificio"], 2)

    def test_braschi_table_ocr_requires_one_valid_integer_per_cell(self):
        content = build_pdf_with_braschi_construyo_floor_table()
        with patch("provider_ocr.pytesseract") as tesseract, patch("provider_ocr.Image") as image:
            image.frombytes.return_value = MagicMock()
            tesseract.image_to_string.side_effect = ["6", "2"]
            with fitz.open(stream=content, filetype="pdf") as document:
                self.assertEqual(extract_braschi_floor_table_ocr(document), (6, 2, 1))

        with patch("provider_ocr.pytesseract") as tesseract, patch("provider_ocr.Image") as image:
            image.frombytes.return_value = MagicMock()
            tesseract.image_to_string.side_effect = ["6 7", "2"]
            with fitz.open(stream=content, filetype="pdf") as document:
                self.assertIsNone(extract_braschi_floor_table_ocr(document))

    def test_unknown_text_keeps_generic_profile(self):
        profile = detect_document_profile(["Informe independiente sin firma de plantilla"])
        self.assertEqual(profile.profile_id, "generic-v1")
        self.assertEqual(profile.confidence, 0.0)

    def test_identifies_registered_provider_without_changing_generic_profile(self):
        providers = {"tasadora-ejemplo": ("Tasadora Ejemplo", ("TASACIONES EJEMPLO S.A.C.",))}
        with patch("profile_registry.load_providers", return_value=providers):
            profile = detect_document_profile(["Informe de TASACIONES EJEMPLO S.A.C."])
        self.assertEqual(profile.profile_id, "generic-v1")
        self.assertEqual(profile.provider_id, "tasadora-ejemplo")
        self.assertEqual(profile.provider_name, "Tasadora Ejemplo")

    def test_uses_the_earliest_unambiguous_provider_signature_by_page(self):
        providers = {
            "tasadora-a": ("Tasadora A", ("TASADORA A S.A.C.",)),
            "tasadora-b": ("Tasadora B", ("TASADORA B S.A.C.",)),
        }
        with patch("profile_registry.load_providers", return_value=providers):
            profile = detect_document_profile([
                "Informe de TASADORA A S.A.C.",
                "Anexo que cita TASADORA B S.A.C.",
            ])
        self.assertEqual(profile.provider_id, "tasadora-a")

    def test_keeps_provider_empty_when_the_first_matching_page_has_two_companies(self):
        providers = {
            "tasadora-a": ("Tasadora A", ("TASADORA A S.A.C.",)),
            "tasadora-b": ("Tasadora B", ("TASADORA B S.A.C.",)),
        }
        with patch("profile_registry.load_providers", return_value=providers):
            profile = detect_document_profile(["TASADORA A S.A.C. / TASADORA B S.A.C."])
        self.assertEqual(profile.provider_id, "")
        self.assertEqual(profile.profile_id, "generic-v1")

    def test_known_provider_signatures_keep_generic_extraction_until_a_profile_is_needed(self):
        cases = (
            ("Braschi Tasaciones", "braschi-tasaciones"),
            ("Layseca Asociados", "layseca-asociados"),
            ("Tinsa", "tinsa-peru"),
            ("Valortec", "valortec-tasaciones"),
            ("IMAX", "imax-ingenieria-maxima"),
            ("IMAX Ingeniería Máxima", "imax-ingenieria-maxima"),
            ("Ingeniería Máxima E.I.R.L.", "imax-ingenieria-maxima"),
            ("EV Inmobiliaria Barrenechea S.A.C.", "ev-inmobiliaria-barrenechea"),
            ("EV Inmobiliaria Barrenechea Sociedad Anónima Cerrada", "ev-inmobiliaria-barrenechea"),
            ("Quantum Valuaciones S.A.C.", "quantum-valuaciones"),
        )
        for signature, provider_id in cases:
            with self.subTest(provider_id=provider_id):
                profile = detect_document_profile([signature])
                self.assertEqual(profile.profile_id, "generic-v1")
                self.assertEqual(profile.provider_id, provider_id)

    def test_provider_ocr_identification_is_local_metadata_only(self):
        with patch("pdf_extractor.iter_provider_ocr_texts", return_value=("Tinsa",)):
            result = extract_pdf(build_pdf(), "D12.pdf")
        self.assertEqual(result["Tasadora id"], "tinsa-peru")
        self.assertEqual(result["Tasadora detectada"], "Tinsa")
        self.assertEqual(result["Origen tasadora"], "OCR_LOCAL")
        self.assertEqual(result["Perfil plantilla"], "generic-v1")

    def test_ambiguous_provider_ocr_keeps_generic_profile(self):
        with patch("pdf_extractor.iter_provider_ocr_texts", return_value=("Tinsa", "Valortec")):
            result = extract_pdf(build_pdf(), "D13.pdf")
        self.assertEqual(result["Tasadora id"], "")
        self.assertEqual(result["Origen tasadora"], "")
        self.assertEqual(result["Perfil plantilla"], "generic-v1")

    def test_profile_and_provider_catalog_pass_reference_data_validation(self):
        self.assertEqual(validate_reference_data(), [])

    def test_reference_data_validator_rejects_a_shared_provider_signature(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            reference = Path(temporary_directory)
            profiles = reference / "profiles"
            profiles.mkdir()
            (profiles / "generic-v1.json").write_text(
                '{"id":"generic-v1","version":"1","default":true,"match":{},"extraction":{}}',
                encoding="utf-8",
            )
            providers = reference / "tasadoras.json"
            providers.write_text(
                '{"providers":['
                '{"id":"tasadora-a","name":"Tasadora A","aliases":["FIRMA MUESTRA"]},'
                '{"id":"tasadora-b","name":"Tasadora B","aliases":["FIRMA MUESTRA"]}'
                ']}',
                encoding="utf-8",
            )
            issues = validate_reference_data(profiles, providers)
        self.assertTrue(any("más de una tasadora" in issue for issue in issues))

    def test_prioritizes_minuta_and_characteristics_table(self):
        content = build_pdf_with_minuta_and_table()
        with fitz.open(stream=content, filetype="pdf") as document:
            address, page_address = extract_address(document)
            pisos, sotanos, page_pisos = extract_pisos_sotanos(document)
        self.assertEqual(address, "Avenida Minuta 123, Miraflores, Lima")
        self.assertEqual(page_address, 1)
        self.assertEqual((pisos, sotanos, page_pisos), (6, 2, 1))

    def test_extracts_pisos_and_sotanos_from_single_pair_cell(self):
        with fitz.open(stream=build_pdf_with_pisos_sotanos_pair(), filetype="pdf") as document:
            self.assertEqual(extract_pisos_sotanos(document), (3, 0, 1))

    def test_does_not_count_azotea_as_an_additional_floor(self):
        with fitz.open(stream=build_pdf_with_azotea_not_counted_as_floor(), filetype="pdf") as document:
            self.assertEqual(extract_pisos_sotanos(document), (5, 2, 1))

    def test_uses_solicitud_construyo_only_for_missing_administrative_location(self):
        result = extract_pdf(build_pdf_with_solicitud_construyo_address(), "D01.pdf")
        self.assertEqual(
            result["Direccion extraida"],
            "AV. AFILADOR S/N MZ. E LT. 7 SECTOR AFILADOR RUPA RUPA - LEONCIO PRADO - HUANUCO",
        )
        self.assertEqual(result["Pagina direccion"], 2)

    def test_derives_zero_age_and_year_for_documented_first_construyo(self):
        result = extract_pdf(build_pdf_with_first_construyo(), "D09.pdf")
        self.assertEqual(result["Nro pisos edificio"], 2)
        self.assertEqual(result["Nro sotanos edificio"], 0)
        self.assertEqual(result["Edad efectiva"], 0)
        self.assertEqual(result["Año expedicion"], 2026)
        self.assertEqual(result["Año construccion"], 2026)

    def test_uses_date_aligned_with_expedition_not_caducidad(self):
        result = extract_pdf(build_pdf_with_positioned_expedition_date(), "D03.pdf")
        self.assertEqual(result["Año expedicion"], 2026)
        self.assertEqual(result["Pagina año expedicion"], 1)

    def test_ignores_geographic_department_when_resolving_property_type(self):
        result = extract_pdf(build_pdf_with_geographic_department(), "D04.pdf")
        self.assertEqual(result["Tipo inmueble"], "CASA")

    def test_builds_historical_pen_row_without_prestamo_or_location_code(self):
        extracted = {
            "PRESTAMO": "12345678901234567890", "Valor elegido US$": 125000.0,
            "Valor elegido S/": 475000.0, "Nro pisos edificio": 6,
            "Nro sotanos edificio": 2, "Año construccion": 2018,
            "Tipo inmueble": "DEPARTAMENTO",
        }
        parsed = {
            "TIPO VIA 1": "AVENIDA", "DOMICILIO 1": "PRUEBA", "N. EXTERIOR": "123",
            "N. INTERIOR": "", "REFERENCIA": "", "UBICACION TIPO": "URBANIZACION",
            "UBICACION 1": "PRUEBA", "DISTRITO": "WANCHAQ", "DISTRITO_COD": "008",
            "PROVINCIA_COD": "01", "DEPARTAMENTO_COD": "08",
        }
        with patch("service.parse_address", return_value=(parsed, [])), \
             patch("service.lookup_property_codes", return_value={"tipo_inmueble": "D", "masivo": "C", "valor_del_bien": "C"}), \
             patch("service.lookup_currency_code", return_value="PEN"), \
             patch("service.lookup_direction_code", side_effect=lambda value: {"AVENIDA": "AV.", "URBANIZACION": "URB"}.get(value, "")), \
             patch("service.lookup_class_code", return_value="2"):
            row, issues = to_macro_row(extracted)
        self.assertEqual(row[MACRO_COLUMNS.index("PRESTAMO")], "")
        self.assertEqual(row[MACRO_COLUMNS.index("TIPO DE INMUEBLE")], "DEPARTAMENTO")
        self.assertEqual(row[MACRO_COLUMNS.index("VALOR DEL BIEN")], "")
        self.assertEqual(row[MACRO_COLUMNS.index("MONEDA")], "PEN")
        self.assertEqual(row[MACRO_COLUMNS.index("IMPORTE")], "475,000.00")
        self.assertEqual(row[MACRO_COLUMNS.index("DIRECCION")], "AVENIDA")
        self.assertEqual(row[MACRO_COLUMNS.index("UBICACION")], "URBANIZACION")
        self.assertEqual(row[MACRO_COLUMNS.index("UBICACION1")], "PRUEBA")
        self.assertEqual(issues, [])

    def test_derives_construction_year_from_effective_age(self):
        result = extract_pdf(build_pdf_with_effective_age(), "D02.pdf")
        self.assertEqual(result["Edad efectiva"], 12)
        self.assertEqual(result["Año expedicion"], 2025)
        self.assertEqual(result["Año construccion"], 2013)
        self.assertEqual(result["Origen año construccion"], "EDAD EFECTIVA")

    def test_extracts_construction_year_from_adjacent_pdf_cell(self):
        result = extract_pdf(build_pdf_with_year_in_adjacent_cell(), "D02.pdf")
        self.assertEqual(result["Año construccion"], 2025)
        self.assertEqual(result["Origen año construccion"], "AÑO DE CONSTRUCCIÓN")

    def test_uses_only_codes_defined_in_datos_catalog(self):
        self.assertEqual(lookup_property_codes("DEPARTAMENTO"), {
            "tipo_inmueble": "D", "masivo": "C", "valor_del_bien": "C",
        })
        self.assertEqual(lookup_property_codes("CASA"), {
            "tipo_inmueble": "C", "masivo": "N", "valor_del_bien": "N",
        })
        self.assertEqual(lookup_currency_code("USD"), "USD")
        self.assertEqual(lookup_currency_code("PEN"), "PEN")
        self.assertEqual(lookup_direction_code("AVENIDA"), "AV.")
        self.assertEqual(lookup_location("CUSCO", "CUSCO", "WANCHAQ")["DISTRITO_COD"], "008")
        self.assertEqual(lookup_location("LAMBAYEQUE", "CHICLAYO", "JOSE LEONARDO ORTIZ")["DEPARTAMENTO_COD"], "15")
        self.assertEqual(lookup_location("HUANUCO", "LEONCIO PRADO", "RUPA RUPA")["DISTRITO_COD"], "004")
        self.assertEqual(
            lookup_location_from_text("Av. Afilador, Rupa Rupa - Leoncio Prado - Huanuco")["DISTRITO_COD"],
            "004",
        )
        self.assertEqual(lookup_class_code(4), "1")
        self.assertEqual(lookup_class_code(5), "2")
        self.assertEqual(lookup_class_code(11), "3")

    def test_routes_unresolved_case_to_revision_ia_and_preserves_control(self):
        ready_rows, review_rows, control_rows = process_zip(
            build_zip(content=build_pdf_with_effective_age()), use_environment_reviewer=False,
        )
        self.assertEqual(ready_rows, [])
        self.assertEqual(len(review_rows), 1)
        self.assertEqual(review_rows[0]["Estado"], "PENDIENTE_IA")
        self.assertTrue(review_rows[0]["ID_CASO"].startswith("TAS-"))
        self.assertEqual(len(control_rows), 1)
        self.assertEqual(control_rows[0]["Ruta final"], "PENDIENTE_IA")
        self.assertEqual(control_rows[0]["Revision IA ejecutada"], "NO")

    def test_records_profile_metadata_in_control_and_revision(self):
        ready_rows, review_rows, control_rows = process_zip(
            build_zip("D11.pdf", build_pdf_with_opd_construyo_profile()), use_environment_reviewer=False,
        )
        self.assertEqual(ready_rows, [])
        self.assertEqual(review_rows[0]["Perfil plantilla"], "opd-construyo-v1")
        self.assertEqual(control_rows[0]["Perfil plantilla"], "opd-construyo-v1")
        workbook = load_workbook(io.BytesIO(build_workbook(ready_rows, review_rows, control_rows).read()), data_only=True)
        self.assertEqual(
            [cell.value for cell in workbook["REVISION_IA"][1]],
            REVIEW_COLUMNS,
        )
        self.assertEqual(
            [cell.value for cell in workbook["CONTROL"][1]],
            CONTROL_COLUMNS,
        )

    def test_assigns_distinct_case_ids_to_duplicate_pdf_content(self):
        archive = io.BytesIO()
        content = build_pdf_with_effective_age()
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
            bundle.writestr("D04.pdf", content)
            bundle.writestr("D05.pdf", content)
        _, review_rows, control_rows = process_zip(archive.getvalue(), use_environment_reviewer=False)
        self.assertEqual(len(review_rows), 2)
        self.assertEqual(len({row["ID_CASO"] for row in control_rows}), 2)

    def test_ia_result_is_revalidated_before_entering_para_procesar(self):
        extracted = {
            "ID / Codigo PDF": "D99", "PDF_Archivo": "D99.pdf", "Año construccion": "",
            "Observacion extraccion": "Año de construcción y edad efectiva no encontrados",
        }
        with patch("service.extract_pdf", return_value=extracted), patch("service.to_macro_row", side_effect=complete_row):
            ready_rows, review_rows, control_rows = process_zip(build_zip("D99.pdf"), reviewer=FakeReviewer())
        self.assertEqual(len(ready_rows), 1)
        self.assertEqual(ready_rows[0][PARA_PROCESAR_COLUMNS.index("AÑO")], 2018)
        self.assertTrue(ready_rows[0][PARA_PROCESAR_COLUMNS.index("ID_CASO")].startswith("TAS-"))
        self.assertEqual(review_rows, [])
        self.assertEqual(control_rows[0]["Ruta final"], "LISTO_IA_VERIFICADO")
        self.assertEqual(control_rows[0]["Revision IA ejecutada"], "SI")
        self.assertIn("Año construccion", control_rows[0]["Correcciones IA"])

    def test_ia_correction_without_evidence_never_enters_para_procesar(self):
        extracted = {
            "ID / Codigo PDF": "D98", "PDF_Archivo": "D98.pdf", "Año construccion": "",
            "Observacion extraccion": "Año de construcción y edad efectiva no encontrados",
        }
        with patch("service.extract_pdf", return_value=extracted), patch("service.to_macro_row", side_effect=complete_row):
            ready_rows, review_rows, control_rows = process_zip(build_zip("D98.pdf"), reviewer=UnsafeReviewer())
        self.assertEqual(ready_rows, [])
        self.assertEqual(review_rows[0]["Estado"], "REVISION_HUMANA")
        self.assertIn("evidencia de página válida", control_rows[0]["Incidencias de validacion"])

    def test_conflict_without_approved_rule_never_reaches_ai_or_para_procesar(self):
        extracted = {
            "ID / Codigo PDF": "D97", "PDF_Archivo": "D97.pdf", "Año construccion": 2018,
            "Observacion extraccion": "Conflicto pisos/sótanos entre descripción y tabla del PDF",
        }
        with patch("service.extract_pdf", return_value=extracted), patch("service.to_macro_row", side_effect=complete_row):
            ready_rows, review_rows, control_rows = process_zip(build_zip("D97.pdf"), reviewer=NeverCalledReviewer())
        self.assertEqual(ready_rows, [])
        self.assertEqual(review_rows[0]["Estado"], "REGLA_NEGOCIO_PENDIENTE")
        self.assertEqual(control_rows[0]["Revision IA ejecutada"], "NO")

    def test_builds_three_named_tables_for_power_automate(self):
        row, _ = complete_row({"Año construccion": 2018})
        control = {"ID_CASO": "TAS-TEST", "PDF_Archivo": "D01.pdf", "Ruta final": "LISTO_DETERMINISTA"}
        workbook = load_workbook(io.BytesIO(build_workbook([[*row, "TAS-TEST"]], [], [control]).read()), data_only=True)
        self.assertEqual(workbook.sheetnames, ["PARA_PROCESAR", "REVISION_IA", "CONTROL"])
        self.assertEqual(workbook["PARA_PROCESAR"]["B2"].value, "DEPARTAMENTO")
        self.assertIn("tblParaProcesar", workbook["PARA_PROCESAR"].tables)
        self.assertIn("tblRevisionIa", workbook["REVISION_IA"].tables)
        self.assertIn("tblControl", workbook["CONTROL"].tables)

    def test_endpoint_accepts_raw_zip_and_returns_xlsx(self):
        app = Flask(__name__)
        with app.test_request_context("/", method="POST", data=build_zip(), content_type="application/zip"):
            response = procesar_tasaciones(request)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("Resultado_Final.xlsx", response.headers["Content-Disposition"])

    def test_creates_temporary_signed_upload_ticket(self):
        credentials = MagicMock()
        credentials.token = "token-prueba"
        blob = MagicMock()
        blob.generate_signed_url.return_value = "https://storage.example/upload"
        client = MagicMock()
        client.bucket.return_value.blob.return_value = blob
        environment = {
            "GCS_UPLOAD_BUCKET": "tasaciones-prueba",
            "GCS_SIGNING_SERVICE_ACCOUNT": "runtime@example.iam.gserviceaccount.com",
        }
        with patch.dict(os.environ, environment, clear=False), \
             patch("service.storage.Client", return_value=client), \
             patch("service.google.auth.default", return_value=(credentials, "project")):
            ticket = build_upload_ticket({"nombre_archivo": "auto-10.zip", "tamano_bytes": 49_810_784})
        self.assertEqual(ticket["operacion"], "subir_zip")
        self.assertTrue(ticket["objeto"].startswith("ingresos/"))
        self.assertTrue(ticket["objeto"].endswith("/auto-10.zip"))
        self.assertEqual(ticket["encabezados_carga"], {"Content-Type": "application/zip"})
        blob.generate_signed_url.assert_called_once()

    def test_accepts_staged_zip_up_to_ninety_megabytes(self):
        supported_size = 86_955_661
        self.assertLess(supported_size, MAX_ZIP_BYTES)
        self.assertEqual(declared_zip_size(supported_size), supported_size)

    def test_rejects_staged_zip_above_ninety_megabytes(self):
        with self.assertRaisesRegex(ValueError, "tamaño máximo"):
            declared_zip_size(MAX_ZIP_BYTES + 1)

    def test_endpoint_processes_staged_zip(self):
        app = Flask(__name__)
        with patch("service.download_staged_zip", return_value=build_zip("lote/D01.pdf")):
            with app.test_request_context(
                "/", method="POST", json={"operacion": "procesar_carga", "objeto": "ingresos/prueba/auto-10.zip"},
            ):
                response = procesar_tasaciones(request)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_downloads_staged_zip_after_loading_metadata(self):
        raw_zip = build_zip()
        blob = MagicMock()
        blob.exists.return_value = True
        blob.size = None
        blob.download_as_bytes.return_value = raw_zip
        blob.reload.side_effect = lambda: setattr(blob, "size", len(raw_zip))
        client = MagicMock()
        client.bucket.return_value.blob.return_value = blob
        with patch.dict(os.environ, {"GCS_UPLOAD_BUCKET": "tasaciones-prueba"}, clear=False), \
             patch("service.storage.Client", return_value=client):
            content = download_staged_zip("ingresos/prueba/auto-10.zip")
        self.assertEqual(content, raw_zip)
        blob.reload.assert_called_once()

    def test_validates_operator_correction_ledger_and_skips_blank_rows(self):
        ledger = (
            "ID_CASO;PDF_Archivo;Tasadora id;Perfil plantilla;Campo;Valor extraido;Valor final;Pagina;Evidencia;Motivo;Operador;Fecha\n"
            "TAS-123;D11.pdf;;opd-construyo-v1;Valor elegido S/;350000;350000;1;Valor de reposicion;Etiqueta alternativa;operador.prueba;2026-10-05\n"
            ";;;;;;;;;;;\n"
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "correcciones.csv"
            path.write_text(ledger, encoding="utf-8")
            summary, issues = validate_corrections(path)
        self.assertEqual(issues, [])
        self.assertEqual(summary["filas_validas"], 1)
        self.assertEqual(summary["por_perfil"], {"opd-construyo-v1": 1})


if __name__ == "__main__":
    unittest.main()
