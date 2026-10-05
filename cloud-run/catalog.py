"""Lectura del catálogo ``DATOS`` usado por la macro MASIVO.

El libro se conserva junto al servicio para que la imagen desplegada no dependa
de un archivo local del operador. Solo se devuelve un código cuando existe una
coincidencia explícita; nunca se completa con Lima u otro valor por defecto.
"""
from __future__ import annotations

from functools import lru_cache
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook


CATALOG_PATH = Path(__file__).parent / "reference-data" / "catalogos_tasaciones.xlsx"


def normalize(value: object) -> str:
    text = "" if value is None else str(value)
    text = "".join(
        char for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )
    return " ".join(text.upper().split())


@lru_cache(maxsize=1)
def load_catalog() -> dict[str, dict]:
    """Carga códigos geográficos y reglas de la hoja ``DATOS``."""
    if not CATALOG_PATH.exists():
        return {
            "departments": {}, "provinces": {}, "districts": {}, "district_names": {},
            "property_types": {}, "masivo_types": {}, "value_types": {},
            "currencies": {}, "directions": {}, "classes": {},
        }

    workbook = load_workbook(CATALOG_PATH, read_only=True, data_only=True)
    sheet = workbook.active
    departments: dict[str, str] = {}
    provinces: dict[tuple[str, str], str] = {}
    districts: dict[tuple[str, str, str], str] = {}
    district_names: dict[str, set[tuple[str, str, str]]] = {}
    province_department = {}
    current_province_department = ""
    current_province = ""
    current_district_department = ""
    current_district_province = ""
    # Los encabezados de los catálogos M:N y P:Q están en la fila 2, igual
    # que los títulos de las listas geográficas. Los datos comienzan en la 3.
    property_section = normalize(sheet["M2"].value)
    direction_section = normalize(sheet["P2"].value)
    property_types: dict[str, str] = {}
    masivo_types: dict[str, str] = {}
    value_types: dict[str, str] = {}
    currencies: dict[str, str] = {}
    directions: dict[str, str] = {}
    classes: dict[str, str] = {}

    # El catálogo tiene tres listas paralelas: B:C (departamentos), E:G
    # (departamento/provincia/código) e I:J (distritos por provincia).
    for row in sheet.iter_rows(min_row=3, values_only=True):
        department, department_code = row[1], row[2]
        group_department, province, province_code = row[4], row[5], row[6]
        district, district_code = row[8], row[9]

        normalized_department = normalize(department)
        if normalized_department and department_code is not None and str(department_code).strip():
            departments[normalized_department] = str(department_code).strip().zfill(2)

        normalized_group_department = normalize(group_department)
        if normalized_group_department:
            current_province_department = normalized_group_department
        normalized_province = normalize(province)
        if normalized_province and current_province_department and province_code is not None and str(province_code).strip():
            key = (current_province_department, normalized_province)
            provinces[key] = str(province_code).strip().zfill(2)
            province_department.setdefault(normalized_province, set()).add(current_province_department)

        normalized_district = normalize(district)
        district_heading = re.fullmatch(r"DISTRITOS?\s+(.+)", normalized_district)
        if district_heading:
            # El libro usa ambos encabezados, "DISTRITOS <provincia>" y
            # "DISTRITO <provincia>". Ambos definen la lista inmediatamente
            # inferior y no representan un distrito seleccionable.
            current_district_province = district_heading.group(1).strip()
            owners = province_department.get(current_district_province, set())
            current_district_department = next(iter(owners)) if len(owners) == 1 else ""
            continue
        if normalized_district and district_code is not None and str(district_code).strip():
            # La primera lista de distritos de Lima no tiene encabezado propio;
            # conserva el último departamento/provincia declarado en E:G.
            if not current_district_province:
                current_district_department = current_province_department
                current_district_province = normalized_province
            if current_district_department and current_district_province:
                code = str(district_code).strip().zfill(3)
                key = (current_district_department, current_district_province, normalized_district)
                districts[key] = code
                district_names.setdefault(normalized_district, set()).add(key)

        # M:N contiene cuatro bloques: tipo de inmueble, valor, MASIVO y
        # clase. Los códigos se leen desde el mismo catálogo que la macro.
        property_label, property_code = normalize(row[12]), row[13]
        if property_label in {"TIPO DE INMUEBLE", "VALOR DEL BIEN", "MASIVO", "CLASE INMUEBLE"}:
            property_section = property_label
        elif property_section and property_label and property_code is not None and str(property_code).strip():
            code = str(property_code).strip()
            if property_section == "TIPO DE INMUEBLE":
                property_types[property_label] = code
            elif property_section == "VALOR DEL BIEN":
                value_types[property_label] = code
            elif property_section == "MASIVO":
                masivo_types[property_label] = code
            elif property_section == "CLASE INMUEBLE":
                classes[property_label] = code

        # P:Q contiene moneda y abreviaturas de vía/ubicación.
        direction_label, direction_code = normalize(row[15]), row[16]
        if direction_label in {"MONEDA", "DIRECCION"}:
            direction_section = direction_label
        elif direction_section and direction_label and direction_code is not None and str(direction_code).strip():
            code = str(direction_code).strip()
            if direction_section == "MONEDA":
                currencies[direction_label] = code
            elif direction_section == "DIRECCION":
                directions[direction_label] = code

    workbook.close()
    return {
        "departments": departments,
        "provinces": provinces,
        "districts": districts,
        "district_names": district_names,
        "property_types": property_types,
        "masivo_types": masivo_types,
        "value_types": value_types,
        "currencies": currencies,
        "directions": directions,
        "classes": classes,
    }


def lookup_location(department: object, province: object, district: object) -> dict[str, str]:
    """Obtiene códigos de departamento, provincia y distrito sin adivinar."""
    catalog = load_catalog()
    dept = normalize(department)
    prov = normalize(province)
    dist = normalize(district)
    result = {
        "DEPARTAMENTO": dept,
        "PROVINCIA": prov,
        "DISTRITO": dist,
        "DEPARTAMENTO_COD": catalog["departments"].get(dept, ""),
        "PROVINCIA_COD": catalog["provinces"].get((dept, prov), ""),
        "DISTRITO_COD": catalog["districts"].get((dept, prov, dist), ""),
    }

    # Un distrito inequívoco puede resolverse aunque el PDF no indique uno de
    # los niveles superiores; si tiene más de una ubicación se deja vacío.
    if not result["DISTRITO_COD"] and dist:
        matches = catalog["district_names"].get(dist, set())
        if len(matches) == 1:
            dep, pro, _ = next(iter(matches))
            result["DEPARTAMENTO"] = result["DEPARTAMENTO"] or dep
            result["PROVINCIA"] = result["PROVINCIA"] or pro
            result["DEPARTAMENTO_COD"] = catalog["departments"].get(dep, "")
            result["PROVINCIA_COD"] = catalog["provinces"].get((dep, pro), "")
            result["DISTRITO_COD"] = catalog["districts"].get((dep, pro, dist), "")
    return result


def lookup_property_codes(property_type: object) -> dict[str, str]:
    """Obtiene exclusivamente los códigos autorizados por ``DATOS``."""
    item = normalize(property_type)
    catalog = load_catalog()
    value_name = "VALOR COMERCIAL" if item == "DEPARTAMENTO" else "VALOR NUEVO" if item == "CASA" else ""
    return {
        "tipo_inmueble": catalog["property_types"].get(item, ""),
        "masivo": catalog["masivo_types"].get(item, ""),
        "valor_del_bien": catalog["value_types"].get(value_name, ""),
    }


def lookup_currency_code(currency: object) -> str:
    """Convierte USD/PEN al código exacto definido por ``DATOS``."""
    aliases = {"USD": "DOLARES", "PEN": "SOLES"}
    return load_catalog()["currencies"].get(aliases.get(normalize(currency), ""), "")


def lookup_direction_code(direction: object) -> str:
    """Devuelve la abreviatura autorizada o vacío si no está en ``DATOS``."""
    return load_catalog()["directions"].get(normalize(direction), "")


def lookup_class_code(floors: object) -> str:
    """Resuelve la clase con los tramos de ``CLASE INMUEBLE``."""
    try:
        total = int(floors)
    except (TypeError, ValueError):
        return ""
    classes = load_catalog()["classes"]
    if total <= 4:
        return classes.get("INM.HASTA 4 PISOS", "")
    if total <= 10:
        return classes.get("INM.ENTRE 5 Y 10 PISOS", "")
    return classes.get("INM.CON MAS DE 10 PISOS", "")
