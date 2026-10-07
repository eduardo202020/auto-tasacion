"""Endpoint Power Automate -> Cloud Run -> OneDrive para tasaciones.

El XLSX separa la cola operable de las excepciones:

* ``PARA_PROCESAR``: única tabla que Power Automate puede enviar a IBM 3270.
* ``REVISION_IA``: casos sin resolver tras validación determinista e IA.
* ``CONTROL``: trazabilidad completa de todos los PDFs del lote.

Ninguna respuesta de IA llega directamente a la cola de operación: se vuelve a
validar contra el contrato y los catálogos antes de incorporarla.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
from pathlib import Path
import re
from datetime import timedelta
from uuid import uuid4
import zipfile
from typing import Any

import functions_framework
import google.auth
import pandas as pd
from flask import Response
from google.auth.transport.requests import Request as GoogleAuthRequest
from google.cloud import storage
from openpyxl.worksheet.table import Table, TableStyleInfo

import address_parser
from ai_reviewer import AI_FIELDS, AiReviewUnavailable, CaseReviewer, reviewer_from_environment
from catalog import (
    lookup_class_code,
    lookup_currency_code,
    lookup_direction_code,
    lookup_location,
    lookup_location_from_text,
    lookup_property_codes,
)
from pdf_extractor import extract_pdf


MAX_PDFS_PER_BATCH = 300
# El endpoint HTTP/1 de Cloud Run solo recibe hasta 32 MiB. Este valor conserva
# margen para que Power Automate use el camino directo sin recibir un 413.
MAX_DIRECT_ZIP_BYTES = 30 * 1024 * 1024
# Los conectores de Power Automate admiten mensajes de hasta 100 MB. El límite
# de 90 MB deja margen para los encabezados y permite procesar lotes grandes
# que llegan desde OneDrive mediante la URL firmada de Cloud Storage.
MAX_ZIP_BYTES = 90 * 1_000_000
MAX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024
# Límites exclusivos del Cloud Run Job de lotes masivos. No se aplican al
# endpoint HTTP ni a la ruta de carga temporal de 90 MB.
MAX_BATCH_ZIP_BYTES = 2 * 1024 * 1024 * 1024
MAX_BATCH_UNCOMPRESSED_BYTES = 6 * 1024 * 1024 * 1024
MAX_BATCH_PDF_BYTES = 100 * 1024 * 1024
# Gemini admite PDFs inline de hasta 50 MB. Los mayores deben pasar por el
# mecanismo corporativo aprobado (por ejemplo, OCR o File API), nunca forzarse.
MAX_INLINE_AI_PDF_BYTES = 50 * 1024 * 1024
ZIP_CONTENT_TYPE = "application/zip"
GCS_UPLOAD_BUCKET_ENV = "GCS_UPLOAD_BUCKET"
GCS_SIGNING_SERVICE_ACCOUNT_ENV = "GCS_SIGNING_SERVICE_ACCOUNT"
GCS_UPLOAD_PREFIX = "ingresos"
GCS_UPLOAD_TTL_SECONDS = 15 * 60
_SAFE_FILENAME = re.compile(r"[^A-Za-z0-9._-]+")

PARA_PROCESAR_SHEET = "PARA_PROCESAR"
REVISION_IA_SHEET = "REVISION_IA"
CONTROL_SHEET = "CONTROL"

MACRO_COLUMNS = [
    "PRESTAMO", "TIPO DE INMUEBLE", "VALOR DEL BIEN", "MONEDA", "IMPORTE", "COL_F",
    "DIRECCION", "DIRECCION1", "EXTERIOR", "INTERIOR", "REFERENCIA", "COL_L",
    "UBICACION", "UBICACION1", "MUNICIPIO", "DIST_COD", "COL_Q", "PROV_COD",
    "COL_S", "DEPT_COD", "CLASE", "PISOS", "SOTANOS", "AÑO",
]
# Se añade al final para no desplazar las 24 columnas heredadas de MASIVO.
# Ahora que no hay macro, Power Automate usa esta llave para trazar el resultado
# de IBM 3270 contra CONTROL.
PARA_PROCESAR_COLUMNS = [*MACRO_COLUMNS, "ID_CASO"]

# PRESTAMO se obtiene después de la operación en IBM 3270. VALOR DEL BIEN se
# conserva vacío por compatibilidad histórica; IMPORTE contiene el valor PEN.
# Las columnas COL_* son reservadas.
REQUIRED_MACRO_COLUMNS = (
    "TIPO DE INMUEBLE", "MONEDA", "IMPORTE",
    "DIRECCION", "DIRECCION1", "MUNICIPIO", "DIST_COD", "PROV_COD",
    "DEPT_COD", "CLASE", "PISOS", "SOTANOS", "AÑO",
)

CONTROL_COLUMNS = [
    "ID_CASO", "ID / Codigo PDF", "PDF_Archivo", "Tasadora id", "Tasadora detectada", "Origen tasadora",
    "Perfil plantilla", "Version perfil", "Confianza perfil", "Coincidencias perfil",
    "PRESTAMO", "Direccion extraida", "Pagina direccion",
    "Tipo inmueble", "Tipo inmueble texto", "Pagina tipo inmueble", "Valor elegido tipo",
    "Valor elegido US$", "Valor elegido S/", "Valor comercial US$", "Valor comercial S/",
    "Pagina valor comercial", "Valor reconstruccion US$", "Valor reconstruccion S/",
    "Pagina valor reconstruccion", "Año construccion", "Pagina año construccion",
    "Origen año construccion", "Edad efectiva", "Pagina edad efectiva",
    "Año expedicion", "Pagina año expedicion", "Codigo tipo inmueble", "Codigo MASIVO",
    "Codigo valor del bien", "Nro pisos edificio", "Nro sotanos edificio", "Pagina pisos/sotanos",
    "Observacion extraccion", "Campos faltantes", "Incidencias de validacion", "Ruta final",
    "Revision IA ejecutada", "Modelo IA", "Campos enviados a IA", "Motivo IA",
    "Correcciones IA", "Evidencia IA", "Fila PARA_PROCESAR",
]

REVIEW_COLUMNS = [
    "ID_CASO", "PDF_Archivo", "Tasadora id", "Tasadora detectada", "Origen tasadora", "Perfil plantilla",
    "Version perfil", "Confianza perfil", "Coincidencias perfil", "Estado", "Campos faltantes", "Incidencias de validacion",
    "Campos enviados a IA", "Motivo IA", "Correcciones IA", "Evidencia IA", "Modelo IA",
    "Siguiente accion",
]

AI_TARGETS_BY_MACRO_FIELD = {
    "TIPO DE INMUEBLE": ("Tipo inmueble",),
    "MONEDA": ("Valor elegido US$", "Valor elegido S/"),
    "IMPORTE": ("Valor elegido US$", "Valor elegido S/"),
    "DIRECCION": ("Direccion extraida",),
    "DIRECCION1": ("Direccion extraida",),
    "MUNICIPIO": ("Direccion extraida",),
    "DIST_COD": ("Direccion extraida",),
    "PROV_COD": ("Direccion extraida",),
    "DEPT_COD": ("Direccion extraida",),
    "CLASE": ("Nro pisos edificio",),
    "PISOS": ("Nro pisos edificio",),
    "SOTANOS": ("Nro sotanos edificio",),
    "AÑO": ("Año construccion",),
}

TRANSIENT_EXTRACTION_OBSERVATIONS = (
    "Dirección no encontrada", "Tipo de inmueble no determinado", "Valor comercial no encontrado",
    "Valor de reconstrucción no encontrado", "Año de construcción", "Número de pisos no encontrado",
    "Número de sótanos no encontrado",
)
RULES_PATH = Path(__file__).parent / "reference-data" / "reglas_operativas.json"

CORS_HEADERS = {
    "Access-Control-Allow-Origin": "*",
    "Access-Control-Allow-Methods": "POST, OPTIONS",
    "Access-Control-Allow-Headers": "Content-Type",
}


def as_excel_value(value: Any) -> Any:
    return "" if value is None else value


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def unique(items: list[str]) -> list[str]:
    return list(dict.fromkeys(item.strip() for item in items if item and item.strip()))


def json_cell(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ": ")) if value else ""


def build_case_id(filename: str, content: bytes) -> str:
    """Identificador estable por entrada para conciliar una fila con 3270.

    Dos PDFs distintos pueden tener el mismo contenido (por ejemplo, una
    copia de un informe con nombres D04 y D05). El nombre dentro del ZIP forma
    parte de la identidad técnica para que cada fila conserve una llave única.
    """
    identity = str(filename or "").encode("utf-8") + b"\0" + content
    return f"TAS-{hashlib.sha256(identity).hexdigest()[:16].upper()}"


def load_approved_rules() -> list[str]:
    """Lee exclusivamente reglas aprobadas y versionadas para la IA."""
    if not RULES_PATH.exists():
        return []
    try:
        payload = json.loads(RULES_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    rules = payload.get("rules", []) if isinstance(payload, dict) else []
    return [str(rule).strip() for rule in rules if str(rule).strip()]


def choose_importe(extracted: dict[str, Any]) -> tuple[Any, str]:
    """Prioriza el importe en PEN, como en la salida histórica de MASIVO."""
    usd = extracted.get("Valor elegido US$")
    pen = extracted.get("Valor elegido S/")
    if not is_blank(pen):
        return pen, "PEN"
    if not is_blank(usd):
        return usd, "USD"
    return "", ""


def format_importe(value: Any) -> str:
    """Emite el importe en el formato textual heredado de MASIVO.

    La macro histórica consume, por ejemplo, ``735,096.00``. Se entrega como
    texto para conservar los separadores al guardar o exportar el Excel a CSV,
    sin depender de la configuración regional del equipo que lo abra.
    """
    try:
        return f"{float(value):,.2f}"
    except (TypeError, ValueError):
        return ""


def parse_address(extracted: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    observations: list[str] = []
    raw_address = str(extracted.get("Direccion extraida") or "")
    if not raw_address:
        return {}, ["No se puede estructurar la dirección porque no fue extraída"]
    try:
        parsed = address_parser.parse_direccion(raw_address)
    except Exception:
        return {}, ["No se pudo estructurar la dirección"]

    parsed.update(lookup_location(
        parsed.get("DEPARTAMENTO", ""), parsed.get("PROVINCIA", ""), parsed.get("DISTRITO", ""),
    ))
    if not all(parsed.get(field) for field in ("DISTRITO_COD", "PROVINCIA_COD", "DEPARTAMENTO_COD")):
        textual_location = lookup_location_from_text(raw_address)
        for field, value in textual_location.items():
            if value:
                parsed[field] = value
    for field, label in (
        ("DISTRITO_COD", "código de distrito"),
        ("PROVINCIA_COD", "código de provincia"),
        ("DEPARTAMENTO_COD", "código de departamento"),
    ):
        if not parsed.get(field):
            observations.append(f"{label.capitalize()} no encontrado en catálogo")
    return parsed, observations


def to_macro_row(extracted: dict[str, Any]) -> tuple[list[Any], list[str]]:
    """Crea una fila compatible con Power Automate sin clasificarla aún."""
    parsed, observations = parse_address(extracted)
    importe, currency = choose_importe(extracted)
    pisos = extracted.get("Nro pisos edificio", "")
    tipo = str(extracted.get("Tipo inmueble") or "")
    property_codes = lookup_property_codes(tipo)
    currency_code = lookup_currency_code(currency)
    via_code = lookup_direction_code(parsed.get("TIPO VIA 1", ""))
    ubicacion_code = lookup_direction_code(parsed.get("UBICACION TIPO", ""))

    extracted["Codigo tipo inmueble"] = property_codes["tipo_inmueble"]
    extracted["Codigo MASIVO"] = property_codes["masivo"]
    extracted["Codigo valor del bien"] = property_codes["valor_del_bien"]
    if tipo and not property_codes["masivo"]:
        observations.append("Tipo de inmueble sin código autorizado en DATOS")
    if currency and not currency_code:
        observations.append("Moneda sin código autorizado en DATOS")
    if parsed.get("TIPO VIA 1") and not via_code:
        observations.append("Tipo de vía sin abreviatura autorizada en DATOS")
    row = [""] * len(MACRO_COLUMNS)
    row[0] = ""
    # La hoja operativa conserva las etiquetas completas que usaba MASIVO;
    # los códigos del catálogo permanecen auditables en CONTROL.
    row[1] = tipo
    row[2] = ""
    row[3] = currency_code
    row[4] = format_importe(importe)
    row[6] = parsed.get("TIPO VIA 1", "") if via_code else ""
    row[7] = parsed.get("DOMICILIO 1", "")
    row[8] = parsed.get("N. EXTERIOR", "")
    row[9] = parsed.get("N. INTERIOR", "")
    row[10] = parsed.get("REFERENCIA", "")
    row[12] = parsed.get("UBICACION TIPO", "") if ubicacion_code else ""
    row[13] = parsed.get("UBICACION 1", "") if ubicacion_code else ""
    row[14] = parsed.get("DISTRITO", "")
    row[15] = parsed.get("DISTRITO_COD", "")
    row[17] = parsed.get("PROVINCIA_COD", "")
    row[19] = parsed.get("DEPARTAMENTO_COD", "")
    row[20] = lookup_class_code(pisos)
    row[21] = as_excel_value(pisos)
    row[22] = as_excel_value(extracted.get("Nro sotanos edificio", ""))
    row[23] = as_excel_value(extracted.get("Año construccion", ""))
    return row, observations


def persistent_extraction_observations(extracted: dict[str, Any]) -> list[str]:
    """Conserva conflictos y errores, no vacíos que la IA podría corregir."""
    observations = [item.strip() for item in str(extracted.get("Observacion extraccion") or "").split(";") if item.strip()]
    return [item for item in observations if not any(item.startswith(prefix) for prefix in TRANSIENT_EXTRACTION_OBSERVATIONS)]


def assess_case(extracted: dict[str, Any]) -> tuple[list[Any], list[str], list[str]]:
    row, address_observations = to_macro_row(extracted)
    field_values = dict(zip(MACRO_COLUMNS, row))
    missing = [field for field in REQUIRED_MACRO_COLUMNS if is_blank(field_values[field])]
    issues = unique([*persistent_extraction_observations(extracted), *address_observations])
    return row, missing, issues


def ai_targets(missing_fields: list[str], issues: list[str]) -> list[str]:
    targets: list[str] = []
    for field in missing_fields:
        targets.extend(AI_TARGETS_BY_MACRO_FIELD.get(field, ()))
    if any(any(word in issue.lower() for word in ("dirección", "distrito", "provincia", "ubicación")) for issue in issues):
        targets.append("Direccion extraida")
    return [target for target in AI_FIELDS if target in set(targets)]


def has_business_conflict(issues: list[str]) -> bool:
    return any(issue.startswith("Conflicto ") for issue in issues)


def _integer_value(value: object, *, lower: int, upper: int) -> int | None:
    try:
        parsed = int(str(value).strip())
    except (TypeError, ValueError):
        return None
    return parsed if lower <= parsed <= upper else None


def _amount_value(value: object) -> float | None:
    try:
        parsed = float(str(value).strip().replace(",", ""))
    except (TypeError, ValueError):
        return None
    return parsed if parsed > 0 else None


def apply_ai_corrections(
    extracted: dict[str, Any],
    response: dict[str, Any],
    *,
    requested_fields: list[str],
    page_count: int,
) -> tuple[list[dict[str, Any]], list[str]]:
    """Acepta solo correcciones solicitadas, tipadas y con evidencia."""
    accepted: list[dict[str, Any]] = []
    rejected: list[str] = []
    raw_corrections = response.get("corrections", [])
    if not isinstance(raw_corrections, list):
        return accepted, ["La IA devolvió correcciones con una estructura inválida"]

    for correction in raw_corrections:
        if not isinstance(correction, dict):
            rejected.append("La IA devolvió una corrección inválida")
            continue
        field = str(correction.get("field", "")).strip()
        evidence = str(correction.get("evidence", "")).strip()
        page = _integer_value(correction.get("page"), lower=1, upper=page_count)
        if field not in requested_fields or field not in AI_FIELDS:
            rejected.append("La IA propuso un campo no solicitado")
            continue
        if not evidence or page is None:
            rejected.append(f"La corrección IA para {field} no tiene evidencia de página válida")
            continue

        value: Any = correction.get("value", "")
        if field == "Tipo inmueble":
            value = str(value).strip().upper()
            if value not in {"CASA", "DEPARTAMENTO"}:
                rejected.append("La IA propuso un tipo de inmueble no autorizado")
                continue
            extracted["Tipo inmueble texto"] = value
            extracted["Valor elegido tipo"] = "VALOR COMERCIAL" if value == "DEPARTAMENTO" else "VALOR DE RECONSTRUCCION"
        elif field in {"Valor elegido US$", "Valor elegido S/"}:
            value = _amount_value(value)
            if value is None:
                rejected.append(f"La corrección IA para {field} no es un importe positivo")
                continue
        elif field == "Nro pisos edificio":
            value = _integer_value(value, lower=1, upper=150)
            if value is None:
                rejected.append("La corrección IA para pisos no es válida")
                continue
        elif field == "Nro sotanos edificio":
            value = _integer_value(value, lower=0, upper=50)
            if value is None:
                rejected.append("La corrección IA para sótanos no es válida")
                continue
        elif field == "Año construccion":
            value = _integer_value(value, lower=1900, upper=2100)
            if value is None:
                rejected.append("La corrección IA para año de construcción no es válida")
                continue
            extracted["Origen año construccion"] = "IA CON EVIDENCIA"
        else:
            value = str(value).strip()
            if len(value) < 8 or len(value) > 500:
                rejected.append("La corrección IA para dirección no es válida")
                continue

        extracted[field] = value
        accepted.append({
            "field": field, "value": value, "page": page, "evidence": evidence,
            "rule": str(correction.get("rule", "")).strip(),
        })
    return accepted, rejected


def pdf_page_count(content: bytes) -> int:
    import fitz

    with fitz.open(stream=content, filetype="pdf") as document:
        return max(len(document), 1)


def next_action(route: str) -> str:
    return {
        "PENDIENTE_IA": "Configurar o reintentar la revisión IA.",
        "REGLA_NEGOCIO_PENDIENTE": "Validar la regla de negocio y reprocesar.",
        "REVISION_HUMANA": "Completar o confirmar la evidencia antes de reprocesar.",
    }.get(route, "")


def validate_zip_archive(
    archive: zipfile.ZipFile,
    *,
    archive_size: int,
    max_archive_bytes: int,
    max_uncompressed_bytes: int,
    max_pdf_bytes: int | None = None,
) -> list[zipfile.ZipInfo]:
    """Valida el índice de un ZIP ya abierto sin cargarlo completo en memoria.

    ``ZipFile`` solo lee el directorio central para obtener ``infolist``. Esto
    permite reutilizar la validación tanto en el endpoint pequeño como en el
    Cloud Run Job que abre un objeto montado desde Cloud Storage.
    """
    if archive_size <= 0:
        raise ValueError("El ZIP está vacío")
    if archive_size > max_archive_bytes:
        raise ValueError("El ZIP excede el tamaño máximo permitido")

    entries = [
        entry for entry in archive.infolist()
        if not entry.is_dir() and entry.filename.lower().endswith(".pdf")
    ]
    if not entries:
        raise ValueError("El ZIP no contiene archivos PDF")
    if len(entries) > MAX_PDFS_PER_BATCH:
        raise ValueError(f"El ZIP supera el máximo de {MAX_PDFS_PER_BATCH} PDFs")
    if sum(entry.file_size for entry in entries) > max_uncompressed_bytes:
        raise ValueError("El contenido descomprimido excede el tamaño máximo permitido")
    if max_pdf_bytes is not None and any(entry.file_size > max_pdf_bytes for entry in entries):
        raise ValueError("El ZIP contiene un PDF que excede el tamaño máximo permitido")
    if any(entry.flag_bits & 0x1 for entry in entries):
        raise ValueError("El ZIP contiene PDFs cifrados y no puede procesarse")
    return entries


def validate_zip(raw_data: bytes) -> list[zipfile.ZipInfo]:
    """Valida el cuerpo en memoria usado por las rutas HTTP de hasta 90 MB."""
    if not raw_data.startswith(b"PK"):
        raise ValueError("El cuerpo debe ser un archivo ZIP válido")
    with zipfile.ZipFile(io.BytesIO(raw_data)) as archive:
        return validate_zip_archive(
            archive,
            archive_size=len(raw_data),
            max_archive_bytes=MAX_ZIP_BYTES,
            max_uncompressed_bytes=MAX_UNCOMPRESSED_BYTES,
        )


def validate_zip_file(
    zip_path: str | Path,
    *,
    max_archive_bytes: int = MAX_BATCH_ZIP_BYTES,
    max_uncompressed_bytes: int = MAX_BATCH_UNCOMPRESSED_BYTES,
    max_pdf_bytes: int = MAX_BATCH_PDF_BYTES,
) -> list[zipfile.ZipInfo]:
    """Valida un ZIP almacenado como archivo, por ejemplo en un volumen GCS FUSE."""
    path = Path(zip_path)
    if not path.is_file():
        raise ValueError("El ZIP de lote no existe")
    with path.open("rb") as source:
        if source.read(2) != b"PK":
            raise ValueError("El archivo de lote no es un ZIP válido")
    with zipfile.ZipFile(path) as archive:
        return validate_zip_archive(
            archive,
            archive_size=path.stat().st_size,
            max_archive_bytes=max_archive_bytes,
            max_uncompressed_bytes=max_uncompressed_bytes,
            max_pdf_bytes=max_pdf_bytes,
        )


def upload_bucket_name() -> str:
    bucket = os.getenv(GCS_UPLOAD_BUCKET_ENV, "").strip()
    if not bucket:
        raise ValueError("La carga mediante Cloud Storage no está configurada")
    return bucket


def safe_zip_filename(value: Any) -> str:
    filename = Path(str(value or "")).name.strip()
    filename = _SAFE_FILENAME.sub("_", filename)
    if not filename or not filename.lower().endswith(".zip"):
        raise ValueError("El nombre de carga debe terminar en .zip")
    return filename[:120]


def declared_zip_size(value: Any) -> int:
    try:
        size = int(value)
    except (TypeError, ValueError) as error:
        raise ValueError("tamano_bytes debe ser un entero válido") from error
    if size <= 0:
        raise ValueError("tamano_bytes debe ser mayor que cero")
    if size > MAX_ZIP_BYTES:
        raise ValueError("El ZIP excede el tamaño máximo permitido")
    return size


def build_upload_ticket(payload: dict[str, Any]) -> dict[str, Any]:
    """Crea una URL PUT temporal sin exponer el bucket al público.

    Power Automate usa el URL para subir el ZIP directamente a Cloud Storage y
    luego invoca la operación ``procesar_carga`` con el nombre retornado.
    """
    filename = safe_zip_filename(payload.get("nombre_archivo"))
    declared_zip_size(payload.get("tamano_bytes"))
    bucket_name = upload_bucket_name()
    object_name = f"{GCS_UPLOAD_PREFIX}/{uuid4().hex}/{filename}"
    blob = storage.Client().bucket(bucket_name).blob(object_name)

    credentials, _ = google.auth.default(scopes=["https://www.googleapis.com/auth/cloud-platform"])
    credentials.refresh(GoogleAuthRequest())
    service_account_email = os.getenv(GCS_SIGNING_SERVICE_ACCOUNT_ENV, "").strip()
    if not service_account_email:
        raise ValueError("La cuenta de firma de Cloud Storage no está configurada")

    upload_url = blob.generate_signed_url(
        version="v4",
        expiration=timedelta(seconds=GCS_UPLOAD_TTL_SECONDS),
        method="PUT",
        content_type=ZIP_CONTENT_TYPE,
        query_parameters={"ifGenerationMatch": "0"},
        service_account_email=service_account_email,
        access_token=credentials.token,
    )
    return {
        "status": "ok",
        "operacion": "subir_zip",
        "bucket": bucket_name,
        "objeto": object_name,
        "url_carga": upload_url,
        "encabezados_carga": {"Content-Type": ZIP_CONTENT_TYPE},
        "vence_en_segundos": GCS_UPLOAD_TTL_SECONDS,
        "solicitud_proceso": {"operacion": "procesar_carga", "objeto": object_name},
    }


def staged_object_name(value: Any) -> str:
    object_name = str(value or "").strip()
    prefix = f"{GCS_UPLOAD_PREFIX}/"
    if not object_name.startswith(prefix) or ".." in object_name or not object_name.lower().endswith(".zip"):
        raise ValueError("El objeto de carga no es válido")
    return object_name


def download_staged_zip(value: Any) -> bytes:
    """Descarga solo objetos creados para esta ruta de carga temporal."""
    object_name = staged_object_name(value)
    blob = storage.Client().bucket(upload_bucket_name()).blob(object_name)
    if not blob.exists():
        raise ValueError("El ZIP temporal no existe o ya venció")
    # exists() solo confirma presencia; reload() obtiene size y el resto de
    # metadatos antes de decidir si el objeto puede descargarse.
    blob.reload()
    size = int(blob.size or 0)
    if not size:
        raise ValueError("El ZIP temporal está vacío")
    if size > MAX_ZIP_BYTES:
        raise ValueError("El ZIP excede el tamaño máximo permitido")
    raw_data = blob.download_as_bytes()
    if len(raw_data) != size:
        raise ValueError("No fue posible leer el ZIP temporal completo")
    return raw_data


def process_zip(
    zip_source: bytes | str | Path,
    reviewer: CaseReviewer | None = None,
    *,
    use_environment_reviewer: bool = True,
    max_archive_bytes: int = MAX_ZIP_BYTES,
    max_uncompressed_bytes: int = MAX_UNCOMPRESSED_BYTES,
    max_pdf_bytes: int | None = None,
) -> tuple[list[list[Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """Clasifica un ZIP en cola operable, excepciones IA y control.

    Las rutas HTTP siguen entregando ``bytes``. El worker masivo entrega una
    ruta a un ZIP montado desde Cloud Storage; ambas rutas comparten las mismas
    reglas y solo leen el contenido de un PDF por iteración.
    """
    if isinstance(zip_source, bytes):
        if not zip_source.startswith(b"PK"):
            raise ValueError("El cuerpo debe ser un archivo ZIP válido")
        archive_input: io.BytesIO | Path = io.BytesIO(zip_source)
        archive_size = len(zip_source)
    else:
        archive_input = Path(zip_source)
        if not archive_input.is_file():
            raise ValueError("El ZIP de lote no existe")
        with archive_input.open("rb") as source:
            if source.read(2) != b"PK":
                raise ValueError("El archivo de lote no es un ZIP válido")
        archive_size = archive_input.stat().st_size

    ready_rows: list[list[Any]] = []
    review_rows: list[dict[str, Any]] = []
    control_rows: list[dict[str, Any]] = []
    reviewer_error = ""
    if reviewer is None and use_environment_reviewer:
        try:
            reviewer = reviewer_from_environment()
        except AiReviewUnavailable as error:
            reviewer_error = str(error)
    approved_rules = load_approved_rules()

    with zipfile.ZipFile(archive_input) as archive:
        entries = validate_zip_archive(
            archive,
            archive_size=archive_size,
            max_archive_bytes=max_archive_bytes,
            max_uncompressed_bytes=max_uncompressed_bytes,
            max_pdf_bytes=max_pdf_bytes,
        )
        for entry in entries:
            # ``ZipFile.read`` descomprime una sola entrada. Nunca carga el
            # archivo ZIP entero, incluso cuando ``archive_input`` es un path
            # montado desde Cloud Storage.
            content = archive.read(entry)
            filename = entry.filename.rsplit("/", 1)[-1]
            case_id = build_case_id(entry.filename, content)
            if not content.startswith(b"%PDF-"):
                extracted: dict[str, Any] = {
                    "ID / Codigo PDF": filename.rsplit(".", 1)[0],
                    "PDF_Archivo": entry.filename,
                    "Tasadora id": "",
                    "Tasadora detectada": "",
                    "Origen tasadora": "",
                    "Perfil plantilla": "generic-v1",
                    "Version perfil": "1",
                    "Confianza perfil": 0.0,
                    "Coincidencias perfil": "",
                    "Observacion extraccion": "El archivo no tiene una firma PDF válida",
                }
            else:
                extracted = extract_pdf(content, filename)
            extracted["ID_CASO"] = case_id

            deterministic = dict(extracted)
            _, initial_missing, initial_issues = assess_case(deterministic)
            working = dict(deterministic)
            macro_row, missing_fields, issues = assess_case(working)
            route = "LISTO_DETERMINISTA" if not missing_fields and not issues else ""
            review_executed = False
            model_name = ""
            ai_reason = ""
            requested_fields = ai_targets(missing_fields, issues)
            accepted: list[dict[str, Any]] = []
            rejected: list[str] = []

            if not route:
                if has_business_conflict(issues):
                    route = "REGLA_NEGOCIO_PENDIENTE"
                    ai_reason = "Existe un conflicto entre fuentes que no cubre una regla operativa activa"
                elif len(content) > MAX_INLINE_AI_PDF_BYTES:
                    route = "PENDIENTE_IA"
                    ai_reason = "El PDF excede el límite de revisión IA inline y requiere una ruta aprobada"
                elif reviewer is None:
                    route = "PENDIENTE_IA"
                    ai_reason = reviewer_error or "La revisión IA no está habilitada"
                elif not requested_fields:
                    route = "REVISION_HUMANA"
                    ai_reason = "No existe un campo autorizable para revisión IA"
                else:
                    review_executed = True
                    model_name = str(getattr(reviewer, "model_name", "IA"))
                    try:
                        response = reviewer.review(
                            content, case_id=case_id, requested_fields=requested_fields,
                            observations=issues, extracted=deterministic, approved_rules=approved_rules,
                        )
                        ai_reason = str(response.get("reason", "")).strip()
                        accepted, rejected = apply_ai_corrections(
                            working, response, requested_fields=requested_fields, page_count=pdf_page_count(content),
                        )
                        macro_row, missing_fields, issues = assess_case(working)
                        if not missing_fields and not issues:
                            route = "LISTO_IA_VERIFICADO"
                        elif has_business_conflict(issues):
                            route = "REGLA_NEGOCIO_PENDIENTE"
                        else:
                            route = "REVISION_HUMANA"
                    except AiReviewUnavailable as error:
                        route = "PENDIENTE_IA"
                        ai_reason = str(error)
                    except Exception:
                        route = "PENDIENTE_IA"
                        ai_reason = "La revisión IA no pudo completarse"

            ready_row_number: int | str = ""
            if route in {"LISTO_DETERMINISTA", "LISTO_IA_VERIFICADO"}:
                ready_rows.append([*macro_row, case_id])
                ready_row_number = len(ready_rows)
            else:
                review_rows.append({
                    "ID_CASO": case_id, "PDF_Archivo": entry.filename,
                    "Tasadora id": deterministic.get("Tasadora id", ""),
                    "Tasadora detectada": deterministic.get("Tasadora detectada", ""),
                    "Origen tasadora": deterministic.get("Origen tasadora", ""),
                    "Perfil plantilla": deterministic.get("Perfil plantilla", "generic-v1"),
                    "Version perfil": deterministic.get("Version perfil", "1"),
                    "Confianza perfil": deterministic.get("Confianza perfil", 0.0),
                    "Coincidencias perfil": deterministic.get("Coincidencias perfil", ""),
                    "Estado": route,
                    "Campos faltantes": "; ".join(missing_fields),
                    "Incidencias de validacion": "; ".join(issues),
                    "Campos enviados a IA": "; ".join(requested_fields), "Motivo IA": ai_reason,
                    "Correcciones IA": json_cell(accepted),
                    "Evidencia IA": json_cell([
                        {"field": item["field"], "page": item["page"], "evidence": item["evidence"]}
                        for item in accepted
                    ]),
                    "Modelo IA": model_name, "Siguiente accion": next_action(route),
                })

            control = dict(deterministic)
            control.update({
                "ID_CASO": case_id, "Campos faltantes": "; ".join(missing_fields),
                "Incidencias de validacion": "; ".join(unique([*issues, *rejected])), "Ruta final": route,
                "Revision IA ejecutada": "SI" if review_executed else "NO", "Modelo IA": model_name,
                "Campos enviados a IA": "; ".join(requested_fields), "Motivo IA": ai_reason,
                "Correcciones IA": json_cell(accepted),
                "Evidencia IA": json_cell([
                    {"field": item["field"], "page": item["page"], "evidence": item["evidence"]}
                    for item in accepted
                ]),
                "Fila PARA_PROCESAR": ready_row_number,
            })
            control["Observacion extraccion"] = "; ".join(unique([
                str(deterministic.get("Observacion extraccion") or ""),
                *(f"Campo inicial faltante: {field}" for field in initial_missing), *initial_issues,
            ]))
            control_rows.append(control)
    return ready_rows, review_rows, control_rows


def process_zip_file(
    zip_path: str | Path,
    reviewer: CaseReviewer | None = None,
    *,
    use_environment_reviewer: bool = True,
    max_archive_bytes: int = MAX_BATCH_ZIP_BYTES,
    max_uncompressed_bytes: int = MAX_BATCH_UNCOMPRESSED_BYTES,
    max_pdf_bytes: int = MAX_BATCH_PDF_BYTES,
) -> tuple[list[list[Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """Procesa un ZIP de lote masivo desde un archivo seekable.

    Cloud Run Job lo recibe como un path de GCS FUSE. El contrato XLSX y las
    rutas de negocio son idénticos a los del endpoint existente.
    """
    return process_zip(
        zip_path,
        reviewer,
        use_environment_reviewer=use_environment_reviewer,
        max_archive_bytes=max_archive_bytes,
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_pdf_bytes=max_pdf_bytes,
    )


def _add_excel_table(worksheet, name: str) -> None:
    """Crea una tabla de Excel con nombre estable para Power Automate."""
    from openpyxl.utils import get_column_letter

    ref = f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    worksheet.add_table(table)


def _format_worksheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column in worksheet.columns:
        values = ["" if cell.value is None else str(cell.value) for cell in column]
        width = min(max(max(map(len, values), default=10) + 2, 12), 42)
        worksheet.column_dimensions[column[0].column_letter].width = width


def build_workbook(
    ready_rows: list[list[Any]], review_rows: list[dict[str, Any]], control_rows: list[dict[str, Any]],
) -> io.BytesIO:
    """Genera las tres tablas que consume el flujo Microsoft."""
    output = io.BytesIO()
    ready_df = pd.DataFrame(ready_rows, columns=PARA_PROCESAR_COLUMNS)
    review_df = pd.DataFrame(review_rows).reindex(columns=REVIEW_COLUMNS)
    control_df = pd.DataFrame(control_rows).reindex(columns=CONTROL_COLUMNS)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        ready_df.to_excel(writer, sheet_name=PARA_PROCESAR_SHEET, index=False)
        review_df.to_excel(writer, sheet_name=REVISION_IA_SHEET, index=False)
        control_df.to_excel(writer, sheet_name=CONTROL_SHEET, index=False)
        table_names = {
            PARA_PROCESAR_SHEET: "tblParaProcesar", REVISION_IA_SHEET: "tblRevisionIa", CONTROL_SHEET: "tblControl",
        }
        for worksheet in writer.book.worksheets:
            _format_worksheet(worksheet)
            _add_excel_table(worksheet, table_names[worksheet.title])
    output.seek(0)
    return output


def json_error(message: str, status: int) -> Response:
    response = Response(json.dumps({"status": "error", "mensaje": message}, ensure_ascii=False), status=status, mimetype="application/json")
    response.headers.update(CORS_HEADERS)
    return response


def json_response(payload: dict[str, Any], status: int = 200) -> Response:
    response = Response(json.dumps(payload, ensure_ascii=False), status=status, mimetype="application/json")
    response.headers.update(CORS_HEADERS)
    return response


def xlsx_response(raw_data: bytes) -> Response:
    ready_rows, review_rows, control_rows = process_zip(raw_data)
    response = Response(
        build_workbook(ready_rows, review_rows, control_rows).getvalue(), status=200,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response.headers["Content-Disposition"] = 'attachment; filename="Resultado_Final.xlsx"'
    response.headers.update(CORS_HEADERS)
    return response


@functions_framework.http
def procesar_tasaciones(request):
    if request.method == "OPTIONS":
        return "", 204, CORS_HEADERS
    if request.method != "POST":
        return json_error("Método no permitido", 405)
    try:
        if request.mimetype == "application/json":
            payload = request.get_json(silent=False)
            if not isinstance(payload, dict):
                raise ValueError("El cuerpo JSON debe ser un objeto")
            operation = str(payload.get("operacion", "")).strip()
            if operation == "iniciar_carga":
                return json_response(build_upload_ticket(payload), 201)
            if operation == "procesar_carga":
                return xlsx_response(download_staged_zip(payload.get("objeto")))
            raise ValueError("operacion debe ser iniciar_carga o procesar_carga")

        raw_data = request.get_data(cache=False)
        if len(raw_data) > MAX_DIRECT_ZIP_BYTES:
            raise ValueError(
                "El ZIP supera 30 MiB para POST directo; use iniciar_carga y procesar_carga mediante Cloud Storage"
            )
        return xlsx_response(raw_data)
    except ValueError as error:
        return json_error(str(error), 400)
    except Exception:
        return json_error("Error interno al procesar el lote de tasaciones", 500)
